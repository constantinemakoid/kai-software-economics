import os
import shutil
import subprocess
import sys
import tempfile
import uuid
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT.parent))
import econ_model as m

OUT = ROOT / "out"
SOFFICE = os.environ.get("SOFFICE", "/Applications/LibreOffice.app/Contents/MacOS/soffice")
COLORS = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4"]
INK, MUTED, GRID = "#0b0b0b", "#52514e", "#e1e0d9"

THIN = Side(style="thin", color="9a9890")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="dce8f7")
INPUT_FILL = PatternFill("solid", fgColor="fff4d6")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="top", horizontal="center")
RIGHT = Alignment(wrap_text=True, vertical="top", horizontal="right")


def style_table(ws, top, left, bottom, right, header=True):
    for r in range(top, bottom + 1):
        for c in range(left, right + 1):
            cell = ws.cell(r, c)
            cell.border = BORDER
            numeric = isinstance(cell.value, (int, float)) or isinstance(cell.value, str) and cell.value.startswith("=")
            cell.alignment = RIGHT if numeric else WRAP
            cell.font = Font(name="Arial", size=10, bold=header and r == top)
            if header and r == top:
                cell.fill = HEAD_FILL
                cell.alignment = CENTER


def page(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4


def widths(ws, values):
    for i, w in enumerate(values):
        ws.column_dimensions[chr(65 + i)].width = w


def name(wb, key, ref):
    wb.defined_names[key] = DefinedName(key, attr_text=ref)


def sheet_params(wb):
    ws = wb.active
    ws.title = "Параметри"
    ws.append(["Ім'я", "Показник", "Значення", "Одиниця", "Тип", "Пояснення"])
    for i, p in enumerate(m.PARAMS, start=2):
        ws.append([f"p_{p.key}", p.label, p.value, p.unit, p.source, p.note])
        ws.cell(i, 3).fill = INPUT_FILL
        name(wb, f"p_{p.key}", f"'Параметри'!$C${i}")
    style_table(ws, 1, 1, len(m.PARAMS) + 1, 6)
    widths(ws, [20, 42, 11, 14, 7, 52])
    page(ws)


def sheet_value(wb):
    ws = wb.create_sheet("Цінність")
    rows = [
        ("Інтервалів на добу", "=1440/p_interval_min", "шт.", "1440 / інтервал"),
        ("Показів на будівлю за місяць", "=p_meters*B2*p_days_month", "шт.", "лічильники × інтервали × доби"),
        ("Річне споживання будівлі", "=p_area_m2*p_kwh_m2_year", "кВт·год", "площа × питоме споживання"),
        ("Річна вартість електроенергії", "=B4*p_tariff", "EUR", "споживання × тариф"),
        ("Втрати від аномалій, як є", "=p_delta_kw*p_anomaly_hours_day*p_days_as_is*p_anomalies_month*12", "кВт·год/рік", "ΔP × год × доби × події × 12"),
        ("Втрати від аномалій, як буде", "=p_delta_kw*p_anomaly_hours_day*p_days_to_be*p_anomalies_month*12", "кВт·год/рік", "ΔP × год × доби × події × 12"),
        ("Частка втрат у споживанні, як є", "=B6/B4", "%", "втрати / споживання"),
        ("Збережено енергії", "=B6-B7", "кВт·год/рік", "різниця втрат"),
        ("Збережено енергії, частка", "=B9/B4", "%", "основа цілі R2"),
        ("Економія на енергії", "=B9*p_tariff", "EUR/рік", "кВт·год × тариф"),
        ("Економія часу техніка", "=(p_tech_hours_as_is-p_tech_hours_to_be)*p_anomalies_month*12*p_tech_rate", "EUR/рік", "Δгод × події × 12 × ставка"),
        ("Економічна цінність на будівлю", "=B11+B12", "EUR/рік", "сума"),
        ("Економічна цінність на будівлю", "=B13/12", "EUR/міс", "для гіпотези ціни"),
        ("Гіпотеза ціни A2 як частка цінності", "=p_price_hypothesis/B14", "%", "ціна / цінність"),
    ]
    ws.append(["Показник", "Значення", "Одиниця", "Логіка"])
    for r in rows:
        ws.append(list(r))
    for rr in (8, 10, 15):
        ws.cell(rr, 2).number_format = "0.00%"
    for rr in (2, 3, 4, 5, 6, 7, 9, 11, 12, 13):
        ws.cell(rr, 2).number_format = "#,##0"
    ws.cell(14, 2).number_format = "#,##0.00"
    style_table(ws, 1, 1, len(rows) + 1, 4)
    widths(ws, [38, 14, 13, 32])
    page(ws)


def sheet_segments(wb):
    ws = wb.create_sheet("Сегменти")
    ws.append(["Сегмент", "Контекст"] + [c for c, _ in m.CRITERIA] + ["Зважений бал", "Ранг"])
    ws.append(["Вага критерію", ""] + [w for _, w in m.CRITERIA] + ["=SUM(C2:F2)", ""])
    n = len(m.SEGMENTS)
    for i, s in enumerate(m.SEGMENTS, start=3):
        ws.append([f"{s.sid} {s.name}", s.context] + list(s.scores) +
                  [f"=SUMPRODUCT($C$2:$F$2,C{i}:F{i})", f"=RANK(G{i},$G$3:$G${n + 2})"])
        for c in range(3, 7):
            ws.cell(i, c).fill = INPUT_FILL
    style_table(ws, 1, 1, n + 2, 8)
    for c in range(3, 7):
        ws.cell(2, c).fill = INPUT_FILL
    widths(ws, [30, 46, 11, 11, 13, 14, 11, 7])
    page(ws)


def sheet_stakeholders(wb):
    ws = wb.create_sheet("Стейкхолдери")
    ws.append(["Стейкхолдер", "Роль", "Інтерес і цінність", "Вплив 1...5", "Зацікавленість 1...5", "Пріоритет залучення", "Квадрант", "Конфлікт, ризик", "Залучення"])
    for i, s in enumerate(m.STAKEHOLDERS, start=2):
        quad = f'=IF(AND(D{i}>=4,E{i}>=4),"керувати тісно",IF(D{i}>=4,"задовольняти",IF(E{i}>=4,"інформувати","спостерігати")))'
        ws.append([s.name, s.role, s.interest, s.influence, s.interest_level, f"=D{i}*E{i}", quad, s.conflict, s.engagement])
    style_table(ws, 1, 1, len(m.STAKEHOLDERS) + 1, 9)
    widths(ws, [22, 22, 32, 9, 16, 11, 15, 32, 30])
    page(ws)


def sheet_lifecycle(wb):
    ws = wb.create_sheet("Витрати ЖЦ")
    ws.append(["Етап", "Робота або ресурс", "Драйвер", "Прямі/непрямі", "Одноразові/періодичні", "Постійні/змінні", "Незворотні"])
    for c in m.LIFECYCLE:
        ws.append([c.stage, c.item, c.driver, c.direct, c.recurrence, c.behavior, "так" if c.sunk else "ні"])
    n = len(m.LIFECYCLE) + 1
    style_table(ws, 1, 1, n, 7)
    ws.cell(n + 2, 1, "Позицій усього")
    ws.cell(n + 2, 2, f"=COUNTA(A2:A{n})")
    for j, st in enumerate(["Discovery", "Delivery", "Operation", "Retirement"]):
        ws.cell(n + 3 + j, 1, st)
        ws.cell(n + 3 + j, 2, f'=COUNTIF($A$2:$A${n},A{n + 3 + j})')
    style_table(ws, n + 2, 1, n + 6, 2, header=False)
    widths(ws, [13, 46, 34, 12, 14, 12, 11])
    page(ws)


def sheet_variable(wb):
    ws = wb.create_sheet("Змінні витрати")
    ws.append(["Складова v (на будівлю за місяць)", "Формула", "EUR"])
    parts = [
        ("Обробка і зберігання показів", "покази / 1000 × ціна", "='Цінність'!B3/1000*p_cost_per_1000"),
        ("Мобільний зв'язок шлюзу", "SIM", "=p_sim_month"),
        ("Платні сповіщення", "кількість × ціна", "=p_alerts_month*p_alert_price"),
        ("Підтримка", "год × ставка", "=p_support_hours*p_support_rate"),
    ]
    for a, b, c in parts:
        ws.append([a, b, c])
    ws.append(["v, змінні витрати на будівлю", "сума", "=SUM(C2:C5)"])
    ws.append(["F, постійна частина", "базова платформа", "=p_fixed_platform"])
    name(wb, "v_unit", "'Змінні витрати'!$C$6")
    name(wb, "f_fixed", "'Змінні витрати'!$C$7")
    for r in range(2, 8):
        ws.cell(r, 3).number_format = "0.00"
    style_table(ws, 1, 1, 7, 3)
    ws.cell(9, 1, "Кількість будівель N")
    ws.cell(9, 2, "C(N) = F + v × N, EUR/міс")
    ws.cell(9, 3, "C(N) / N, EUR на будівлю")
    ws.cell(9, 4, "Частка F у C(N)")
    for i, n in enumerate(SCALE, start=10):
        ws.cell(i, 1, n)
        ws.cell(i, 2, f"=f_fixed+v_unit*A{i}")
        ws.cell(i, 3, f"=B{i}/A{i}")
        ws.cell(i, 4, f"=f_fixed/B{i}")
        ws.cell(i, 2).number_format = "#,##0.00"
        ws.cell(i, 3).number_format = "0.00"
        ws.cell(i, 4).number_format = "0.0%"
    style_table(ws, 9, 1, 9 + len(SCALE), 4)
    widths(ws, [34, 26, 24, 16])
    page(ws, landscape=False)


def sheet_assumptions(wb):
    ws = wb.create_sheet("Припущення")
    ws.append(["ID", "Категорія", "Припущення", "Вплив", "Невизнач.", "Пріоритет", "Ключ сортування", "Ранг", "Топ 5"])
    n = len(m.ASSUMPTIONS) + 1
    for i, a in enumerate(m.ASSUMPTIONS, start=2):
        ws.append([a.aid, a.category, a.text, a.impact, a.uncertainty, f"=D{i}*E{i}",
                   f"=F{i}*100+D{i}*10+(100-ROW())/1000", f"=RANK(G{i},$G$2:$G${n})", f'=IF(H{i}<=5,"так","ні")'])
        ws.cell(i, 4).fill = INPUT_FILL
        ws.cell(i, 5).fill = INPUT_FILL
        ws.cell(i, 7).number_format = "0.000"
    style_table(ws, 1, 1, n, 9)
    widths(ws, [6, 15, 60, 8, 10, 10, 12, 7, 7])
    page(ws)


def sheet_passport(wb):
    ws = wb.create_sheet("Паспорт", 0)
    ws.append(["Поле", "Зміст"])
    for k, v in passport_rows():
        ws.append([k, v])
    style_table(ws, 1, 1, len(passport_rows()) + 1, 2)
    widths(ws, [28, 110])
    page(ws)


def passport_rows():
    vm = m.value_model()
    return [
        ("Назва продукту", m.PRODUCT_FULL),
        ("Проблема і контекст", "Енергоменеджер багатоорендного офісного центру дізнається про аномальне споживання лише з місячного рахунку, у середньому через 12 діб (П); причину перевитрат встановити складно, бо немає інтервальних даних по точках обліку. Втрати від аномалій оцінено в 3,4% річного споживання (П)."),
        ("Первинний сегмент", "S1: багатоорендні офісні центри класу B/B+ площею 8...25 тис. м² з енергоменеджером або головним інженером і субобліком орендарів"),
        ("Користувач, покупець, платник", "Користувачі: енергоменеджер, технік; вторинний: орендар. Покупець: власник (asset manager) або керуюча компанія. Платник: власник з операційного бюджету, частково відшкодовується орендарями через сервісний платіж"),
        ("Поточна альтернатива", "Місячні рахунки і ручне зняття показів у Excel, обходи техніка, погодинні дані комерційного обліку лише по вводу з затримкою, BMS без аналітики"),
        ("Ціннісна пропозиція", "Медіанний час виявлення аномалії не більше 2 год замість 12 діб і встановлення її причини до рахунку; оцінена цінність " + f"{vm['total_eur_year']:,.0f}".replace(",", " ") + " EUR на будівлю за рік (П)"),
        ("Ключова метрика", "Медіана часу від початку аномалії до її фіксації: база 288 год (П), ціль: медіана не більше 2 год за 3 міс. пілоту; додатковий контроль хвоста розподілу: P80 не більше 4 год"),
        ("Межі MVP", "Одна будівля, лише електроенергія, 30 точок обліку, 11 функцій і 5 NFR; без керування обладнанням, білінгу, мобільного застосунку, погодної моделі"),
        ("Головне припущення", "A1: аномалії, які усуваються після рахунку, становлять не менше 2% річного споживання будівлі (вплив 5, невизначеність 4)"),
        ("Критерій переходу далі", "A1, A2 і A5 підтверджено тестами: перевитрата не менше 2% у 2 з 3 будівель, не менше 3 листів про намір за ціною 130 EUR, дозавантаження не менше 99,5% інтервалів"),
    ]


SCALE = [5, 10, 20, 30, 40, 50, 60, 70, 85]


def build_workbook(path):
    wb = Workbook()
    sheet_params(wb)
    sheet_value(wb)
    sheet_segments(wb)
    sheet_stakeholders(wb)
    sheet_lifecycle(wb)
    sheet_variable(wb)
    sheet_assumptions(wb)
    sheet_passport(wb)
    wb.save(path)


def soffice(src, fmt, outdir):
    prof = Path(tempfile.gettempdir()) / f"lo_{uuid.uuid4().hex}"
    try:
        subprocess.run([SOFFICE, "--headless", "--norestore", f"-env:UserInstallation=file://{prof}",
                        "--convert-to", fmt, "--outdir", str(outdir), str(src)], check=True, capture_output=True, timeout=300)
    finally:
        shutil.rmtree(prof, ignore_errors=True)


def verify(path):
    tmp = Path(tempfile.mkdtemp())
    soffice(path, "xlsx", tmp)
    wb = load_workbook(tmp / path.name, data_only=True)
    vm = m.value_model()
    parts, v = m.unit_variable_cost()
    checks = {
        "Цінність!B13": (wb["Цінність"]["B13"].value, vm["total_eur_year"]),
        "Цінність!B9": (wb["Цінність"]["B9"].value, vm["saved_kwh"]),
        "Цінність!B10": (wb["Цінність"]["B10"].value, vm["saved_share"]),
        "Змінні витрати!C6": (wb["Змінні витрати"]["C6"].value, v),
        "Змінні витрати!B18": (wb["Змінні витрати"]["B18"].value, m.monthly_cost(85)),
        "Сегменти!G3": (wb["Сегменти"]["G3"].value, m.segment_score(m.SEGMENTS[0])),
        "Припущення!H2": (wb["Припущення"]["H2"].value, 1),
        "Витрати ЖЦ!B21": (wb["Витрати ЖЦ"]["B21"].value, len(m.LIFECYCLE)),
    }
    quads = {wb["Стейкхолдери"].cell(r, 1).value: wb["Стейкхолдери"].cell(r, 7).value for r in range(2, len(m.STAKEHOLDERS) + 2)}
    quads_ok = quads == {st.name: m.quadrant(st) for st in m.STAKEHOLDERS}
    print("quadrants match:", quads_ok)
    ok = all(a is not None and abs(a - b) < 1e-6 for a, b in checks.values())
    for k, (a, b) in checks.items():
        print(f"{k}: LibreOffice={a} Python={b}")
    ranks = {wb["Припущення"].cell(r, 1).value: wb["Припущення"].cell(r, 8).value for r in range(2, 14)}
    expected = {a.aid: i for i, a in enumerate(m.ranked_assumptions(), start=1)}
    ok = ok and ranks == expected and quads_ok
    print("ranks match:", ranks == expected)
    shutil.copy(tmp / path.name, OUT / "lab1_recalculated.xlsx")
    shutil.rmtree(tmp, ignore_errors=True)
    return ok


PRINT_VIEWS = {
    "Параметри": ("params", ["F"]),
    "Цінність": ("value", []),
    "Сегменти": ("segments", ["B"]),
    "Змінні витрати": ("variable", []),
    "Припущення": ("assumptions", ["B", "C"]),
    "Стейкхолдери": ("stakeholders", ["B", "C", "H", "I"]),
}

FIXED_ROWS = {"params": 18, "assumptions": 18, "stakeholders": 34}


def split_columns(ws):
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.dimensions import ColumnDimension
    for key, dim in list(ws.column_dimensions.items()):
        if dim.min and dim.max and dim.max > dim.min:
            for idx in range(dim.min, dim.max + 1):
                letter = get_column_letter(idx)
                ws.column_dimensions[letter] = ColumnDimension(ws, index=letter, width=dim.width, hidden=dim.hidden)


def render_sheets(path):
    import pymupdf
    from PIL import Image
    files = {}
    for title, (key, hidden) in PRINT_VIEWS.items():
        tmp = Path(tempfile.mkdtemp())
        wb = load_workbook(path)
        for ws in wb.worksheets:
            ws.sheet_state = "visible" if ws.title == title else "hidden"
        wb.active = wb.sheetnames.index(title)
        ws = wb[title]
        split_columns(ws)
        for col in hidden:
            ws.column_dimensions[col].hidden = True
        if key in FIXED_ROWS:
            for r in range(2, ws.max_row + 1):
                ws.row_dimensions[r].height = FIXED_ROWS[key]
        ws.sheet_properties.pageSetUpPr.fitToPage = False
        ws.page_setup.scale = 100
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.orientation = "portrait"
        src = tmp / f"view_{key}.xlsx"
        wb.save(src)
        soffice(src, "pdf", tmp)
        doc = pymupdf.open(tmp / f"view_{key}.pdf")
        pages = []
        for i, pg in enumerate(doc):
            pix = pg.get_pixmap(dpi=170)
            part = tmp / f"p{i}.png"
            pix.save(part)
            crop(part)
            pages.append(Image.open(part).convert("RGB"))
        width = max(im.width for im in pages)
        canvas = Image.new("RGB", (width, sum(im.height for im in pages)), "white")
        y = 0
        for im in pages:
            canvas.paste(im, (0, y))
            y += im.height
        out = OUT / f"sheet_{key}.png"
        canvas.save(out)
        files[key] = out
        shutil.rmtree(tmp, ignore_errors=True)
    return files


def crop(png):
    from PIL import Image, ImageChops
    im = Image.open(png).convert("RGB")
    bg = Image.new("RGB", im.size, (255, 255, 255))
    box = ImageChops.difference(im, bg).getbbox()
    if box:
        pad = 12
        box = (max(box[0] - pad, 0), max(box[1] - pad, 0), min(box[2] + pad, im.width), min(box[3] + pad, im.height))
        im.crop(box).save(png)


def axes(ax):
    ax.grid(True, color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    for s in ("left", "bottom"):
        ax.spines[s].set_color(MUTED)
    ax.tick_params(colors=MUTED)
    ax.xaxis.label.set_color(INK)
    ax.yaxis.label.set_color(INK)
    ax.title.set_color(INK)


def chart_segments():
    fig, ax = plt.subplots(figsize=(9, 4.2), dpi=200)
    labels = [f"{s.sid} {s.name}" for s in m.SEGMENTS][::-1]
    segs = m.SEGMENTS[::-1]
    left = [0.0] * len(segs)
    for k, (crit, w) in enumerate(m.CRITERIA):
        vals = [w * s.scores[k] for s in segs]
        ax.barh(labels, vals, left=left, color=COLORS[k], label=f"{crit} (вага {w:.2f})".replace(".", ","), height=0.6)
        left = [a + b for a, b in zip(left, vals)]
    for y, total in enumerate(left):
        ax.text(total + 0.05, y, f"{total:.2f}".replace(".", ","), va="center", color=INK, fontsize=10)
    ax.set_xlim(0, 5)
    ax.set_xlabel("Зважений бал, max 5")
    ax.set_title("Вибір первинного сегмента: внесок критеріїв у зважений бал")
    axes(ax)
    ax.grid(axis="y", visible=False)
    ax.legend(loc="lower center", bbox_to_anchor=(0.45, -0.42), ncol=2, frameon=False, fontsize=9)
    fig.tight_layout()
    fig.savefig(OUT / "fig_segments.png")
    plt.close(fig)


LABEL_POS = {
    "Керуюча компанія": (0, -22, "center"),
    "Енергоменеджер": (0, 14, "center"),
    "Постачальник лічильників": (0, -22, "center"),
    "Орендар": (12, 6, "left"),
    "Власник будівлі (asset manager)": (12, 6, "left"),
}


def chart_stakeholders():
    fig, ax = plt.subplots(figsize=(8, 5.6), dpi=200)
    roles = {"користувач": COLORS[0], "покупець": COLORS[1], "партнер": COLORS[2], "зовнішній": COLORS[3]}
    for s in m.STAKEHOLDERS:
        key = next((k for k in roles if s.role.startswith(k)), "користувач")
        if "вторинний" in s.role:
            key = "користувач"
        if s.role.startswith("покупець"):
            key = "покупець"
        pos = (s.interest_level, s.influence)
        ax.scatter(*pos, s=260, color=roles[key], edgecolor="white", zorder=3)
        dx, dy, ha = LABEL_POS.get(s.name, (10, 6, "left"))
        ax.annotate(s.name, pos, xytext=(dx, dy), textcoords="offset points", fontsize=9, color=INK, ha=ha)
    ax.axvline(3.5, color=MUTED, linewidth=0.8, linestyle=":")
    ax.axhline(3.5, color=MUTED, linewidth=0.8, linestyle=":")
    for (x, y, t) in [(0.6, 5.35, "задовольняти"), (3.6, 5.35, "керувати тісно"), (0.6, 0.55, "спостерігати"), (3.6, 0.55, "інформувати")]:
        ax.text(x, y, t, color=MUTED, fontsize=9, style="italic")
    for k, c in roles.items():
        ax.scatter([], [], color=c, s=80, label=k)
    ax.set_xlim(0.5, 6.2)
    ax.set_xticks(range(1, 6))
    ax.set_ylim(0.4, 5.6)
    ax.set_xlabel("Зацікавленість, 1...5")
    ax.set_ylabel("Вплив, 1...5")
    ax.set_title("Карта стейкхолдерів: вплив і зацікавленість")
    axes(ax)
    ax.legend(loc="lower right", frameon=False, fontsize=9)
    fig.tight_layout()
    fig.savefig(OUT / "fig_stakeholders.png")
    plt.close(fig)


def chart_costs():
    ns = list(range(1, 101))
    _, v = m.unit_variable_cost()
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(10, 4.2), dpi=200)
    a1.plot(ns, [m.monthly_cost(n) for n in ns], color=COLORS[0], linewidth=2, label="C(N) = F + v × N")
    a1.plot(ns, [m.P["fixed_platform"]] * len(ns), color=COLORS[1], linewidth=1.5, linestyle="--", label="F = 600 EUR")
    a1.scatter([85], [m.monthly_cost(85)], color=COLORS[0], zorder=3)
    a1.annotate(f"N = 85: {m.monthly_cost(85):,.0f} EUR".replace(",", " "), (85, m.monthly_cost(85)), xytext=(-110, 10), textcoords="offset points", fontsize=9, color=INK)
    a1.set_xlabel("Кількість будівель N")
    a1.set_ylabel("EUR на місяць")
    a1.set_title("Місячні витрати експлуатації")
    a1.set_ylim(0, 2400)
    a1.legend(frameon=False, fontsize=9, loc="upper left")
    axes(a1)
    a2.plot(ns, [m.average_cost(n) for n in ns], color=COLORS[0], linewidth=2, label="C(N) / N")
    a2.plot(ns, [v] * len(ns), color=COLORS[1], linewidth=1.5, linestyle="--", label="v = 16 EUR")
    a2.scatter([85], [m.average_cost(85)], color=COLORS[0], zorder=3)
    a2.annotate(f"N = 85: {m.average_cost(85):.2f} EUR".replace(".", ","), (85, m.average_cost(85)), xytext=(-100, 18), textcoords="offset points", fontsize=9, color=INK)
    a2.set_ylim(0, 120)
    a2.set_xlabel("Кількість будівель N")
    a2.set_ylabel("EUR на будівлю за місяць")
    a2.set_title("Середні витрати на будівлю")
    a2.legend(frameon=False, fontsize=9)
    axes(a2)
    fig.tight_layout()
    fig.savefig(OUT / "fig_costs.png")
    plt.close(fig)


def chart_value():
    vm = m.value_model()
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(10, 4), dpi=200)
    bars = a1.bar(["Як є, 12 діб", "Як буде, 1 доба"], [vm["loss_as_is_kwh"], vm["loss_to_be_kwh"]], color=[COLORS[0], COLORS[1]], width=0.55)
    for b in bars:
        a1.text(b.get_x() + b.get_width() / 2, b.get_height() + 1200, f"{b.get_height():,.0f}".replace(",", " "), ha="center", color=INK, fontsize=10)
    a1.set_ylabel("кВт·год на рік")
    a1.set_title("Втрати від аномалій на будівлю")
    a1.set_ylim(0, 75000)
    axes(a1)
    a1.grid(axis="x", visible=False)
    items = [("Енергія", vm["energy_eur"]), ("Час техніка", vm["labor_eur"]), ("Разом", vm["total_eur_year"])]
    bars = a2.bar([i[0] for i in items], [i[1] for i in items], color=[COLORS[0], COLORS[1], COLORS[2]], width=0.55)
    for b in bars:
        a2.text(b.get_x() + b.get_width() / 2, b.get_height() + 150, f"{b.get_height():,.0f}".replace(",", " "), ha="center", color=INK, fontsize=10)
    a2.set_ylabel("EUR на рік")
    a2.set_title("Оцінена економічна цінність (умовні дані)")
    a2.set_ylim(0, 12500)
    axes(a2)
    a2.grid(axis="x", visible=False)
    fig.tight_layout()
    fig.savefig(OUT / "fig_value.png")
    plt.close(fig)


def chart_assumptions():
    fig, ax = plt.subplots(figsize=(8, 6), dpi=200)
    top = {a.aid for a in m.top_assumptions()}
    groups = {}
    for a in m.ASSUMPTIONS:
        groups.setdefault((a.uncertainty, a.impact), []).append(a.aid)
    for x in range(1, 6):
        for y in range(1, 6):
            ax.add_patch(plt.Rectangle((x - 0.5, y - 0.5), 1, 1, color=COLORS[1] if x * y >= 16 else (COLORS[3] if x * y >= 9 else "#f4f3ee"), alpha=0.18 if x * y >= 9 else 1, linewidth=0))
    for (x, y), ids in groups.items():
        col = COLORS[0] if ids[0] in top else COLORS[1]
        ax.scatter(x, y, s=500 + 220 * (len(ids) - 1), color=col, zorder=3, edgecolor="white")
        ax.text(x, y, "\n".join(ids), ha="center", va="center", color="white", fontsize=8, fontweight="bold", zorder=4)
    from matplotlib.patches import Patch
    ax.add_artist(ax.legend(handles=[Patch(color=COLORS[1], alpha=0.18, label="пріоритет від 16"), Patch(color=COLORS[3], alpha=0.18, label="пріоритет від 9 до 15"), Patch(color="#f4f3ee", label="пріоритет до 8")], loc="upper left", frameon=False, fontsize=9))
    ax.scatter([], [], color=COLORS[0], s=80, label="топ 5 за пріоритетом")
    ax.scatter([], [], color=COLORS[1], s=80, label="інші припущення")
    ax.set_xlim(0.5, 5.5)
    ax.set_ylim(0.5, 5.5)
    ax.set_xticks(range(1, 6))
    ax.set_yticks(range(1, 6))
    ax.set_xlabel("Невизначеність, 1...5")
    ax.set_ylabel("Вплив, 1...5")
    ax.set_title("Реєстр припущень: пріоритет = вплив × невизначеність")
    ax.legend(loc="lower left", frameon=False, fontsize=9)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    fig.tight_layout()
    fig.savefig(OUT / "fig_assumptions.png")
    plt.close(fig)


def dot(name, src):
    p = OUT / f"{name}.dot"
    p.write_text(src, encoding="utf-8")
    subprocess.run(["dot", "-Tpng", "-Gdpi=200", str(p), "-o", str(OUT / f"{name}.png")], check=True)


def diagram_as_is():
    import textwrap
    nodes = "\n".join(f'n{i} [label="{i + 1}. ' + "\\n".join(textwrap.wrap(s, 20)) + '"];' for i, s in enumerate(m.AS_IS))
    dot("fig_as_is", f'''digraph G {{
rankdir=LR; nodesep=0.5; ranksep=0.45;
node [shape=box, style="rounded,filled", fillcolor="#dce8f7", color="#2a78d6", fontname="Arial", fontsize=12, width=2.6, height=0.9, margin="0.15,0.08"];
edge [color="#52514e"];
{nodes}
n6 [fillcolor="#fbe1d6", color="#eb6834"];
{{rank=same; n0; n6;}}
{{rank=same; n1; n5;}}
{{rank=same; n2; n4;}}
n0 -> n1 -> n2 -> n3;
n6 -> n5 [dir=back];
n5 -> n4 [dir=back];
n4 -> n3 [dir=back];
}}''')


def diagram_scenario():
    dot("fig_scenario", '''digraph G {
rankdir=LR; nodesep=0.3; ranksep=0.35; newrank=true;
node [shape=box, style="rounded,filled", fillcolor="#dce8f7", color="#2a78d6", fontname="Arial", fontsize=11, width=2.3, height=0.8];
edge [color="#52514e", fontname="Arial", fontsize=10];
m [label="30 лічильників\\nModbus RTU,\\nімпульсні входи"];
g [label="Шлюз:\\nопитування\\nкожні 15 хв"];
c [label="Зв'язок\\nє?", shape=diamond, fillcolor="#fff4d6", color="#eda100", width=1.2];
b [label="Буфер до 72 год,\\nстатус missing до\\nдозавантаження", fillcolor="#fbe1d6", color="#eb6834"];
q [label="MQTT 5, QoS 1,\\nконтроль якості\\nпоказів"];
s [label="Інтервальне\\nспоживання: різниця\\nнакопичувальних показів"];
a [label="Порівняння з базою:\\n4 інтервали > база + max(3σ; 10%)"];
n [label="Сповіщення\\nз оцінкою кВт·год\\nі EUR на добу"];
i [label="Інцидент:\\nтехнік, причина,\\nзакриття"];
r [label="Повернення до бази,\\nзбережений обсяг,\\nмісячний звіт"];
{rank=same; m; r;}
{rank=same; g; i;}
{rank=same; c; b; n;}
{rank=same; q; a;}
m -> g -> c;
c -> q [label="так"];
c -> b [label="ні"];
b -> q [label="після\\nвідновлення", style=dashed];
q -> s;
a -> s [dir=back];
n -> a [dir=back, label="аномалія"];
i -> n [dir=back];
r -> i [dir=back];
}''')


def main():
    OUT.mkdir(exist_ok=True)
    xlsx = OUT / "lab1_passport_model.xlsx"
    build_workbook(xlsx)
    ok = verify(xlsx)
    print("workbook verified:", ok)
    render_sheets(OUT / "lab1_recalculated.xlsx")
    chart_segments()
    chart_stakeholders()
    chart_costs()
    chart_value()
    chart_assumptions()
    diagram_as_is()
    diagram_scenario()
    if not ok:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
