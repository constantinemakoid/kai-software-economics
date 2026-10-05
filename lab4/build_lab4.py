import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT.parent))
sys.path.insert(0, str(ROOT.parent / "lab1"))
import econ_model as m
import build_lab1 as k

OUT = ROOT / "out"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
plt.rcParams.update({"font.size": 12.5})
COLORS = k.COLORS
INK, MUTED, GRID = k.INK, k.MUTED, k.GRID
TCOL = {"B": "C", "S": "D", "P": "E"}
LAYOUT = {}


def col_widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def sheet_params(wb):
    ws = wb.active
    ws.title = "Параметри ЛР4"
    ws.append(["Ім'я", "Показник", "Значення", "Одиниця", "Тип", "Пояснення"])
    parts, _ = m.unit_variable_cost()
    p4 = m.P4
    vm = m.value_model()
    rows = [
        ("x4_cogs_var", "COGS на будівлю за варіантом", p4["cogs_unit"], "EUR/міс.", "Д", "варіант 14, ЛР № 4"),
        ("x4_gm", "Target gross margin", p4["target_gm"], "частка", "Д", "варіант 14"),
        ("x4_scale", "Масштаб Y3", p4["scale_y3"], "будівель/міс.", "Д", "варіант 14"),
        ("x4_readings", "Обробка і зберігання показів", parts["Обробка і зберігання показів"], "EUR/буд./міс.", "Д", "ЛР № 1: 86 400 показів × 0,075 / 1000"),
        ("x4_sim", "Мобільний зв'язок шлюзу", parts["Мобільний зв'язок шлюзу"], "EUR/буд./міс.", "Д", "ЛР № 1"),
        ("x4_sms_cost", "Собівартість SMS", m.P["alert_price"], "EUR", "Д", "ЛР № 1"),
        ("x4_sms_typ", "SMS на будівлю, типово", m.P["alerts_month"], "шт./міс.", "Д", "ЛР № 1"),
        ("x4_sup_h", "Підтримка на будівлю, типово", m.P["support_hours"], "год/міс.", "Д", "ЛР № 1"),
        ("x4_sup_rate", "Вартість години підтримки", m.P["support_rate"], "EUR/год", "Д", "ЛР № 1"),
        ("x4_sms_incl", "SMS, включені на будівлю", p4["sms_included"], "шт./міс.", "П", "тарифне правило"),
        ("x4_sms_price", "Ціна SMS понад ліміт", p4["sms_price"], "EUR", "П", "не нижче floor 0,20"),
        ("x4_sms_high", "SMS у будівлі з high usage", p4["sms_high"], "шт./міс.", "П", "A3"),
        ("x4_sup_high", "Підтримка будівлі з high usage", p4["support_high_h"], "год/міс.", "П", "A3, нестабільні лічильники"),
        ("x4_fee", "Доплата за нестабільний облік", p4["unstable_fee"], "EUR/буд./міс.", "П", "захисне правило"),
        ("x4_high_share", "Частка будівель з high usage", p4["high_share"], "частка", "П", "A3"),
        ("x4_value", "Цінність на будівлю", vm["total_eur_month"], "EUR/міс.", "Д", "ЛР № 1"),
        ("x4_energy", "Економія енергії на будівлю", vm["energy_eur"] / 12, "EUR/міс.", "Д", "ЛР № 1"),
        ("x4_gain", "Частка економії в gain-sharing", p4["gain_share"], "частка", "П", "альтернатива 2"),
    ]
    for i, r in enumerate(rows, start=2):
        ws.append(list(r))
        ws.cell(i, 3).fill = k.INPUT_FILL
        k.name(wb, r[0], f"'Параметри ЛР4'!$C${i}")
        ws.cell(i, 3).number_format = "0.00" if isinstance(r[2], float) and r[2] % 1 else "#,##0"
    n = len(rows) + 1
    ws.cell(n + 2, 1, "x4_cogs")
    ws.cell(n + 2, 2, "COGS на будівлю з компонентів ЛР № 1")
    ws.cell(n + 2, 3, "=x4_readings+x4_sim+x4_sms_cost*x4_sms_typ+x4_sup_h*x4_sup_rate")
    ws.cell(n + 2, 4, "EUR/буд./міс.")
    ws.cell(n + 2, 5, "Р")
    ws.cell(n + 2, 6, '=IF(ABS(C' + str(n + 2) + '-x4_cogs_var)<0.001,"збігається з варіантом","розбіжність")')
    ws.cell(n + 3, 1, "x4_cogs_high")
    ws.cell(n + 3, 2, "COGS будівлі з high usage")
    ws.cell(n + 3, 3, "=x4_cogs+(x4_sms_high-x4_sms_typ)*x4_sms_cost+(x4_sup_high-x4_sup_h)*x4_sup_rate")
    ws.cell(n + 3, 4, "EUR/буд./міс.")
    ws.cell(n + 3, 5, "Р")
    ws.cell(n + 3, 6, "A3")
    ws.cell(n + 4, 1, "x4_over")
    ws.cell(n + 4, 2, "Overage з будівлі з high usage")
    ws.cell(n + 4, 3, "=x4_fee+MAX(x4_sms_high-x4_sms_incl,0)*x4_sms_price")
    ws.cell(n + 4, 4, "EUR/буд./міс.")
    ws.cell(n + 4, 5, "Р")
    ws.cell(n + 4, 6, "доплата + SMS понад ліміт")
    ws.cell(n + 5, 1, "x4_floor_unit")
    ws.cell(n + 5, 2, "Floor price на будівлю")
    ws.cell(n + 5, 3, "=x4_cogs_var/(1-x4_gm)")
    ws.cell(n + 5, 4, "EUR/буд./міс.")
    ws.cell(n + 5, 5, "Р")
    ws.cell(n + 5, 6, "COGS / (1 − GM)")
    for j, key in enumerate(("x4_cogs", "x4_cogs_high", "x4_over", "x4_floor_unit")):
        k.name(wb, key, f"'Параметри ЛР4'!$C${n + 2 + j}")
        ws.cell(n + 2 + j, 3).number_format = "0.00"
    k.style_table(ws, 1, 1, n + 5, 6)
    ws.cell(n + 2, 6).alignment = Alignment(horizontal="left", vertical="center")
    col_widths(ws, [14, 36, 11, 14, 6, 40])
    k.page(ws, landscape=False)
    LAYOUT["params_extra"] = n + 2


def sheet_assumptions(wb):
    ws = wb.create_sheet("Журнал припущень")
    ws.append(["ID", "Припущення", "Значення", "Підстава", "Що змінює"])
    for a in m.LR4_ASSUMPTIONS:
        ws.append(list(a))
    k.style_table(ws, 1, 1, len(m.LR4_ASSUMPTIONS) + 1, 5)
    col_widths(ws, [6, 24, 50, 44, 16])
    k.page(ws)


def sheet_models(wb):
    ws = wb.create_sheet("Моделі")
    ws.append(["Модель", "Pricing metric"] + m.MODEL_CRITERIA + ["Оцінка 1...5", "Ранг"])
    nm = len(m.MODELS)
    for i, (name, metric, sc, _, _) in enumerate(m.MODELS, start=2):
        ws.append([name, metric] + list(sc) + [f"=AVERAGE(C{i}:G{i})", f"=RANK(H{i},$H$2:$H${nm + 1})"])
        for c in range(3, 8):
            ws.cell(i, c).fill = k.INPUT_FILL
        ws.cell(i, 8).number_format = "0.0"
    k.style_table(ws, 1, 1, nm + 1, 9)
    col_widths(ws, [34, 30, 18, 16, 12, 14, 16, 10, 7])
    k.page(ws)


TARIFF_ROWS = [
    ("Базовий платіж на клієнта, EUR/міс.", lambda t, c: t.base, "input"),
    ("Ціна за активну будівлю, EUR/міс.", lambda t, c: t.per_building, "input"),
    ("Максимум будівель у тарифі", lambda t, c: t.max_buildings, "input"),
    ("Частка клієнтів (A1)", lambda t, c: t.mix, "input"),
    ("Середнє usage, будівель на клієнта (A2)", lambda t, c: t.usage, "input"),
    ("Год підтримки на клієнта (A4)", lambda t, c: t.client_hours, "input"),
    ("COGS на клієнта (база), EUR", lambda t, c: f"={c}7*x4_sup_rate", "calc"),
    ("Очікуваний рахунок клієнта, EUR/міс.", lambda t, c: f"={c}2+{c}3*{c}6", "calc"),
    ("Очікуваний COGS клієнта, EUR/міс.", lambda t, c: f"={c}8+{c}6*x4_cogs", "calc"),
    ("Gross profit, EUR/міс.", lambda t, c: f"={c}9-{c}10", "calc"),
    ("Gross margin", lambda t, c: f"=({c}9-{c}10)/{c}9", "calc"),
    ("Floor price пакета, EUR/міс.", lambda t, c: f"={c}10/(1-x4_gm)", "calc"),
    ("Ціна не нижча за floor", lambda t, c: f'=IF({c}9>={c}13,"так","ні")', "calc"),
    ("Markup на COGS", lambda t, c: f"=({c}9-{c}10)/{c}10", "calc"),
    ("Ціна на будівлю, EUR", lambda t, c: f"={c}9/{c}6", "calc"),
    ("Частка цінності 900 EUR", lambda t, c: f"={c}16/x4_value", "calc"),
    ("Ціна будівлі не нижча за floor одиниці", lambda t, c: f'=IF({c}3>=x4_floor_unit,"так","ні")', "calc"),
    ("Частка COGS бази в COGS клієнта", lambda t, c: f"={c}8/{c}10", "calc"),
]

PCT_ROWS = {5, 12, 15, 17, 19}


def sheet_tariffs(wb):
    ws = wb.create_sheet("Тарифи")
    ws.append(["Показник", "Формула або підстава"] + [t.name for t in m.TIERS])
    notes = ["вхід", "вхід", "захисне правило", "A1", "A2", "A4", "год × 30", "база + ціна × usage", "COGS бази + usage × 16",
             "рахунок мінус COGS", "(Revenue − COGS) / Revenue", "COGS / (1 − 0,70)", "порівняння", "(ціна − COGS) / COGS",
             "рахунок / usage", "ціна будівлі / цінність", "ціна будівлі ≥ 53,33", "COGS бази / COGS клієнта"]
    for i, (label, f, kind) in enumerate(TARIFF_ROWS, start=2):
        ws.cell(i, 1, label)
        ws.cell(i, 2, notes[i - 2])
        for t in m.TIERS:
            c = TCOL[t.tid]
            ws[f"{c}{i}"] = f(t, c)
            if kind == "input":
                ws[f"{c}{i}"].fill = k.INPUT_FILL
            ws[f"{c}{i}"].number_format = "0.0%" if i in PCT_ROWS else ("0.00" if i in (7,) else "#,##0.00")
    for t in m.TIERS:
        c = TCOL[t.tid]
        ws[f"{c}4"].number_format = "0"
        ws[f"{c}6"].number_format = "0.0"
        k.name(wb, f"x4_mix_{t.tid}", f"'Тарифи'!${c}$5")
        k.name(wb, f"x4_use_{t.tid}", f"'Тарифи'!${c}$6")
        k.name(wb, f"x4_base_{t.tid}", f"'Тарифи'!${c}$2")
        k.name(wb, f"x4_pb_{t.tid}", f"'Тарифи'!${c}$3")
        k.name(wb, f"x4_cc_{t.tid}", f"'Тарифи'!${c}$8")
    last = len(TARIFF_ROWS) + 1
    k.style_table(ws, 1, 1, last, 5)
    w0 = last + 2
    ws.cell(w0, 1, "Перевірка формули floor")
    ws.cell(w0, 2, "Значення, EUR")
    ws.cell(w0, 3, "GM")
    ws.cell(w0 + 1, 1, "Правильно: COGS / (1 − GM)")
    ws.cell(w0 + 1, 2, "=x4_cogs_var/(1-x4_gm)")
    ws.cell(w0 + 1, 3, f"=(B{w0 + 1}-x4_cogs_var)/B{w0 + 1}")
    ws.cell(w0 + 2, 1, "Помилково: COGS × (1 + GM)")
    ws.cell(w0 + 2, 2, "=x4_cogs_var*(1+x4_gm)")
    ws.cell(w0 + 2, 3, f"=(B{w0 + 2}-x4_cogs_var)/B{w0 + 2}")
    for r in (w0 + 1, w0 + 2):
        ws.cell(r, 2).number_format = "0.00"
        ws.cell(r, 3).number_format = "0.0%"
    k.style_table(ws, w0, 1, w0 + 2, 3)
    col_widths(ws, [40, 25, 13, 13, 13])
    k.page(ws, landscape=False)
    LAYOUT["floor_row"] = w0


def sheet_high(wb):
    ws = wb.create_sheet("High usage")
    ws.append(["Показник", "Формула"] + [t.name for t in m.TIERS])
    rows = [
        ("Будівель у клієнта з high usage", "максимум тарифу", lambda t, c: t.high_buildings),
        ("Рахунок без захисного правила, EUR", "база + ціна × n", lambda t, c: f"='Тарифи'!{c}2+'Тарифи'!{c}3*{c}2"),
        ("COGS, EUR", "COGS бази + n × COGS high", lambda t, c: f"='Тарифи'!{c}8+{c}2*x4_cogs_high"),
        ("Gross margin без правила", "", lambda t, c: f"=({c}3-{c}4)/{c}3"),
        ("Overage за правилом, EUR", "n × (доплата + SMS)", lambda t, c: f"={c}2*x4_over"),
        ("Рахунок з правилом, EUR", "", lambda t, c: f"={c}3+{c}6"),
        ("Gross margin з правилом", "", lambda t, c: f"=({c}7-{c}4)/{c}7"),
        ("Межа 70% виконана без правила", "", lambda t, c: f'=IF({c}5>=x4_gm,"так","ні")'),
        ("Межа 70% виконана з правилом", "", lambda t, c: f'=IF({c}8>=x4_gm,"так","ні")'),
    ]
    for i, (label, note, f) in enumerate(rows, start=2):
        ws.cell(i, 1, label)
        ws.cell(i, 2, note)
        for t in m.TIERS:
            c = TCOL[t.tid]
            ws[f"{c}{i}"] = f(t, c)
            ws[f"{c}{i}"].number_format = "0.0%" if i in (5, 8) else "#,##0.00"
        if i == 2:
            for t in m.TIERS:
                ws[f"{TCOL[t.tid]}2"].fill = k.INPUT_FILL
                ws[f"{TCOL[t.tid]}2"].number_format = "0"
    k.style_table(ws, 1, 1, len(rows) + 1, 5)
    col_widths(ws, [34, 26, 12, 12, 12])
    k.page(ws, landscape=False)


def sheet_scenarios(wb):
    ws = wb.create_sheet("Сценарії")
    ws.append(["Показник за місяць, EUR", "Низький 70%", "Базовий 100%", "Високий 140%"])
    ws.append(["Коефіцієнт масштабу", 0.7, 1.0, 1.4])
    lines = [
        ("Платні будівлі (billable units)", lambda c: f"=x4_scale*{c}2"),
        ("Середнє usage, будівель на клієнта", lambda c: "=" + "+".join(f"x4_mix_{t.tid}*x4_use_{t.tid}" for t in m.TIERS)),
        ("Платні клієнти", lambda c: f"={c}3/{c}4"),
        ("Клієнти Basic", lambda c: f"={c}5*x4_mix_B"),
        ("Клієнти Standard", lambda c: f"={c}5*x4_mix_S"),
        ("Клієнти Pro", lambda c: f"={c}5*x4_mix_P"),
        ("Revenue: базові платежі", lambda c: f"={c}6*x4_base_B+{c}7*x4_base_S+{c}8*x4_base_P"),
        ("Revenue: будівлі (usage)", lambda c: f"={c}6*x4_use_B*x4_pb_B+{c}7*x4_use_S*x4_pb_S+{c}8*x4_use_P*x4_pb_P"),
        ("Будівлі з high usage", lambda c: f"={c}3*x4_high_share"),
        ("Revenue: overage", lambda c: f"={c}11*x4_over"),
        ("Revenue", lambda c: f"={c}9+{c}10+{c}12"),
        ("COGS: база клієнтів", lambda c: f"={c}6*x4_cc_B+{c}7*x4_cc_S+{c}8*x4_cc_P"),
        ("COGS: будівлі", lambda c: f"=({c}3-{c}11)*x4_cogs+{c}11*x4_cogs_high"),
        ("COGS", lambda c: f"={c}14+{c}15"),
        ("Gross profit", lambda c: f"={c}13-{c}16"),
        ("Gross margin", lambda c: f"={c}17/{c}13"),
        ("ARPA, EUR на клієнта", lambda c: f"={c}13/{c}5"),
        ("Revenue на будівлю", lambda c: f"={c}13/{c}3"),
    ]
    for i, (label, f) in enumerate(lines, start=3):
        ws.cell(i, 1, label)
        for c in "BCD":
            ws[f"{c}{i}"] = f(c)
            ws[f"{c}{i}"].number_format = "0.0%" if i == 18 else ("#,##0.0" if i in (3, 4, 5, 6, 7, 8, 11) else "#,##0.00")
    for c in "BCD":
        ws[f"{c}2"].fill = k.INPUT_FILL
        ws[f"{c}2"].number_format = "0%"
    n = len(lines) + 2
    k.style_table(ws, 1, 1, n, 4)
    ch = BarChart()
    ch.type = "col"
    ch.grouping = "clustered"
    ch.title = "Revenue, COGS і gross profit за місяць, EUR"
    ch.y_axis.title = "EUR за місяць"
    ch.y_axis.number_format = "#,##0"
    for r, color in ((13, "2a78d6"), (16, "eb6834"), (17, "1baf7a")):
        ref = Reference(ws, min_col=1, max_col=4, min_row=r, max_row=r)
        ch.add_data(ref, titles_from_data=True, from_rows=True)
        ch.series[-1].graphicalProperties.solidFill = color
        ch.series[-1].graphicalProperties.line.solidFill = color
    ch.set_categories(Reference(ws, min_col=2, max_col=4, min_row=1, max_row=1))
    ch.height = 8.5
    ch.width = 17
    ch.legend.position = "b"
    ws.add_chart(ch, f"A{n + 2}")
    col_widths(ws, [38, 15, 15, 15])
    k.page(ws, landscape=False)
    LAYOUT["scen_rows"] = n


def build_workbook(path):
    wb = Workbook()
    sheet_params(wb)
    sheet_assumptions(wb)
    sheet_models(wb)
    sheet_tariffs(wb)
    sheet_high(wb)
    sheet_scenarios(wb)
    wb.save(path)


def close(a, b, tol=1e-6):
    return a is not None and not isinstance(a, str) and abs(a - b) < tol


def verify(path):
    tmp = Path(tempfile.mkdtemp())
    k.soffice(path, "xlsx", tmp)
    wb = load_workbook(tmp / path.name, data_only=True)
    ts, hs, ss, ms, ps = wb["Тарифи"], wb["High usage"], wb["Сценарії"], wb["Моделі"], wb["Параметри ЛР4"]
    e = LAYOUT["params_extra"]
    checks = {
        "cogs": (ps[f"C{e}"].value, m.cogs_building()),
        "cogs_high": (ps[f"C{e + 1}"].value, m.cogs_building(True)),
        "overage": (ps[f"C{e + 2}"].value, m.overage_high()),
        "floor_unit": (ps[f"C{e + 3}"].value, m.floor_price(m.P4["cogs_unit"])),
        "floor_wrong_gm": (ts[f"C{LAYOUT['floor_row'] + 2}"].value, (m.P4["cogs_unit"] * 1.7 - 16) / (m.P4["cogs_unit"] * 1.7)),
    }
    for t in m.TIERS:
        c = TCOL[t.tid]
        ex = m.tier_expected(t)
        hu = m.high_usage_check(t)
        checks[f"rev_{t.tid}"] = (ts[f"{c}9"].value, ex["revenue"])
        checks[f"cogs_{t.tid}"] = (ts[f"{c}10"].value, ex["cogs"])
        checks[f"gm_{t.tid}"] = (ts[f"{c}12"].value, ex["gm"])
        checks[f"floor_{t.tid}"] = (ts[f"{c}13"].value, ex["floor"])
        checks[f"markup_{t.tid}"] = (ts[f"{c}15"].value, ex["markup"])
        checks[f"hu_norule_{t.tid}"] = (hs[f"{c}5"].value, hu["norule"]["gm"])
        checks[f"hu_rule_{t.tid}"] = (hs[f"{c}8"].value, hu["rule"]["gm"])
    for (name, s), c in zip(m.scenarios(), "BCD"):
        for key, row in (("revenue", 13), ("cogs", 16), ("gp", 17), ("gm", 18), ("clients", 5), ("arpa", 19)):
            checks[f"{name}_{key}"] = (ss[f"{c}{row}"].value, s[key])
    for i, (name, sc) in enumerate(m.model_scores(), start=2):
        checks[f"model_{i}"] = (ms[f"H{i}"].value, sc)
    texts = [ts[f"{c}14"].value for c in "CDE"] + [ts[f"{c}18"].value for c in "CDE"] + [ps[f"F{e}"].value]
    ok = all(v in ("так", "збігається з варіантом") for v in texts)
    print("text checks:", texts)
    for key, (a, b) in checks.items():
        good = close(a, b, 1e-6)
        ok = ok and good
        print(f"{key}: LibreOffice={a} Python={b} {'ok' if good else 'MISMATCH'}")
    print("checks:", len(checks) + len(texts))
    shutil.copy(tmp / path.name, OUT / "lab4_recalculated.xlsx")
    shutil.rmtree(tmp, ignore_errors=True)
    return ok


PRINT_VIEWS = {
    "Тарифи": ("tariffs", "portrait"),
    "High usage": ("high", "portrait"),
    "Сценарії": ("scen", "portrait"),
    "Моделі": ("models", "landscape"),
    "Параметри ЛР4": ("params", "portrait"),
}


def render_sheets(path):
    import pymupdf
    from PIL import Image
    files = {}
    for title, (key, orient) in PRINT_VIEWS.items():
        tmp = Path(tempfile.mkdtemp())
        wb = load_workbook(path)
        for ws in wb.worksheets:
            ws.sheet_state = "visible" if ws.title == title else "hidden"
        wb.active = wb.sheetnames.index(title)
        ws = wb[title]
        k.split_columns(ws)
        for r in range(1, ws.max_row + 1):
            ws.row_dimensions[r].height = 16
        ws.row_dimensions[1].height = 30 if key != "models" else 46
        if key == "models":
            for r in range(2, ws.max_row + 1):
                ws.row_dimensions[r].height = 32
        if key == "params":
            for r in range(2, ws.max_row + 1):
                if len(str(ws.cell(r, 6).value or "")) > 40 or len(str(ws.cell(r, 2).value or "")) > 36:
                    ws.row_dimensions[r].height = 28
        ws.sheet_properties.pageSetUpPr.fitToPage = False
        ws.page_setup.scale = 100
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.orientation = orient
        src = tmp / f"view_{key}.xlsx"
        wb.save(src)
        k.soffice(src, "pdf", tmp)
        doc = pymupdf.open(tmp / f"view_{key}.pdf")
        pages = []
        for i, pg in enumerate(doc):
            pix = pg.get_pixmap(dpi=170)
            part = tmp / f"p{i}.png"
            pix.save(part)
            k.crop(part)
            pages.append(Image.open(part).convert("RGB"))
        width = max(im.width for im in pages)
        canvas = Image.new("RGB", (width, sum(im.height for im in pages)), "white")
        y = 0
        for im in pages:
            canvas.paste(im, (0, y))
            y += im.height
        out = OUT / f"sheet_{key}.png"
        canvas.save(out)
        files[key] = out
        shutil.rmtree(tmp, ignore_errors=True)
    return files


def eur(v, d=0):
    return f"{v:,.{d}f}".replace(",", " ").replace(".", ",")


def chart_scenarios():
    sc = m.scenarios()
    labels = [f"{n}\n{int(s['k'] * 100)}%: {eur(s['buildings'], 1)} буд., {eur(s['clients'])} клієнтів" for n, s in sc]
    x = np.arange(3)
    w = 0.26
    fig, ax = plt.subplots(figsize=(10, 5.2), dpi=200)
    for i, (key, lbl) in enumerate((("revenue", "Revenue"), ("cogs", "COGS"), ("gp", "Gross profit"))):
        vals = [s[key] for _, s in sc]
        ax.bar(x + (i - 1) * w, vals, width=w, color=COLORS[i], label=lbl)
        for j, v in enumerate(vals):
            ax.text(x[j] + (i - 1) * w, v + 250, eur(v), ha="center", fontsize=9.5, color=INK)
    for j, (_, s) in enumerate(sc):
        ax.text(x[j], s["revenue"] + 1500, f"GM {eur(s['gm'] * 100, 1)}%", ha="center", fontsize=11, color=INK, fontweight="bold")
    ax.set_xticks(x)
    ax.set_xticklabels(labels, fontsize=10.5)
    ax.set_ylim(0, 19500)
    ax.set_ylabel("EUR за місяць")
    ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: eur(v)))
    ax.set_title("Сценарії попиту третього року: Revenue і COGS за місяць")
    ax.legend(loc="upper left", frameon=False)
    k.axes(ax)
    ax.grid(axis="x", visible=False)
    fig.tight_layout()
    fig.savefig(OUT / "fig_scenarios.png")
    plt.close(fig)


def chart_bills():
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(10.5, 4.9), dpi=200, gridspec_kw={"width_ratios": [1.15, 1]})
    for i, t in enumerate(m.TIERS):
        pts = [(n, b) for n, b in m.bill_curve(t, 10)]
        a1.plot([p[0] for p in pts], [p[1] for p in pts], marker="o", color=COLORS[i], linewidth=2, label=t.name)
    floor = [m.floor_price(m.client_cogs(m.TIER["P"], n)) for n in range(1, 11)]
    a1.plot(range(1, 11), floor, color=MUTED, linestyle=":", linewidth=1.6, label="floor пакета Pro")
    a1.axvline(5, color=GRID, linewidth=1.2)
    a1.annotate("Standard = Pro\nпри 5 будівлях", xy=(5, 600), xytext=(5.9, 170), fontsize=9.5, color=INK, arrowprops={"arrowstyle": "->", "color": MUTED})
    a1.set_xlabel("Активних будівель у клієнта")
    a1.set_ylabel("Рахунок клієнта, EUR за місяць")
    a1.set_title("Місячний рахунок за тарифами", fontsize=12.5)
    a1.set_xticks(range(1, 11))
    a1.legend(loc="upper left", frameon=False, fontsize=10)
    a1.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: eur(v)))
    k.axes(a1)
    names = [t.name for t in m.TIERS]
    vals = [[m.tier_expected(t)["gm"] for t in m.TIERS],
            [m.high_usage_check(t)["norule"]["gm"] for t in m.TIERS],
            [m.high_usage_check(t)["rule"]["gm"] for t in m.TIERS]]
    lbls = ["очікуване usage", "high usage без правила", "high usage з правилом"]
    x = np.arange(3)
    w = 0.26
    for i, (v, l) in enumerate(zip(vals, lbls)):
        a2.bar(x + (i - 1) * w, [q * 100 for q in v], width=w, color=COLORS[i], label=l)
        for j, q in enumerate(v):
            a2.text(x[j] + (i - 1) * w, q * 100 + 1, eur(q * 100, 1), ha="center", fontsize=8.5, color=INK)
    a2.axhline(70, color=COLORS[3], linewidth=1.6)
    a2.text(2.45, 71, "70%", fontsize=10, color=INK)
    a2.set_ylim(50, 95)
    a2.set_xticks(x)
    a2.set_xticklabels(names)
    a2.set_ylabel("Gross margin, %")
    a2.set_title("Gross margin тарифів і межа 70%", fontsize=12.5)
    a2.legend(loc="upper center", bbox_to_anchor=(0.5, -0.1), ncol=1, frameon=False, fontsize=9.5)
    k.axes(a2)
    a2.grid(axis="x", visible=False)
    fig.tight_layout()
    fig.savefig(OUT / "fig_bills.png")
    plt.close(fig)


def chart_sens():
    base, rows = m.sensitivity4()
    rows = sorted(rows, key=lambda r: abs(r["rev_hi"] - r["rev_lo"]))
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(10.5, 4.3), dpi=200)
    r0, g0 = base["revenue"], base["gm"] * 100
    for i, r in enumerate(rows):
        a1.barh(i, r["rev_lo"] - r0, left=r0, color=COLORS[0], height=0.6)
        a1.barh(i, r["rev_hi"] - r0, left=r0, color=COLORS[1], height=0.6)
        a2.barh(i, r["gm_lo"] * 100 - g0, left=g0, color=COLORS[0], height=0.6)
        a2.barh(i, r["gm_hi"] * 100 - g0, left=g0, color=COLORS[1], height=0.6)
    for a in (a1, a2):
        a.set_yticks(range(len(rows)))
        a.set_yticklabels([r["label"] for r in rows], fontsize=10)
        k.axes(a)
        a.grid(axis="y", visible=False)
    a2.set_yticklabels([])
    a1.axvline(r0, color=INK, linewidth=0.8)
    a2.axvline(g0, color=INK, linewidth=0.8)
    a1.set_xlabel("Revenue, EUR за місяць")
    a1.xaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: eur(v)))
    a1.set_title(f"Revenue: база {eur(r0)} EUR", fontsize=12)
    a2.set_xlabel("Gross margin, %")
    a2.set_title(f"Gross margin: база {eur(g0, 1)}%", fontsize=12)
    from matplotlib.patches import Patch
    fig.legend(handles=[Patch(color=COLORS[0], label="параметр мінус 20%"), Patch(color=COLORS[1], label="параметр плюс 20%")],
               loc="lower center", ncol=2, frameon=False, fontsize=10.5)
    fig.tight_layout(rect=(0, 0.08, 1, 1))
    fig.savefig(OUT / "fig_sens.png")
    plt.close(fig)


def tariff_html():
    cards = []
    for i, t in enumerate(m.TIERS):
        ex = m.tier_expected(t)
        limit = f"до {t.max_buildings} будівель" if t.tid != "P" else "до 20 будівель, більше: індивідуальний договір"
        funcs = "моніторинг, аномалії, сповіщення, інциденти, дашборд, звірка з рахунком" if t.functions == "F1...F10" else "усе з Basic і звіти орендарям з ролями доступу"
        items = [f"{t.functions}: {funcs}", f"користувачів {t.users}", limit, "до 30 точок обліку і 20 SMS на будівлю",
                 f"відповідь підтримки: {t.response}"]
        if t.tid == "P":
            items.append("квартальний огляд енергоспоживання")
        li = "".join(f"<li>{x}</li>" for x in items)
        cards.append(f"""<div class="card" style="border-top:6px solid {k.COLORS[i]}"><h2>{t.name}</h2><p class="who">{t.target}</p>
<p class="price"><b>{eur(t.base)} EUR</b> на клієнта<br><b>+ {eur(t.per_building)} EUR</b> за активну будівлю</p>
<p class="ex">приклад: {eur(t.usage)} буд. = {eur(ex['revenue'])} EUR на місяць</p><ul>{li}</ul></div>""")
    rules = ("SMS понад 20 на будівлю: 0,25 EUR за SMS. Будівля зі статусом «нестабільний облік» (понад 10% пропущених "
             "інтервалів 2 місяці поспіль, після 1 пільгового місяця): доплата 50 EUR на місяць. Будівля є платною, якщо в місяці "
             "отримано хоча б один валідний 15-хвилинний інтервал. Ціни в EUR за календарний місяць, без ПДВ. Підключення шлюзу "
             "оплачується окремо за собівартістю.")
    html = f"""<!doctype html><html><head><meta charset="utf-8"><style>
body{{font-family:Arial,Helvetica,sans-serif;color:#0b0b0b;margin:0;background:#fff;width:900px}}
.wrap{{padding:20px 22px}}h1{{font-size:30px;margin:0 0 4px}}.sub{{color:#52514e;font-size:18px;margin:0 0 18px}}
.grid{{display:flex;gap:12px}}.card{{flex:1;border:1px solid #e1e0d9;border-radius:8px;padding:12px 14px;background:#fbfbf9}}
h2{{margin:0 0 6px;font-size:26px}}.who{{color:#52514e;font-size:16px;min-height:84px;margin:0}}
.price{{font-size:19px;margin:10px 0 6px;line-height:1.4}}.ex{{font-size:16px;color:#52514e;margin:0 0 8px}}
ul{{padding-left:18px;margin:0;font-size:17px;line-height:1.4}}.rules{{margin-top:16px;font-size:16.5px;color:#0b0b0b;border-top:1px solid #e1e0d9;padding-top:12px;line-height:1.5}}
</style></head><body><div class="wrap"><h1>ЕнергоПульс: тарифна сітка</h1>
<p class="sub">Hybrid base + usage: базовий платіж клієнта і ціна за активну будівлю на місяць</p>
<div class="grid">{''.join(cards)}</div><p class="rules">{rules}</p></div></body></html>"""
    path = OUT / "tariff_grid.html"
    path.write_text(html, encoding="utf-8")
    png = OUT / "tariff_grid.png"
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--hide-scrollbars", f"--screenshot={png}",
                    "--window-size=900,1300", "--force-device-scale-factor=2", path.as_uri()], check=True, capture_output=True, timeout=120)
    k.crop(png)
    return png


def main():
    OUT.mkdir(exist_ok=True)
    xlsx = OUT / "lab4_pricing.xlsx"
    build_workbook(xlsx)
    ok = verify(xlsx)
    print("workbook verified:", ok)
    render_sheets(OUT / "lab4_recalculated.xlsx")
    chart_scenarios()
    chart_bills()
    chart_sens()
    tariff_html()
    if not ok:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
