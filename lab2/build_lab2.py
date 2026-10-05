import math
import shutil
import subprocess
import sys
import tempfile
import textwrap
from datetime import date, datetime
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from matplotlib.ticker import FuncFormatter
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT.parent))
sys.path.insert(0, str(ROOT.parent / "lab1"))
import econ_model as m
import build_lab1 as k

OUT = ROOT / "out"
plt.rcParams.update({"font.size": 12.5})
COLORS = k.COLORS
INK, MUTED, GRID = k.INK, k.MUTED, k.GRID
ROLE_COLS = {"BA": "E", "D1": "F", "D2": "G", "QA": "H", "OPS": "I"}
FIRST = 2
ROW = {w.wid: FIRST + i for i, w in enumerate(m.WBS)}
LAST = FIRST + len(m.WBS) - 1
END_ROW = LAST + 2
VEL_ROWS = [0, 0]


UA_MONTHS = ["січ", "лют", "бер", "кві", "тра", "чер", "лип", "сер", "вер", "жов", "лис", "гру"]


def iso(s):
    return datetime.strptime(str(s), "%Y-%m-%d")


def col_widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def sheet_params(wb):
    ws = wb.active
    ws.title = "Параметри ЛР2"
    ws.append(["Ім'я", "Показник", "Значення", "Одиниця", "Тип", "Пояснення"])
    rows = [
        ("p2_start", "Дата старту проєкту", iso(m.PROJECT_START), "дата", "П", "після тестів A1, A2, A5 з ЛР № 1"),
        ("p2_hours", "Годин у робочому дні", m.HOURS_DAY, "год", "П", "повний робочий день"),
        ("p2_focus", "Частка продуктивного часу", m.FOCUS, "частка", "П", "30% наради, церемонії, комунікація; 3,75 FTE × 40 × 0,7 = 105"),
        ("p2_days_week", "Робочих днів на тиждень", m.DAYS_WEEK, "днів", "П", ""),
        ("p2_z80", "Квантиль N(0;1) рівня 0,8", "=NORMSINV(0.8)", "", "Д", "методичка: 0,842"),
        ("p2_backlog", "Backlog", m.BACKLOG_SP, "SP", "Д", "варіант 14"),
        ("p2_sprint", "Тривалість спринту", m.SPRINT_WEEKS, "тиж.", "Д", "усі спринти 2 тижні"),
        ("p2_tcf", "TCF", m.UCP_TCF, "", "Д", "варіант 14"),
        ("p2_ecf", "ECF", m.UCP_ECF, "", "Д", "варіант 14"),
        ("p2_pf", "PF", m.UCP_PF, "люд.-год/UCP", "Д", "варіант 14"),
        ("p2_util", "Досяжне завантаження вузької ролі", m.BOTTLENECK_UTIL, "частка", "П", "залежності і очікування; WBS дає 0,80 для розробників"),
        ("p2_mr", "Management reserve", m.MR_DAYS, "роб. днів", "П", "системний сценарій нестабільних лічильників"),
        ("p2_mc_p50", "P50 за симуляцією Монте-Карло", round(m.baseline()["p50_mc"], 2), "роб. днів", "Р", "20 000 прогонів, seed 14, Python"),
        ("p2_mc_p80", "P80 за симуляцією Монте-Карло", round(m.baseline()["p80_mc"], 2), "роб. днів", "Р", "20 000 прогонів, seed 14, Python"),
    ]
    for i, r in enumerate(rows, start=2):
        ws.append(list(r))
        if r[4] in ("П", "Д"):
            ws.cell(i, 3).fill = k.INPUT_FILL
        k.name(wb, r[0], f"'Параметри ЛР2'!$C${i}")
    ws.cell(2, 3).number_format = "DD.MM.YYYY"
    ws.cell(6, 3).number_format = "0.000"
    k.style_table(ws, 1, 1, len(rows) + 1, 6)
    top = len(rows) + 3
    ws.cell(top, 1, "Неробочі дні")
    for j, h in enumerate(m.HOLIDAYS):
        c = ws.cell(top + 1 + j, 1, iso(h))
        c.number_format = "DD.MM.YYYY"
        c.fill = k.INPUT_FILL
    k.name(wb, "p2_holidays", f"'Параметри ЛР2'!$A${top + 1}:$A${top + len(m.HOLIDAYS)}")
    k.style_table(ws, top, 1, top + len(m.HOLIDAYS), 1)
    ws.cell(top + 1, 2, "команда бере 2 дні відпустки; державні свята в умовах воєнного стану не є вихідними")
    rt = top + len(m.HOLIDAYS) + 2
    ws.append([])
    hdr = ["Роль", "Обов'язки", "FTE", "Год/день", "Год/тиждень"]
    for j, h in enumerate(hdr, start=1):
        ws.cell(rt, j, h)
    for i, r in enumerate(m.ROLES, start=rt + 1):
        ws.cell(i, 1, r.name)
        ws.cell(i, 2, r.duties)
        ws.cell(i, 3, r.fte).fill = k.INPUT_FILL
        ws.cell(i, 4, f"=C{i}*p2_hours*p2_focus")
        ws.cell(i, 5, f"=D{i}*p2_days_week")
        k.name(wb, f"cap_{r.rid}", f"'Параметри ЛР2'!$E${i}")
        k.name(wb, f"capd_{r.rid}", f"'Параметри ЛР2'!$D${i}")
    tot = rt + len(m.ROLES) + 1
    ws.cell(tot, 1, "Разом")
    ws.cell(tot, 3, f"=SUM(C{rt + 1}:C{tot - 1})")
    ws.cell(tot, 4, f"=SUM(D{rt + 1}:D{tot - 1})")
    ws.cell(tot, 5, f"=SUM(E{rt + 1}:E{tot - 1})")
    k.name(wb, "team_week", f"'Параметри ЛР2'!$E${tot}")
    k.name(wb, "team_day", f"'Параметри ЛР2'!$D${tot}")
    for r in range(rt + 1, tot + 1):
        ws.cell(r, 4).number_format = "0.00"
        ws.cell(r, 5).number_format = "0.0"
    k.style_table(ws, rt, 1, tot, 5)
    col_widths(ws, [20, 40, 12, 14, 12, 50])
    k.page(ws)
    return tot


def preds_text(w):
    parts = []
    for pid, kind, lag in w.preds:
        parts.append(pid + ("" if kind == "FS" and lag == 0 else f" {kind}+{lag}" if lag else f" {kind}"))
    return ", ".join(parts) if parts else "немає"


def sheet_wbs(wb):
    ws = wb.create_sheet("WBS і CPM")
    hdr = ["ID", "Пакет / результат", "Попередник", "Лідер", "BA/PM, год", "Розр. 1, год", "Розр. 2, год", "QA, год", "DevOps, год",
           "Effort M, год", "O, днів", "M, днів", "P, днів", "E, днів", "σ, днів", "Var, дні²", "Effort E, год",
           "ES", "EF", "LS", "LF", "Резерв", "Критичний", "Старт", "Фініш", "DoD"]
    ws.append(hdr)
    succ = m.successors()
    for w in m.WBS:
        r = ROW[w.wid]
        role = m.ROLE[w.lead].name
        vals = [w.wid, w.name, preds_text(w), role] + [w.hours.get(rid, 0) for rid in ROLE_COLS]
        vals += [f"=SUM(E{r}:I{r})", w.o, w.m, w.p, f"=(K{r}+4*L{r}+M{r})/6", f"=(M{r}-K{r})/6", f"=O{r}^2", f"=J{r}*N{r}/L{r}"]
        terms = []
        for pid, kind, lag in w.preds:
            pr = ROW[pid]
            terms.append(f"S{pr}+{lag}" if kind == "FS" else f"R{pr}+{lag}")
        vals.append("=MAX(0," + ",".join(terms) + ")" if terms else 0)
        vals.append(f"=R{r}+N{r}")
        vals.append(f"=U{r}-N{r}")
        lt = []
        for sid, kind, lag in succ[w.wid]:
            sr = ROW[sid]
            lt.append(f"T{sr}-{lag}" if kind == "FS" else f"T{sr}-{lag}+N{r}")
        vals.append("=MIN(proj_end," + ",".join(lt) + ")" if lt else "=proj_end")
        vals.append(f"=T{r}-R{r}")
        vals.append(f'=IF(ABS(V{r})<0.001,"так","ні")')
        vals.append(f"=WORKDAY(p2_start,ROUNDDOWN(R{r},0),p2_holidays)")
        vals.append(f"=WORKDAY(p2_start,ROUNDUP(S{r},0)-1,p2_holidays)")
        vals.append(w.dod)
        ws.append(vals)
        for c in range(5, 10):
            ws.cell(r, c).fill = k.INPUT_FILL
        for c in (11, 12, 13):
            ws.cell(r, c).fill = k.INPUT_FILL
    ws.cell(END_ROW, 1, "Разом")
    ws.cell(END_ROW, 2, "тривалість проєкту за E, роб. днів; effort, год")
    for c in "EFGHIJQ":
        ws[f"{c}{END_ROW}"] = f"=SUM({c}{FIRST}:{c}{LAST})"
    ws[f"S{END_ROW}"] = f"=MAX(S{FIRST}:S{LAST})"
    k.name(wb, "proj_end", f"'WBS і CPM'!$S${END_ROW}")
    ws[f"V{END_ROW}"] = f'=COUNTIF(W{FIRST}:W{LAST},"так")'
    ws[f"W{END_ROW}"] = "критичних"
    for r in range(FIRST, END_ROW + 1):
        for c in range(14, 23):
            ws.cell(r, c).number_format = "0.00"
        ws.cell(r, 17).number_format = "0"
        ws.cell(r, 24).number_format = "DD.MM.YY"
        ws.cell(r, 25).number_format = "DD.MM.YY"
    ws[f"V{END_ROW}"].number_format = "0"
    k.style_table(ws, 1, 1, LAST, 26)
    k.style_table(ws, END_ROW, 1, END_ROW, 26, header=False)
    for r in range(FIRST, LAST + 1):
        if m.cpm()[0][m.WBS[r - FIRST].wid]["critical"]:
            for c in (1, 23):
                ws.cell(r, c).font = Font(name="Arial", size=10, bold=True, color="C2410C")
    col_widths(ws, [7, 36, 11, 12, 7, 7, 7, 7, 7, 8, 6, 6, 6, 7, 6, 7, 8, 7, 7, 7, 7, 7, 11, 9, 9, 60])
    k.page(ws)


def sheet_pert(wb):
    ws = wb.create_sheet("PERT і резерв")
    ws.append(["ID", "Пакет", "O", "M", "P", "E", "σ", "Var", "Критичний", "Пояснення P"])
    for i, wid in enumerate(m.KEY_PERT, start=2):
        w = m.WBS_BY_ID[wid]
        r = ROW[wid]
        ws.append([wid, w.name, f"='WBS і CPM'!K{r}", f"='WBS і CPM'!L{r}", f"='WBS і CPM'!M{r}",
                   f"=(C{i}+4*D{i}+E{i})/6", f"=(E{i}-C{i})/6", f"=G{i}^2", f"='WBS і CPM'!W{r}", w.p_note])
        for c in (6, 7, 8):
            ws.cell(i, c).number_format = "0.00"
    n = len(m.KEY_PERT) + 1
    k.style_table(ws, 1, 1, n, 10)
    b = n + 2
    rng = f"'WBS і CPM'!$W${FIRST}:$W${LAST}"
    rows = [
        ("μ = ΣE критичного шляху", f"=SUMIF({rng},\"так\",'WBS і CPM'!$N${FIRST}:$N${LAST})", "роб. днів", "P50 ≈ μ"),
        ("ΣVar критичного шляху", f"=SUMIF({rng},\"так\",'WBS і CPM'!$P${FIRST}:$P${LAST})", "дні²", ""),
        ("σ шляху", f"=SQRT(B{b + 2})", "роб. днів", "√ΣVar"),
        ("P50", f"=B{b + 1}", "роб. днів", "аналітично"),
        ("P80", f"=B{b + 1}+p2_z80*B{b + 3}", "роб. днів", "μ + 0,842σ"),
        ("Contingency за формулою", f"=B{b + 5}-B{b + 4}", "роб. днів", "P80 − P50"),
        ("Контроль: кінець мережі за CPM", "=proj_end", "роб. днів", "має дорівнювати μ"),
        ("P80 за Монте-Карло", "=p2_mc_p80", "роб. днів", "враховує злиття шляхів"),
        ("Baseline", f"=ROUNDUP(MAX(B{b + 5},B{b + 8}),0)", "роб. днів", "більше з двох P80, ціле"),
        ("Contingency у baseline", f"=B{b + 9}-B{b + 4}", "роб. днів", "baseline − P50"),
        ("Management reserve", "=p2_mr", "роб. днів", "поза baseline"),
        ("Дата старту", "=p2_start", "дата", ""),
        ("Дата P50", f"=WORKDAY(p2_start,ROUNDUP(B{b + 4},0)-1,p2_holidays)", "дата", ""),
        ("Baseline-дата", f"=WORKDAY(p2_start,B{b + 9}-1,p2_holidays)", "дата", "з contingency"),
        ("Дата з management reserve", f"=WORKDAY(p2_start,B{b + 9}+B{b + 11}-1,p2_holidays)", "дата", "межа рішення спонсора"),
        ("Baseline, тижнів", f"=B{b + 9}/p2_days_week", "тиж.", ""),
    ]
    ws.cell(b, 1, "Показник")
    ws.cell(b, 2, "Значення")
    ws.cell(b, 3, "Одиниця")
    ws.cell(b, 4, "Пояснення")
    for j, (a, f, u, e) in enumerate(rows, start=b + 1):
        ws.cell(j, 1, a)
        ws.cell(j, 2, f)
        ws.cell(j, 3, u)
        ws.cell(j, 4, e)
        ws.cell(j, 2).number_format = "DD.MM.YYYY" if u == "дата" else "0.00"
    k.name(wb, "bl_days", f"'PERT і резерв'!$B${b + 9}")
    k.name(wb, "bl_p50", f"'PERT і резерв'!$B${b + 4}")
    k.name(wb, "bl_p80", f"'PERT і резерв'!$B${b + 5}")
    k.style_table(ws, b, 1, b + len(rows), 4)
    col_widths(ws, [30, 40, 9, 22, 7, 7, 7, 7, 13, 60])
    k.page(ws)
    return b


def sheet_velocity(wb):
    ws = wb.create_sheet("Backlog і velocity")
    ws.append(["ID", "Backlog item", "SP", "Пакет WBS", None, "Чому така оцінка"])
    for b in m.BACKLOG:
        ws.append([b.bid, b.name, b.sp, b.wbs, None, b.reason])
        ws.cell(ws.max_row, 3).fill = k.INPUT_FILL
    n = len(m.BACKLOG) + 1
    ws.cell(n + 1, 1, "Разом")
    ws.cell(n + 1, 3, f"=SUM(C2:C{n})")
    ws.cell(n + 1, 2, '=IF(C' + str(n + 1) + '=p2_backlog,"Разом: збігається з варіантом","Разом: розбіжність")')
    k.style_table(ws, 1, 1, n + 1, 4)
    k.style_table(ws, 1, 6, n + 1, 6)
    t = n + 3
    ws.cell(t, 1, "Спринт")
    ws.cell(t, 2, "Velocity, SP")
    for j, v in enumerate(m.VELOCITY):
        ws.cell(t + 1 + j, 1, f"S{j + 1}")
        ws.cell(t + 1 + j, 2, v).fill = k.INPUT_FILL
    k.style_table(ws, t, 1, t + 6, 2)
    f = t + 8
    ws.cell(f, 1, "Сценарій")
    ws.cell(f, 2, "Velocity, SP")
    ws.cell(f, 3, "Backlog / v")
    ws.cell(f, 4, "Спринтів")
    ws.cell(f, 5, "Тижні")
    vr = f"B{t + 1}:B{t + 6}"
    for j, (lbl, fn) in enumerate([("Fast: max", "MAX"), ("Central: медіана", "MEDIAN"), ("Slow: min", "MIN")], start=f + 1):
        ws.cell(j, 1, lbl)
        ws.cell(j, 2, f"={fn}({vr})")
        ws.cell(j, 3, f"=p2_backlog/B{j}")
        ws.cell(j, 4, f"=ROUNDUP(p2_backlog/B{j},0)")
        ws.cell(j, 5, f"=D{j}*p2_sprint")
        ws.cell(j, 3).number_format = "0.00"
    k.style_table(ws, f, 1, f + 3, 5)
    s = f + 5
    sched, end = m.cpm()
    ws.cell(s, 1, "Повна тривалість")
    ws.cell(s, 2, "До спринтів (1.1, 1.2), роб. днів")
    ws.cell(s, 3, "Спринти, днів")
    ws.cell(s, 4, "Після FC, днів")
    ws.cell(s, 5, "Тижнів")
    fc_cells = ",".join(f"'WBS і CPM'!S{ROW[i]}" for i in m.MILESTONES[2][1])
    for j, lbl in enumerate(["Fast", "Central", "Slow"], start=s + 1):
        ws.cell(j, 1, lbl)
        ws.cell(j, 2, f"='WBS і CPM'!S{ROW['1.2']}")
        ws.cell(j, 3, f"=E{f + 1 + j - s - 1}*p2_days_week")
        ws.cell(j, 4, f"=proj_end-MAX({fc_cells})")
        ws.cell(j, 5, f"=(B{j}+C{j}+D{j})/p2_days_week")
        for c in (2, 4, 5):
            ws.cell(j, c).number_format = "0.00"
    k.style_table(ws, s, 1, s + 3, 5)
    col_widths(ws, [22, 46, 14, 14, 12, 50])
    k.page(ws)
    VEL_ROWS[:] = [f, s]
    return f, s


def sheet_ucp(wb):
    ws = wb.create_sheet("UCP")
    ws.append(["Актор", "Тип S/A/C", "Правило", "Вага"])
    for a in m.UCP_ACTORS:
        r = ws.max_row + 1
        ws.append([a.name, a.kind, a.rule, f'=IF(B{r}="S",1,IF(B{r}="A",2,3))'])
    na = len(m.UCP_ACTORS) + 1
    k.style_table(ws, 1, 1, na, 4)
    u0 = na + 2
    ws.cell(u0, 1, "Use case")
    ws.cell(u0, 2, "Транзакцій")
    ws.cell(u0, 3, "Функція")
    ws.cell(u0, 4, "Тип S/A/C")
    ws.cell(u0, 5, "Вага")
    for i, u in enumerate(m.USE_CASES, start=u0 + 1):
        ws.cell(i, 1, f"{u.uid} {u.name}")
        ws.cell(i, 2, u.tx).fill = k.INPUT_FILL
        ws.cell(i, 3, u.func)
        ws.cell(i, 4, f'=IF(B{i}<=3,"S",IF(B{i}<=7,"A","C"))')
        ws.cell(i, 5, f'=IF(D{i}="S",5,IF(D{i}="A",10,15))')
    ul = u0 + len(m.USE_CASES)
    k.style_table(ws, u0, 1, ul, 5)
    s = ul + 2
    hdr = ["Компонент", "Кількість", "Вага", "Результат"]
    for j, h in enumerate(hdr, start=1):
        ws.cell(s, j, h)
    rows = [
        ("Прості актори", f'=COUNTIF(B2:B{na},"S")', 1, f"=B{s + 1}*C{s + 1}"),
        ("Середні актори", f'=COUNTIF(B2:B{na},"A")', 2, f"=B{s + 2}*C{s + 2}"),
        ("Складні актори", f'=COUNTIF(B2:B{na},"C")', 3, f"=B{s + 3}*C{s + 3}"),
        ("UAW", "", "", f"=SUM(D{s + 1}:D{s + 3})"),
        ("Прості UC", f'=COUNTIF(D{u0 + 1}:D{ul},"S")', 5, f"=B{s + 5}*C{s + 5}"),
        ("Середні UC", f'=COUNTIF(D{u0 + 1}:D{ul},"A")', 10, f"=B{s + 6}*C{s + 6}"),
        ("Складні UC", f'=COUNTIF(D{u0 + 1}:D{ul},"C")', 15, f"=B{s + 7}*C{s + 7}"),
        ("UUCW", "", "", f"=SUM(D{s + 5}:D{s + 7})"),
        ("UUCP", "", "", f"=D{s + 4}+D{s + 8}"),
        ("TCF × ECF", "=p2_tcf", "=p2_ecf", f"=B{s + 10}*C{s + 10}"),
        ("UCP", "", "", f"=D{s + 9}*D{s + 10}"),
        ("PF, люд.-год/UCP", "=p2_pf", "", "=p2_pf"),
        ("Effort, люд.-год", "", "", f"=D{s + 11}*D{s + 12}"),
        ("Effort при PF −20%", "", "", f"=D{s + 11}*p2_pf*0.8"),
        ("Effort при PF +20%", "", "", f"=D{s + 11}*p2_pf*1.2"),
    ]
    for j, r in enumerate(rows, start=s + 1):
        for c, v in enumerate(r, start=1):
            ws.cell(j, c, v)
        ws.cell(j, 4).number_format = "0.00"
    k.name(wb, "ucp_val", f"UCP!$D${s + 11}")
    k.name(wb, "ucp_effort", f"UCP!$D${s + 13}")
    k.style_table(ws, s, 1, s + len(rows), 4)
    c0 = s + len(rows) + 2
    hdr = ["Роль", "Частка в WBS", "Effort, год", "Тижнів, PF 22", "Тижнів, PF −20%", "Тижнів, PF +20%"]
    for j, h in enumerate(hdr, start=1):
        ws.cell(c0, j, h)
    for i, r in enumerate(m.ROLES, start=c0 + 1):
        col = ROLE_COLS[r.rid]
        ws.cell(i, 1, r.name)
        ws.cell(i, 2, f"='WBS і CPM'!{col}{END_ROW}/'WBS і CPM'!J{END_ROW}")
        ws.cell(i, 3, f"=ucp_effort*B{i}")
        ws.cell(i, 4, f"=C{i}/cap_{r.rid}")
        ws.cell(i, 5, f"=D{i}*0.8")
        ws.cell(i, 6, f"=D{i}*1.2")
        ws.cell(i, 2).number_format = "0.0%"
        for c in (3, 4, 5, 6):
            ws.cell(i, c).number_format = "0.0"
    d = c0 + len(m.ROLES) + 1
    r1, r2 = c0 + 2, c0 + 3
    extra = [
        ("Пул розробників", f"=B{r1}+B{r2}", f"=C{r1}+C{r2}", f"=C{d}/(cap_D1+cap_D2)"),
        ("Механічно: effort / 105", "", "=ucp_effort", "=ucp_effort/team_week"),
        ("Вузьке місце, тижнів", "", "", f"=MAX(D{d},D{c0 + 1},D{c0 + 4},D{c0 + 5})"),
        ("Строк з урахуванням завантаження", "", "", f"=D{d + 2}/p2_util"),
    ]
    for j, r in enumerate(extra, start=d):
        for c, v in enumerate(r, start=1):
            ws.cell(j, c, v)
        ws.cell(j, 5, f"=D{j}*0.8")
        ws.cell(j, 6, f"=D{j}*1.2")
        ws.cell(j, 2).number_format = "0.0%"
        for c in (3, 4, 5, 6):
            ws.cell(j, c).number_format = "0.0"
    k.name(wb, "ucp_weeks", f"UCP!$D${d + 3}")
    k.style_table(ws, c0, 1, d + 3, 6)
    col_widths(ws, [52, 14, 30, 17, 17, 17])
    k.page(ws)
    return s, c0, d


def sheet_calendar(wb):
    ws = wb.create_sheet("Календар")
    ws.append(["Роль", "FTE", "Год/тиждень", "Год у WBS (E)", "Ємність за baseline, год", "Завантаження"])
    for i, r in enumerate(m.ROLES, start=2):
        col = ROLE_COLS[r.rid]
        ws.append([r.name, r.fte, f"=cap_{r.rid}",
                   f"=SUMPRODUCT('WBS і CPM'!{col}{FIRST}:{col}{LAST},'WBS і CPM'!N{FIRST}:N{LAST}/'WBS і CPM'!L{FIRST}:L{LAST})",
                   f"=C{i}*bl_days/p2_days_week", f"=D{i}/E{i}"])
        for c in (3, 4, 5):
            ws.cell(i, c).number_format = "0.0"
        ws.cell(i, 6).number_format = "0.0%"
    n = len(m.ROLES) + 2
    ws.append(["Команда", f"=SUM(B2:B{n - 1})", f"=SUM(C2:C{n - 1})", f"=SUM(D2:D{n - 1})", f"=SUM(E2:E{n - 1})", f"=D{n}/E{n}"])
    for c in (3, 4, 5):
        ws.cell(n, c).number_format = "0.0"
    ws.cell(n, 6).number_format = "0.0%"
    k.style_table(ws, 1, 1, n, 6)
    t = n + 2
    for j, h in enumerate(["Віха", "День (E)", "Дата"], start=1):
        ws.cell(t, j, h)
    for i, (name, ids) in enumerate(m.MILESTONES, start=t + 1):
        cells = ",".join(f"'WBS і CPM'!S{ROW[x]}" for x in ids)
        ws.cell(i, 1, name)
        ws.cell(i, 2, f"=MAX({cells})")
        ws.cell(i, 3, f"=WORKDAY(p2_start,ROUNDUP(B{i},0)-1,p2_holidays)")
        ws.cell(i, 2).number_format = "0.00"
        ws.cell(i, 3).number_format = "DD.MM.YYYY"
    e = t + len(m.MILESTONES) + 1
    ws.cell(e, 1, "Baseline (P80, з contingency)")
    ws.cell(e, 2, "=bl_days")
    ws.cell(e, 3, "=WORKDAY(p2_start,bl_days-1,p2_holidays)")
    ws.cell(e + 1, 1, "Межа з management reserve")
    ws.cell(e + 1, 2, "=bl_days+p2_mr")
    ws.cell(e + 1, 3, "=WORKDAY(p2_start,bl_days+p2_mr-1,p2_holidays)")
    for r in (e, e + 1):
        ws.cell(r, 3).number_format = "DD.MM.YYYY"
    k.style_table(ws, t, 1, e + 1, 3)
    col_widths(ws, [32, 12, 14, 16, 18, 14])
    k.page(ws, landscape=False)
    return t


def sheet_compare(wb, vel_rows, ucp_rows):
    f, s = vel_rows
    us, uc0, ud = ucp_rows
    ws = wb.create_sheet("Порівняння")
    ws.append(["Метод", "Розмір", "Effort, год", "Тижнів: низ", "Тижнів: центр", "Тижнів: верх"])
    ws.append(["PERT/WBS", f"=COUNTA('WBS і CPM'!A{FIRST}:A{LAST})&\" пакетів\"", f"='WBS і CPM'!Q{END_ROW}",
               f"=SUMIF('WBS і CPM'!W{FIRST}:W{LAST},\"так\",'WBS і CPM'!K{FIRST}:K{LAST})/p2_days_week",
               "=bl_p50/p2_days_week", "=bl_days/p2_days_week"])
    ws.append(["SP/velocity", f"='Backlog і velocity'!C{len(m.BACKLOG) + 2}&\" SP\"", "не переводиться",
               f"='Backlog і velocity'!E{s + 1}", f"='Backlog і velocity'!E{s + 2}", f"='Backlog і velocity'!E{s + 3}"])
    ws.append(["UCP/parametric", "=TEXT(ucp_val,\"0.00\")&\" UCP\"", "=ucp_effort",
               f"=UCP!E{ud + 3}", "=ucp_weeks", f"=UCP!F{ud + 3}"])
    for r in range(2, 5):
        for c in (3, 4, 5, 6):
            ws.cell(r, c).number_format = "0.0"
    k.style_table(ws, 1, 1, 4, 6)
    col_widths(ws, [18, 16, 16, 14, 14, 14])
    k.page(ws, landscape=False)


def build_workbook(path):
    wb = Workbook()
    sheet_params(wb)
    sheet_wbs(wb)
    sheet_pert(wb)
    vel = sheet_velocity(wb)
    ucp_rows = sheet_ucp(wb)
    sheet_calendar(wb)
    sheet_compare(wb, vel, ucp_rows)
    ws = wb.create_sheet("Журнал scope", 1)
    ws.append(["ID", "Елемент ЛР № 1", "Статус", "Зміна / причина", "Вплив"])
    for row in m.SCOPE_LOG:
        ws.append(list(row))
    k.style_table(ws, 1, 1, len(m.SCOPE_LOG) + 1, 5)
    col_widths(ws, [6, 26, 12, 60, 44])
    k.page(ws)
    wb.save(path)
    return vel, ucp_rows


def close(a, b, tol=1e-6):
    return a is not None and abs(a - b) < tol


def verify(path):
    tmp = Path(tempfile.mkdtemp())
    k.soffice(path, "xlsx", tmp)
    wb = load_workbook(tmp / path.name, data_only=True)
    sched, end = m.cpm()
    ws = wb["WBS і CPM"]
    ok = True
    for w in m.WBS:
        r = ROW[w.wid]
        s = sched[w.wid]
        for col, key in (("R", "es"), ("S", "ef"), ("T", "ls"), ("U", "lf"), ("V", "float")):
            if not close(ws[f"{col}{r}"].value, s[key]):
                ok = False
                print("mismatch", w.wid, key, ws[f"{col}{r}"].value, s[key])
        if (ws[f"W{r}"].value == "так") != s["critical"]:
            ok = False
            print("critical mismatch", w.wid)
    bl = m.baseline()
    pp = m.p50_p80()
    pr = wb["PERT і резерв"]
    b = len(m.KEY_PERT) + 3
    checks = {
        "proj_end": (ws[f"S{END_ROW}"].value, end),
        "mu": (pr[f"B{b + 1}"].value, pp["mu"]),
        "var": (pr[f"B{b + 2}"].value, pp["var"]),
        "p80": (pr[f"B{b + 5}"].value, pp["p80"]),
        "baseline": (pr[f"B{b + 9}"].value, bl["days"]),
        "effort_e": (ws[f"Q{END_ROW}"].value, sum(w.effort_e for w in m.WBS)),
    }
    u = m.ucp()
    uc = m.ucp_calendar()
    names = {dn: wb.defined_names[dn] for dn in ("ucp_effort", "ucp_weeks", "ucp_val")}
    for key, dn in names.items():
        sheet, cell = list(dn.destinations)[0]
        val = wb[sheet][cell.replace("$", "")].value
        checks[key] = (val, {"ucp_effort": u["effort"], "ucp_weeks": uc["duration_weeks"], "ucp_val": u["ucp"]}[key])
    vs = wb["Backlog і velocity"]
    vf = m.velocity_forecast()
    f = VEL_ROWS[0]
    for j, kk in enumerate(("fast", "central", "slow"), start=f + 1):
        checks[f"sprints_{kk}"] = (vs[f"D{j}"].value, vf[kk]["sprints"])
    sp = m.sp_phase_totals()
    for j, kk in enumerate(("fast", "central", "slow"), start=VEL_ROWS[1] + 1):
        checks[f"sp_total_{kk}"] = (vs[f"E{j}"].value, sp[kk] / m.DAYS_WEEK)
    dates = {
        "baseline_date": (pr[f"B{b + 14}"].value, bl["finish"]),
        "mr_date": (pr[f"B{b + 15}"].value, bl["finish_mr"]),
        "p50_date": (pr[f"B{b + 13}"].value, bl["finish_p50"]),
    }
    for key, (a, e) in checks.items():
        good = close(a, e, 1e-6)
        ok = ok and good
        print(f"{key}: LibreOffice={a} Python={e} {'ok' if good else 'MISMATCH'}")
    for key, (a, e) in dates.items():
        good = a is not None and a.strftime("%Y-%m-%d") == e
        ok = ok and good
        print(f"{key}: LibreOffice={a} Python={e} {'ok' if good else 'MISMATCH'}")
    shutil.copy(tmp / path.name, OUT / "lab2_recalculated.xlsx")
    shutil.rmtree(tmp, ignore_errors=True)
    return ok


PRINT_VIEWS = {
    "WBS і CPM": ("wbs", ["B", "C", "X", "Y", "Z"], 0),
    "PERT і резерв": ("pert", ["J"], 0),
    "Backlog і velocity": ("velocity", ["F"], 0),
    "UCP": ("ucp", [], 0),
    "Календар": ("calendar", [], 0),
    "Порівняння": ("compare", [], 0),
    "Параметри ЛР2": ("params2", ["F", "B"], 0),
}


def render_sheets(path):
    import pymupdf
    from PIL import Image
    files = {}
    for title, (key, hidden, _) in PRINT_VIEWS.items():
        tmp = Path(tempfile.mkdtemp())
        wb = load_workbook(path)
        for ws in wb.worksheets:
            ws.sheet_state = "visible" if ws.title == title else "hidden"
        wb.active = wb.sheetnames.index(title)
        ws = wb[title]
        k.split_columns(ws)
        for col in hidden:
            ws.column_dimensions[col].hidden = True
        if key == "params2":
            ws.column_dimensions["B"].hidden = False
            ws.column_dimensions["B"].width = 34
            for r in range(1, ws.max_row + 1):
                ws.row_dimensions[r].height = 15
            for r in range(17, ws.max_row + 1):
                c = ws.cell(r, 2)
                if c.value and len(str(c.value)) > 30:
                    c.value = None
        if key == "ucp":
            for r in range(1, ws.max_row + 1):
                ws.row_dimensions[r].height = 15
        if key == "velocity":
            for r in range(1, ws.max_row + 1):
                ws.row_dimensions[r].height = 15
        if key == "wbs":
            for r in range(2, ws.max_row + 1):
                ws.row_dimensions[r].height = 15
        ws.sheet_properties.pageSetUpPr.fitToPage = False
        ws.page_setup.scale = 100
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.orientation = "landscape" if key in ("wbs", "ucp") else "portrait"
        src = tmp / f"view_{key}.xlsx"
        wb.save(src)
        k.soffice(src, "pdf", tmp)
        doc = pymupdf.open(tmp / f"view_{key}.pdf")
        pages = []
        for i, pg in enumerate(doc):
            pix = pg.get_pixmap(dpi=170)
            part = tmp / f"p{i}.png"
            pix.save(part)
            k.crop(part)
            pages.append(Image.open(part).convert("RGB"))
        width = max(im.width for im in pages)
        canvas = Image.new("RGB", (width, sum(im.height for im in pages)), "white")
        y = 0
        for im in pages:
            canvas.paste(im, (0, y))
            y += im.height
        out = OUT / f"sheet_{key}.png"
        canvas.save(out)
        files[key] = out
        shutil.rmtree(tmp, ignore_errors=True)
    return files


def to_date(day):
    return np.busday_offset(m.PROJECT_START, int(math.floor(day + 1e-9)), roll="forward", holidays=m.HOLIDAYS).astype("datetime64[D]").astype(date)


def to_date_end(day):
    return to_date(max(math.ceil(day - 1e-9) - 1, 0))


def chart_gantt():
    sched, end = m.cpm()
    bl = m.baseline()
    order = [w for w in m.WBS]
    fig, ax = plt.subplots(figsize=(10, 8.4), dpi=200)
    for i, w in enumerate(order):
        s = sched[w.wid]
        y = len(order) - 1 - i
        start = mdates.date2num(to_date(s["es"]))
        fin = mdates.date2num(to_date_end(s["ef"])) + 1
        col = COLORS[1] if s["critical"] else COLORS[0]
        ax.barh(y, fin - start, left=start, height=0.62, color=col)
        if s["float"] > 0.5:
            lf = mdates.date2num(to_date_end(s["lf"])) + 1
            ax.plot([fin, lf], [y, y], color=MUTED, linewidth=1.2)
            ax.plot([lf], [y], marker="|", color=MUTED, markersize=8)
    labels = [f"{w.wid} {w.name.split(':')[0]}" for w in order][::-1]
    ax.set_yticks(range(len(order)))
    ax.set_yticklabels(labels, fontsize=10)
    p50 = mdates.date2num(to_date_end(bl["p50"])) + 1
    blv = mdates.date2num(iso(bl["finish"]).date()) + 1
    mrv = mdates.date2num(iso(bl["finish_mr"]).date()) + 1
    ax.barh(-1.3, blv - p50, left=p50, height=0.62, color=COLORS[3])
    ax.barh(-2.3, mrv - blv, left=blv, height=0.62, color=COLORS[4])
    ax.text(p50 - 2, -1.3, "Contingency (у baseline)", ha="right", va="center", fontsize=11, color=INK)
    ax.text(blv - 2, -2.3, "Management reserve (поза baseline)", ha="right", va="center", fontsize=11, color=INK)
    for name, d in m.milestones():
        x = mdates.date2num(to_date_end(d)) + 1
        ax.plot([x], [len(order) + 0.4], marker="D", color=COLORS[2], markersize=7, clip_on=False)
        ax.axvline(x, color=COLORS[2], linewidth=0.7, linestyle=":")
    ms = m.milestones()
    offs = {"Scope freeze": (0, 0.6), "Architecture spike": (0, 1.4), "Feature complete": (0, 0.6), "Stabilization": (0, 1.4), "Release readiness": (0, 2.2)}
    for name, d in ms:
        x = mdates.date2num(to_date_end(d)) + 1
        ax.text(x, len(order) + 0.4 + offs[name][1], name, ha="center", fontsize=11, color=INK)
    ax.axvline(blv, color=COLORS[3], linewidth=1.4)
    ax.axvline(mrv, color=COLORS[4], linewidth=1.4)
    ax.text(blv + 1, -3.5, f"Baseline {iso(bl['finish']).strftime('%d.%m.%Y')}", fontsize=11, color=INK)
    ax.text(mrv + 1, -4.3, f"З резервом {iso(bl['finish_mr']).strftime('%d.%m.%Y')}", fontsize=11, color=INK)
    sprint0 = sched["1.2"]["ef"]
    vf = m.velocity_forecast()
    for j in range(vf["central"]["sprints"]):
        a = mdates.date2num(to_date(sprint0 + 10 * j))
        b = mdates.date2num(to_date(sprint0 + 10 * (j + 1)))
        if j % 2 == 0:
            ax.axvspan(a, b, color="#f4f3ee", zorder=0)
        ax.text((a + b) / 2, -0.85, f"S{j + 1}", ha="center", fontsize=10, color=MUTED)
    ax.set_ylim(-4.8, len(order) + 3.3)
    ax.xaxis_date()
    ax.xaxis.set_major_locator(mdates.MonthLocator())
    ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: UA_MONTHS[mdates.num2date(v).month - 1] + "\n" + str(mdates.num2date(v).year)))
    k.axes(ax)
    ax.grid(axis="y", visible=False)
    from matplotlib.patches import Patch
    from matplotlib.lines import Line2D
    handles = [Patch(color=COLORS[1], label="критичний шлях"), Patch(color=COLORS[0], label="некритичний пакет"),
               Line2D([0], [0], color=MUTED, label="резерв часу пакета"), Line2D([0], [0], marker="D", color=COLORS[2], linestyle="", label="віха"),
               Patch(color=COLORS[3], label="contingency"), Patch(color=COLORS[4], label="management reserve"),
               Patch(color="#f4f3ee", label="спринти центрального прогнозу")]
    ax.legend(handles=handles, loc="upper center", bbox_to_anchor=(0.35, -0.09), ncol=3, frameon=False, fontsize=10.5)
    ax.set_title("Календарний план MVP (тривалості E, ранній старт)", color=INK)
    fig.tight_layout()
    fig.savefig(OUT / "fig_gantt.png")
    plt.close(fig)


def chart_load():
    load = m.weekly_load()
    fig, axs = plt.subplots(len(m.ROLES), 1, figsize=(10, 8.6), dpi=200, sharex=True)
    for i, (ax, r) in enumerate(zip(axs, m.ROLES)):
        vals = load[r.rid]
        x = np.arange(1, len(vals) + 1)
        ax.bar(x, vals, color=COLORS[i], width=0.75)
        ax.axhline(r.hours_week, color=INK, linewidth=1, linestyle="--")
        ax.text(len(vals) + 0.7, r.hours_week, f"ємність {r.hours_week:.1f}".replace(".", ","), va="center", fontsize=10.5, color=INK)
        ax.set_ylabel(r.name, fontsize=11, rotation=0, ha="right", va="center")
        ax.set_ylim(0, r.hours_week * 1.25)
        ax.set_xlim(0.3, len(vals) + 4.2)
        k.axes(ax)
        ax.grid(axis="x", visible=False)
        ax.tick_params(labelsize=10)
    axs[-1].set_xlabel("Тиждень проєкту")
    axs[-1].set_xticks(range(1, len(load["BA"]) + 1, 2))
    axs[0].set_title("Тижневе навантаження ролей, продуктивні люд.-год (ранній старт, E)", color=INK)
    fig.tight_layout()
    fig.savefig(OUT / "fig_load.png")
    plt.close(fig)


def chart_velocity():
    vf = m.velocity_forecast()
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(10, 4.8), dpi=200)
    xs = [f"S{i}" for i in range(1, 7)]
    bars = a1.bar(xs, m.VELOCITY, color=COLORS[0], width=0.6)
    for b in bars:
        a1.text(b.get_x() + b.get_width() / 2, b.get_height() + 0.3, str(int(b.get_height())), ha="center", fontsize=9, color=INK)
    a1.axhline(vf["median"], color=COLORS[1], linewidth=1.6, label=f"медіана {vf['median']:.1f}".replace(".", ","))
    a1.axhline(max(m.VELOCITY), color=COLORS[2], linewidth=1, linestyle="--", label="max 21")
    a1.axhline(min(m.VELOCITY), color=COLORS[3], linewidth=1, linestyle="--", label="min 16")
    a1.set_ylim(0, 25)
    a1.set_ylabel("SP за спринт")
    a1.set_title("Історична velocity команди (варіант 14)")
    a1.legend(frameon=False, fontsize=11, loc="lower right")
    k.axes(a1)
    a1.grid(axis="x", visible=False)
    for i, (kk, lbl) in enumerate((("fast", "fast, 21 SP"), ("central", "central, 18,5 SP"), ("slow", "slow, 16 SP"))):
        v = vf[kk]["velocity"]
        n = vf[kk]["sprints"]
        xs2 = list(range(0, n + 1))
        ys = [min(v * j, m.BACKLOG_SP) for j in xs2]
        a2.plot(xs2, ys, marker="o", markersize=3.5, color=[COLORS[2], COLORS[1], COLORS[3]][i], linewidth=1.8, label=f"{lbl}: {n} спринтів")
    a2.axhline(m.BACKLOG_SP, color=INK, linewidth=1, linestyle=":")
    a2.text(0.2, m.BACKLOG_SP + 3, "backlog 116 SP", fontsize=11, color=INK)
    a2.set_xlabel("Спринт")
    a2.set_ylabel("Виконано SP наростаючим підсумком")
    a2.set_title("Прогноз завершення backlog")
    a2.set_ylim(0, 135)
    a2.set_xticks(range(0, 9))
    a2.legend(frameon=False, fontsize=11, loc="lower right")
    k.axes(a2)
    fig.tight_layout()
    fig.savefig(OUT / "fig_velocity.png")
    plt.close(fig)


def chart_mc():
    mc = m.monte_carlo()
    pp = m.p50_p80()
    bl = m.baseline()
    tot = mc["total"]
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(10, 6.0), dpi=200)
    a1.hist(tot, bins=50, color=COLORS[0], alpha=0.85)
    for v, c, lbl in ((pp["p50"], COLORS[1], "P50 аналітично"), (pp["p80"], COLORS[2], "P80 аналітично"), (mc["p80"], COLORS[3], "P80 Монте-Карло")):
        a1.axvline(v, color=c, linewidth=1.6, label=f"{lbl}: {v:.1f}".replace(".", ","))
    a1.set_xlabel("Тривалість проєкту, роб. днів")
    a1.set_ylabel("Кількість прогонів")
    a1.set_title("Розподіл тривалості, " + f"{mc['n']:,}".replace(",", " ") + " прогонів")
    a1.legend(loc="upper center", bbox_to_anchor=(0.5, -0.17), frameon=False, fontsize=10.5)
    k.axes(a1)
    xs = np.sort(tot)
    ys = np.arange(1, len(xs) + 1) / len(xs)
    a2.plot(xs, ys, color=COLORS[0], linewidth=2, label="Монте-Карло")
    from statistics import NormalDist
    nd = NormalDist(pp["mu"], pp["sigma"])
    gx = np.linspace(xs.min(), xs.max(), 200)
    a2.plot(gx, [nd.cdf(v) for v in gx], color=COLORS[1], linewidth=1.6, linestyle="--", label="нормальне наближення μ, σ шляху")
    a2.axhline(0.8, color=MUTED, linewidth=0.8, linestyle=":")
    a2.axvline(bl["days"], color=COLORS[3], linewidth=1.4, label=f"baseline {bl['days']} днів")
    a2.set_xlabel("Тривалість проєкту, роб. днів")
    a2.set_ylabel("Імовірність завершити не пізніше")
    a2.set_title("S-крива")
    a2.legend(loc="upper center", bbox_to_anchor=(0.5, -0.17), frameon=False, fontsize=10.5)
    k.axes(a2)
    fig.tight_layout()
    fig.savefig(OUT / "fig_mc.png")
    plt.close(fig)


def chart_methods():
    bl = m.baseline()
    sp = m.sp_phase_totals()
    uc = m.ucp_calendar()
    pp = m.p50_p80()
    mc = m.monte_carlo()
    rows = [
        ("PERT/WBS, P5...P95", float(np.percentile(mc["total"], 5)) / 5, pp["p50"] / 5, float(np.percentile(mc["total"], 95)) / 5),
        ("SP/velocity", sp["fast"] / 5, sp["central"] / 5, sp["slow"] / 5),
        ("UCP, PF ±20%", uc["duration_weeks"] * 0.8, uc["duration_weeks"], uc["duration_weeks"] * 1.2),
    ]
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(10, 4.8), dpi=200, gridspec_kw={"width_ratios": [1.35, 1]})
    for i, (lbl, lo, mid, hi) in enumerate(rows[::-1]):
        a1.plot([lo, hi], [i, i], color=COLORS[2 - i], linewidth=6, solid_capstyle="butt", alpha=0.55)
        a1.plot([mid], [i], marker="o", color=COLORS[2 - i], markersize=9)
        a1.text(hi + 0.4, i, f"від {lo:.1f} до {hi:.1f}".replace(".", ","), va="center", fontsize=11, color=INK)
    a1.axvline(bl["days"] / 5, color=COLORS[3], linewidth=1.3, linestyle="--")
    a1.text(bl["days"] / 5, 2.45, "baseline", ha="center", fontsize=11, color=INK)
    env = (bl["days"] + bl["mr"]) / 5
    a1.axvline(env, color=COLORS[4], linewidth=1.3, linestyle="--")
    a1.text(env, 2.75, "з резервом", ha="center", fontsize=11, color=INK)
    a1.set_yticks(range(3))
    a1.set_yticklabels([r[0] for r in rows[::-1]])
    a1.set_xlim(15, 50)
    a1.set_ylim(-0.6, 3.0)
    a1.set_xlabel("Тривалість, тижнів (точка: центральна оцінка)")
    a1.set_title("Тривалість за трьома методами")
    k.axes(a1)
    a1.grid(axis="y", visible=False)
    effs = [("WBS (E)", sum(w.effort_e for w in m.WBS)), ("PF 17,6", m.ucp(m.UCP_PF * 0.8)["effort"]),
            ("PF 22", m.ucp()["effort"]), ("PF 26,4", m.ucp(m.UCP_PF * 1.2)["effort"])]
    bars = a2.bar([e[0] for e in effs], [e[1] for e in effs], color=[COLORS[0], COLORS[1], COLORS[2], COLORS[3]], width=0.6)
    for b in bars:
        a2.text(b.get_x() + b.get_width() / 2, b.get_height() + 40, f"{b.get_height():,.0f}".replace(",", " "), ha="center", fontsize=11, color=INK)
    a2.set_ylabel("Продуктивні люд.-год")
    a2.set_title("Effort: WBS і UCP")
    a2.set_ylim(0, 3700)
    a2.tick_params(axis="x", labelsize=10)
    k.axes(a2)
    a2.grid(axis="x", visible=False)
    fig.tight_layout()
    fig.savefig(OUT / "fig_methods.png")
    plt.close(fig)


def node_label(w, s):
    name = "\\n".join(textwrap.wrap(w.name.split(":")[0], 20))
    nums = f"E {s['d']:.1f} | ES {s['es']:.1f} | EF {s['ef']:.1f}\\nрезерв {abs(s['float']) + 0.0:.1f}".replace(".", ",")
    return f"{w.wid} {name}\\n{nums}"


EXTERNAL = {"1.2": "карти регістрів\\nвід постачальника", "5.1": "постачання\\nлічильників", "5.4": "доступ до щитових,\\nпогодження власника"}


def diagram_network():
    sched, _ = m.cpm()
    lines = []
    for w in m.WBS:
        s = sched[w.wid]
        if s["critical"]:
            style = 'fillcolor="#fbe1d6", color="#eb6834", penwidth=2.2'
        else:
            style = 'fillcolor="#dce8f7", color="#2a78d6"'
        lines.append(f'n{w.wid.replace(".", "_")} [label="{node_label(w, s)}", {style}];')
    for wid, txt in EXTERNAL.items():
        lines.append(f'x{wid.replace(".", "_")} [label="{txt}", shape=note, fillcolor="#fff4d6", color="#eda100", fontsize=14];')
        lines.append(f'x{wid.replace(".", "_")} -> n{wid.replace(".", "_")} [style=dotted, color="#eda100", arrowhead=none];')
    for w in m.WBS:
        for pid, kind, lag in w.preds:
            crit = sched[pid]["critical"] and sched[w.wid]["critical"] and kind == "FS" and abs(sched[pid]["ef"] - sched[w.wid]["es"]) < 1e-9
            attrs = 'color="#eb6834", penwidth=2.4' if crit else 'color="#52514e"'
            if kind == "SS":
                attrs += f', style=dashed, label=" SS+{lag}", fontsize=14'
            lines.append(f'n{pid.replace(".", "_")} -> n{w.wid.replace(".", "_")} [{attrs}];')
    legend = '''subgraph cluster_leg { label="Легенда"; fontsize=15; color="#9a9890"; fontname="Arial";
l1 [label="критичний пакет", fillcolor="#fbe1d6", color="#eb6834", penwidth=2.2, fontsize=14];
l2 [label="пакет з резервом часу", fillcolor="#dce8f7", color="#2a78d6", fontsize=14];
l3 [label="зовнішнє очікування", shape=note, fillcolor="#fff4d6", color="#eda100", fontsize=14];
l4 [label="суцільна стрілка: FS\\nштрихова стрілка: SS\\nчисла: робочі дні від старту", shape=plaintext, style="", fontsize=14];
l1 -> l2 -> l3 -> l4 [style=invis]; {rank=same; l1; l2; l3; l4;} }'''
    src = f'''digraph G {{
rankdir=TB; nodesep=0.25; ranksep=0.42; newrank=true;
node [shape=box, style="rounded,filled", fontname="Arial", fontsize=15, margin="0.08,0.05"];
edge [arrowsize=0.7, fontname="Arial"];
{chr(10).join(lines)}
{legend}
{{rank=same; n2_1; n2_2; n1_3; n5_1;}}\nn5_6 -> l1 [style=invis];
}}'''
    p = OUT / "fig_network.dot"
    p.write_text(src, encoding="utf-8")
    subprocess.run(["dot", "-Tpng", "-Gdpi=200", str(p), "-o", str(OUT / "fig_network.png")], check=True)


def diagram_wbs():
    lines = ['root [label="MVP «ЕнергоПульс»\\n20 пакетів, 116 SP", fillcolor="#2a78d6", fontcolor="white", color="#2a78d6", fontsize=14];']
    ranks = []
    for gid, gname in m.WBS_GROUPS:
        glabel = "\\n".join(textwrap.wrap(gname, 16))
        lines.append(f'g{gid} [label="{gid} {glabel}", fillcolor="#dce8f7", color="#2a78d6", fontsize=14];')
        lines.append(f"root -> g{gid};")
        prev = f"g{gid}"
        for w in [x for x in m.WBS if x.wid.startswith(gid + ".")]:
            nid = "p" + w.wid.replace(".", "_")
            txt = "\\n".join(textwrap.wrap(w.name.split(":")[0], 17))
            lines.append(f'{nid} [label="{w.wid} {txt}\\n{m.ROLE[w.lead].name}", fillcolor="white", color="#52514e", fontsize=14];')
            lines.append(f"{prev} -> {nid} [style={'solid' if prev.startswith('g') else 'invis'}];")
            prev = nid
        ranks.append(f"g{gid}")
    src = f'''digraph W {{
rankdir=TB; nodesep=0.12; ranksep=0.22; splines=ortho;
node [shape=box, style="rounded,filled", fontname="Arial", width=1.9, margin="0.05,0.04"];
edge [color="#52514e", arrowhead=none];
{chr(10).join(lines)}
{{rank=same; {"; ".join(ranks)};}}
}}'''
    p = OUT / "fig_wbs.dot"
    p.write_text(src, encoding="utf-8")
    subprocess.run(["dot", "-Tpng", "-Gdpi=200", str(p), "-o", str(OUT / "fig_wbs.png")], check=True)


def main():
    OUT.mkdir(exist_ok=True)
    xlsx = OUT / "lab2_estimation_model.xlsx"
    build_workbook(xlsx)
    ok = verify(xlsx)
    print("workbook verified:", ok)
    render_sheets(OUT / "lab2_recalculated.xlsx")
    chart_gantt()
    chart_load()
    chart_velocity()
    chart_mc()
    chart_methods()
    diagram_network()
    diagram_wbs()
    if not ok:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
