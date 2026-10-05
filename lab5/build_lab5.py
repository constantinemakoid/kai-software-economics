import json
import shutil
import sys
import tempfile
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.patches import Patch, Polygon
from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT.parent))
sys.path.insert(0, str(ROOT.parent / "lab1"))
import econ_model as m
import build_lab1 as k

OUT = ROOT / "out"
plt.rcParams.update({"font.size": 12.5})
COLORS = k.COLORS
INK, MUTED, GRID = k.INK, k.MUTED, k.GRID
P5 = m.P5
LAYOUT = {}
NEG = '#,##0;"−"#,##0'


def col_widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def put_rows(wb, ws, rows, start, names=True):
    for i, r in enumerate(rows, start=start):
        key, label, value, unit, src, note = r
        ws.cell(i, 1, key)
        ws.cell(i, 2, label)
        ws.cell(i, 3, value)
        ws.cell(i, 4, unit)
        ws.cell(i, 5, src)
        ws.cell(i, 6, note)
        if names:
            k.name(wb, key, f"'{ws.title}'!$C${i}")
        if src in ("Д", "П"):
            ws.cell(i, 3).fill = k.INPUT_FILL
        fmt = unit_format(unit, value)
        ws.cell(i, 3).number_format = fmt
    return start + len(rows)


def unit_format(unit, value):
    if unit in ("частка",):
        return "0.00%"
    if unit in ("міс.", "клієнтів", "буд.", "раз"):
        return "0.00"
    if isinstance(value, int):
        return "#,##0"
    return "#,##0.00"


def sheet_params(wb):
    ws = wb.active
    ws.title = "Параметри ЛР5"
    ws.append(["Ім'я", "Показник", "Значення", "Одиниця", "Тип", "Підстава"])
    s = m.base_month_y3()
    rows = [
        ("x5_spend", "Acquisition spend за місяць", P5["acq_spend"], "EUR/міс.", "Д", "варіант 14, ЛР № 5"),
        ("x5_new", "Нові платні клієнти за місяць", P5["new_clients"], "клієнтів/міс.", "Д", "варіант 14, ЛР № 5"),
        ("x5_churn", "Monthly logo churn", P5["churn"], "частка", "Д", "варіант 14, ЛР № 5"),
        ("x5_fixed", "Fixed product OPEX за місяць", P5["fixed_opex"], "EUR/міс.", "Д", "варіант 14, ЛР № 5"),
        ("x5_scale", "Масштаб Y3, будівель", P4_SCALE, "будівель/міс.", "Д", "варіант 14, ЛР № 4 і № 5"),
        ("x5_rev", "Revenue базового місяця Y3", "=x5_rev_bu+x5_over", "EUR/міс.", "Р", "ЛР № 4, сценарій 100%: сума двох рядків нижче"),
        ("x5_rev_bu", "з неї база і будівлі (без overage)", s["base_rev"] + s["usage_rev"], "EUR/міс.", "Д", "ЛР № 4: 2 250 + 8 850"),
        ("x5_over", "з неї overage", s["over_rev"], "EUR/міс.", "Д", "ЛР № 4: 8,5 буд. × 55"),
        ("x5_cogs", "COGS базового місяця Y3", s["cogs"], "EUR/міс.", "Д", "ЛР № 4, з урахуванням A3"),
        ("x5_clients", "Активні клієнти (акаунти) Y3", s["clients"], "клієнтів", "Д", "ЛР № 4: 85 / 1,7"),
        ("x5_usage", "Будівель на клієнта", m.avg_usage(), "буд.", "Д", "ЛР № 4, A2"),
        ("x5_cogs_unit", "COGS будівлі", m.P4["cogs_unit"], "EUR/буд./міс.", "Д", "ЛР № 1, ЛР № 4"),
        ("x5_episode", "Епізод нестабільного обліку", m.P4["episode_months"], "міс.", "П", "A9 ЛР № 4"),
        ("x5_disc", "Річна ставка дисконтування", m.P3["discount"], "частка", "Д", "ЛР № 3"),
        ("x5_leads", "Кваліфіковані ліди за місяць", P5["leads"], "лідів/міс.", "П", "B1"),
        ("x5_act", "Activation rate: лід у пілот", P5["activation"], "частка", "П", "B2, A10 ЛР № 1"),
        ("x5_kept", "Пілоти, активні в 3-му місяці", P5["pilots_retained"], "пілотів", "П", "B4"),
        ("x5_pilot_m", "Тривалість пілоту", P5["pilot_months"], "міс.", "П", "A10 ЛР № 1"),
        ("x5_check", "Місяць відсіву пілотів", P5["pilot_check_month"], "міс.", "П", "B4"),
        ("x5_events", "Події і виїзні демонстрації", P5["spend_events"], "EUR/міс.", "П", "B5"),
        ("x5_pay", "Платежі і банк, частка Revenue", P5["other_var_payments"], "частка", "П", "B3"),
        ("x5_bad", "Резерв боргів, частка Revenue", P5["other_var_bad_debt"], "частка", "П", "B3"),
        ("x5_partner", "Партнерам, частка Revenue", P5["other_var_partner"], "частка", "П", "B3: 10% × 35%"),
        ("x5_cap", "Горизонт обмеженого LTV", P5["ltv_cap_months"], "міс.", "П", "3 роки, як TCO3 ЛР № 3"),
        ("x5_h", "Горизонт прогнозу", P5["horizon"], "міс.", "Д", "методичка, 3.3"),
        ("x5_delta", "Крок чутливості", P5["delta"], "частка", "Д", "методичка, крок 5"),
    ]
    end = put_rows(wb, ws, rows, 2)
    calc = [
        ("x5_ovshare", "Змінні витрати поза COGS, частка", "=x5_pay+x5_bad+x5_partner", "частка", "Р", "B3"),
        ("x5_dm", "Місячна ставка дисконтування", "=(1+x5_disc)^(1/12)-1", "частка", "Р", "(1 + 0,10)^(1/12) − 1"),
    ]
    end2 = put_rows(wb, ws, calc, end + 1)
    ws.cell(end + 1, 3).number_format = "0.00%"
    ws.cell(end + 2, 3).number_format = "0.0000%"
    k.style_table(ws, 1, 1, end2 - 1, 6)
    col_widths(ws, [12, 36, 12, 15, 6, 30])
    k.page(ws, landscape=False)


P4_SCALE = m.P4["scale_y3"]


def sheet_assumptions(wb):
    ws = wb.create_sheet("Журнал припущень ЛР5")
    ws.append(["ID", "Припущення", "Значення", "Підстава", "Що змінює"])
    for a in m.LR5_ASSUMPTIONS:
        ws.append(list(a))
    k.style_table(ws, 1, 1, len(m.LR5_ASSUMPTIONS) + 1, 5)
    col_widths(ws, [6, 24, 50, 40, 18])
    k.page(ws)


def sheet_funnel(wb):
    ws = wb.create_sheet("Воронка")
    ws.append(["Етап", "Подія", "Кількість за місяць", "Конверсія", "Джерело / припущення"])
    rows = [
        ("Acquisition", "кваліфікований лід: дзвінок з розрахунком економії відбувся", "=x5_leads", "", "B1 (П)"),
        ("Activation", "пілот запущено: шлюз передав першу добу 15-хв даних", "=C2*x5_act", "=C3/C2", "B2 (П), A10 ЛР № 1"),
        ("Retention", "пілот активний у 3-му місяці: реакція на алерт підтверджена", "=x5_kept", "=C4/C3", "B4 (П)"),
        ("Revenue", "перший оплачений рахунок (новий платний клієнт)", "=x5_new", "=C5/C4", "Д: варіант 14"),
    ]
    for r in rows:
        ws.append(list(r))
    for i in range(2, 6):
        ws.cell(i, 3).number_format = "0"
        ws.cell(i, 4).number_format = "0.0%"
    k.style_table(ws, 1, 1, 5, 5)
    t0 = 7
    ws.cell(t0, 1, "Показник")
    ws.cell(t0, 2, "Формула")
    ws.cell(t0, 3, "Значення")
    extra = [
        ("Конверсія лід у платного клієнта", "Revenue / Acquisition", "=C5/C2", "0.0%"),
        ("Конверсія пілот у платного клієнта", "Revenue / Activation", "=C5/C3", "0.0%"),
        ("Будівле-місяці пілотів", "пілоти × 1 + утримані × (3 − 1)", "=C3*x5_check+C4*(x5_pilot_m-x5_check)", "0"),
        ("Обслуговування пілотів, EUR", "буд.-міс. × COGS 16", f"=C{t0 + 3}*x5_cogs_unit", "#,##0.00"),
        ("Події і демонстрації, EUR", "B5", "=x5_events", "#,##0.00"),
        ("Цифровий маркетинг, EUR", "spend − пілоти − події", f"=x5_spend-C{t0 + 4}-C{t0 + 5}", "#,##0.00"),
        ("Acquisition spend, EUR", "сума", f"=C{t0 + 4}+C{t0 + 5}+C{t0 + 6}", "#,##0.00"),
        ("CPL, EUR за лід", "spend / ліди", "=x5_spend/C2", "#,##0.00"),
        ("Вартість пілоту, EUR", "spend / пілоти", "=x5_spend/C3", "#,##0.00"),
    ]
    for i, (a, b, c, f) in enumerate(extra, start=t0 + 1):
        ws.cell(i, 1, a)
        ws.cell(i, 2, b)
        ws.cell(i, 3, c)
        ws.cell(i, 3).number_format = f
    k.style_table(ws, t0, 1, t0 + len(extra), 3)
    LAYOUT["funnel_extra"] = t0
    col_widths(ws, [30, 44, 14, 11, 22])
    k.page(ws, landscape=False)


UE_ROWS = [
    ("x5_cac", "CAC", "acquisition spend / нові платні клієнти", "=x5_spend/x5_new", "EUR на нового клієнта, місяць", "#,##0.00"),
    ("x5_arpa", "ARPA", "Revenue / активні акаунти", "=x5_rev/x5_clients", "EUR на акаунт за місяць", "#,##0.00"),
    ("x5_arpu_b", "Довідково: Revenue на будівлю", "Revenue / будівлі (не змішується з ARPA)", "=x5_rev/x5_scale", "EUR на будівлю за місяць", "#,##0.00"),
    ("x5_cogs_pa", "COGS на акаунт", "COGS / активні акаунти", "=x5_cogs/x5_clients", "EUR на акаунт за місяць", "#,##0.00"),
    ("x5_gp", "Gross profit на акаунт", "ARPA − COGS на акаунт", "=x5_arpa-x5_cogs_pa", "EUR на акаунт за місяць", "#,##0.00"),
    ("x5_gm", "Gross margin", "(Revenue − COGS) / Revenue", "=(x5_rev-x5_cogs)/x5_rev", "частка", "0.00%"),
    ("x5_ov", "Інші змінні витрати на акаунт", "ARPA × 6%", "=x5_arpa*x5_ovshare", "EUR на акаунт за місяць", "#,##0.00"),
    ("x5_cm", "Contribution margin на акаунт", "ARPA − COGS − інші змінні", "=x5_arpa-x5_cogs_pa-x5_ov", "EUR на акаунт за місяць", "#,##0.00"),
    ("x5_cmr", "Contribution margin ratio", "CM / ARPA", "=x5_cm/x5_arpa", "частка", "0.00%"),
    ("x5_lost_n", "Втрачені клієнти за місяць", "churn × клієнти на початок місяця", "=x5_churn*x5_clients", "клієнтів", "0.00"),
    ("x5_lost", "Втрачений recurring revenue", "churn × MRR на початок (B7)", "=x5_churn*x5_rev", "EUR/міс.", "#,##0.00"),
    ("x5_contr", "Скорочення: завершені епізоди overage", "overage / 12", "=x5_over/x5_episode", "EUR/міс.", "#,##0.00"),
    ("x5_exp", "Expansion: нові епізоди overage", "дорівнює скороченню (B7)", "=x5_contr", "EUR/міс.", "#,##0.00"),
    ("x5_grc", "Gross revenue churn", "(втрачений + скорочений) / MRR", "=(x5_lost+x5_contr)/x5_rev", "частка за місяць", "0.000%"),
    ("x5_nrc", "Net revenue churn", "(втрачений + скорочений − expansion) / MRR", "=(x5_lost+x5_contr-x5_exp)/x5_rev", "частка за місяць", "0.000%"),
    ("x5_ltv", "Просте LTV", "ARPA × GM / logo churn", "=x5_arpa*x5_gm/x5_churn", "EUR на клієнта", "#,##0.00"),
    ("x5_ltv_cac", "LTV / CAC", "просте LTV / CAC", "=x5_ltv/x5_cac", "раз", "0.00"),
    ("x5_payback", "CAC payback", "CAC / gross profit нового клієнта", "=x5_cac/x5_gp", "міс.", "0.00"),
    ("x5_life", "Середній строк життя клієнта", "1 / churn", "=1/x5_churn", "міс.", "0.00"),
    ("x5_ltv_cm", "LTV за contribution margin", "CM / logo churn", "=x5_cm/x5_churn", "EUR на клієнта", "#,##0.00"),
    ("x5_ltv_rev", "LTV за gross revenue churn", "GP / gross revenue churn", "=x5_gp/x5_grc", "EUR на клієнта", "#,##0.00"),
    ("x5_ltv_disc", "LTV з дисконтуванням 10% річних", "GP × (1 + d) / (churn + d)", "=x5_gp*(1+x5_dm)/(x5_churn+x5_dm)", "EUR на клієнта", "#,##0.00"),
    ("x5_ltv_cap", "LTV за 36 міс.", "GP × (1 − (1 − churn)^36) / churn", "=x5_gp*(1-(1-x5_churn)^x5_cap)/x5_churn", "EUR на клієнта", "#,##0.00"),
    ("x5_payback_cm", "CAC payback за CM", "CAC / CM на акаунт", "=x5_cac/x5_cm", "міс.", "0.00"),
]


def sheet_ue(wb):
    ws = wb.create_sheet("Unit economics")
    ws.append(["Ім'я", "Метрика", "Формула", "Значення", "Період / одиниця"])
    for i, (key, label, note, f, unit, fmt) in enumerate(UE_ROWS, start=2):
        ws.append([key, label, note, f, unit])
        ws.cell(i, 4).number_format = fmt
        k.name(wb, key, f"'Unit economics'!$D${i}")
        LAYOUT["ue_" + key] = i
    k.style_table(ws, 1, 1, len(UE_ROWS) + 1, 5)
    col_widths(ws, [12, 34, 38, 12, 24])
    k.page(ws, landscape=False)


BE_ROWS = [
    ("x5_be_cm", "Contribution margin per unit", "CM на акаунт (одиниця BE: активний клієнт)", "=x5_cm", "#,##0.00"),
    ("x5_be_fixed", "Fixed product OPEX", "варіант 14, EUR на місяць", "=x5_fixed", "#,##0.00"),
    ("x5_be_acq", "Регулярний acquisition spend", "варіант 14, EUR на місяць, B10", "=x5_spend", "#,##0.00"),
    ("x5_be_p", "Product break-even units", "fixed OPEX / CM", "=x5_be_fixed/x5_be_cm", "0.00"),
    ("x5_be_p_int", "Product break-even, цілих клієнтів", "ROUNDUP", "=ROUNDUP(x5_be_p-1E-9,0)", "0"),
    ("x5_be_o", "Operating break-even units", "(fixed OPEX + acquisition spend) / CM", "=(x5_be_fixed+x5_be_acq)/x5_be_cm", "0.00"),
    ("x5_be_o_int", "Operating break-even, цілих клієнтів", "ROUNDUP", "=ROUNDUP(x5_be_o-1E-9,0)", "0"),
    ("x5_be_p_rev", "Product break-even revenue", "fixed OPEX / CM ratio", "=x5_be_fixed/x5_cmr", "#,##0.00"),
    ("x5_be_o_rev", "Operating break-even revenue", "(fixed OPEX + spend) / CM ratio", "=(x5_be_fixed+x5_be_acq)/x5_cmr", "#,##0.00"),
    ("x5_be_p_b", "Product break-even, будівель", "units × 1,7", "=x5_be_p*x5_usage", "0.00"),
    ("x5_be_o_b", "Operating break-even, будівель", "units × 1,7", "=x5_be_o*x5_usage", "0.00"),
    ("x5_be_mos", "Запас безпеки Y3 ЛР № 4 над product BE", "(клієнти Y3 − BE) / клієнти Y3", "=(x5_clients-x5_be_p)/x5_clients", "0.00%"),
    ("x5_be_mos_o", "Запас безпеки M12 над operating BE", "(клієнти M12 − BE) / клієнти M12", "=('Прогноз 12 міс.'!$F$13-x5_be_o)/'Прогноз 12 міс.'!$F$13", "0.00%"),
    ("x5_be_steady", "Стаціонарна база клієнтів", "нові / churn", "=x5_new/x5_churn", "0.00"),
]


def sheet_be(wb):
    ws = wb.create_sheet("Break-even")
    ws.append(["Ім'я", "Показник", "Формула / припущення", "Результат"])
    for i, (key, label, note, f, fmt) in enumerate(BE_ROWS, start=2):
        ws.append([key, label, note, f])
        ws.cell(i, 4).number_format = fmt
        k.name(wb, key, f"'Break-even'!$D${i}")
        LAYOUT["be_" + key] = i
    n = len(BE_ROWS) + 2
    ws.cell(n, 1, "x5_m_p")
    ws.cell(n, 2, "Перший місяць product BE")
    ws.cell(n, 3, "аркуш «Прогноз 12 міс.»")
    ws.cell(n, 4, "='Прогноз 12 міс.'!$F$16")
    ws.cell(n + 1, 1, "x5_m_o")
    ws.cell(n + 1, 2, "Перший місяць operating BE")
    ws.cell(n + 1, 3, "аркуш «Прогноз 12 міс.»")
    ws.cell(n + 1, 4, "='Прогноз 12 міс.'!$F$17")
    LAYOUT["be_month_row"] = n
    k.style_table(ws, 1, 1, n + 1, 4)
    col_widths(ws, [12, 36, 40, 13])
    k.page(ws, landscape=False)


FC_HEAD = ["Міс.", "Календарний місяць", "Клієнти на початок", "Нові", "Втрачені", "Клієнти на кінець", "Будівлі",
           "Revenue", "COGS", "Інші змінні", "Contribution margin", "Fixed OPEX", "Product результат",
           "Acquisition spend", "Operating результат", "Наростаючий operating", "Прапорець product", "Прапорець operating"]


def sheet_forecast(wb):
    ws = wb.create_sheet("Прогноз 12 міс.")
    ws.append(FC_HEAD)
    h = P5["horizon"]
    for i in range(h):
        r = i + 2
        prev = "0" if i == 0 else f"F{r - 1}"
        ws.append([
            i + 1, m.month_label(i), f"={prev}", "=x5_new", f"=C{r}*x5_churn", f"=C{r}-E{r}+D{r}", f"=F{r}*x5_usage",
            f"=F{r}*x5_arpa", f"=F{r}*x5_cogs_pa", f"=F{r}*x5_ov", f"=H{r}-I{r}-J{r}", "=x5_fixed", f"=K{r}-L{r}",
            "=x5_spend", f"=M{r}-N{r}", f"=O{r}" if i == 0 else f"=P{r - 1}+O{r}",
            f"=IF(K{r}>=L{r}-1E-9,A{r},99)", f"=IF(K{r}>=L{r}+N{r}-1E-9,A{r},99)",
        ])
        for c in range(3, 8):
            ws.cell(r, c).number_format = "0.00"
        for c in range(8, 17):
            ws.cell(r, c).number_format = NEG
    last = h + 1
    k.style_table(ws, 1, 1, last, 18)
    s0 = last + 2
    ws.cell(s0, 1, "Підсумок")
    summary = [
        ("Перший місяць product BE", f'=IF(MIN(Q2:Q{last})=99,"не досягнуто",MIN(Q2:Q{last}))', "0"),
        ("Перший місяць operating BE", f'=IF(MIN(R2:R{last})=99,"не досягнуто",MIN(R2:R{last}))', "0"),
        ("Клієнтів на кінець M12", f"=F{last}", "0.00"),
        ("Наростаючий operating результат M1...M12, EUR", f"=P{last}", NEG),
        ("Втрата трьох місяців пілотів до M1, EUR", "=3*(x5_fixed+x5_spend)", "#,##0"),
        ("Cost baseline MVP ЛР № 3, EUR", m.mvp_budget()["baseline"], "#,##0"),
        ("Залишок неокупленої інвестиції після M12, EUR", f"=F{s0 + 6}+F{s0 + 5}-F{s0 + 4}", "#,##0"),
    ]
    for i, (a, f, fmt) in enumerate(summary, start=s0 + 1):
        ws.cell(i, 1, a)
        ws.cell(i, 6, f)
        ws.cell(i, 6).number_format = fmt
    k.style_table(ws, s0, 1, s0 + len(summary), 6)
    for i in range(s0, s0 + len(summary) + 1):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=5)
        ws.cell(i, 6).alignment = Alignment(horizontal="right", vertical="top")
    ws.cell(s0, 6, "Значення")
    LAYOUT["fc_summary"] = s0
    ws.cell(1, 20, "Product BE")
    ws.cell(1, 21, "Operating BE")
    k.style_table(ws, 1, 20, last, 21)
    for r in range(2, last + 1):
        ws.cell(r, 20, "=x5_be_p")
        ws.cell(r, 21, "=x5_be_o")
        ws.cell(r, 20).number_format = "0.00"
        ws.cell(r, 21).number_format = "0.00"
    ch = LineChart()
    ch.title = "Клієнти на кінець місяця і межі break-even"
    ch.y_axis.title = "Активні клієнти"
    ch.x_axis.title = "Місяць прогнозу"
    for col, color in ((6, "2a78d6"), (20, "eb6834"), (21, "1baf7a")):
        ch.add_data(Reference(ws, min_col=col, min_row=1, max_row=last), titles_from_data=True)
        ch.series[-1].graphicalProperties.line.solidFill = color
        ch.series[-1].graphicalProperties.line.width = 28000
        ch.series[-1].smooth = False
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    ch.height = 8
    ch.width = 18
    ch.legend.position = "b"
    ws.add_chart(ch, f"H{s0 + 1}")
    col_widths(ws, [5, 13, 10, 6, 9, 10, 9, 10, 9, 9, 12, 9, 11, 11, 11, 14, 12.5, 12.5, 2, 10, 11])
    k.page(ws)


SENS = [
    ("churn", "Churn", [
        ("Monthly logo churn", lambda c, kk: f"=x5_churn*{kk}", "0.00%"),
        ("Просте LTV, EUR", lambda c, kk: f"=x5_arpa*x5_gm/{c}{{r0}}", "#,##0.00"),
        ("LTV / CAC", lambda c, kk: f"={c}{{r1}}/x5_cac", "0.00"),
        ("Стаціонарна база, клієнтів", lambda c, kk: f"=x5_new/{c}{{r0}}", "0.00"),
        ("Місяць product BE", lambda c, kk: month_formula("x5_be_p", f"{c}{{r0}}"), "0"),
        ("Місяць operating BE", lambda c, kk: month_formula("x5_be_o", f"{c}{{r0}}"), "0"),
    ]),
    ("price", "Ціна", [
        ("ARPA, EUR", lambda c, kk: f"=(x5_rev_bu+x5_over)*{kk}/x5_clients", "#,##0.00"),
        ("Gross margin", lambda c, kk: f"=({c}{{r0}}-x5_cogs_pa)/{c}{{r0}}", "0.00%"),
        ("CM на акаунт, EUR", lambda c, kk: f"={c}{{r0}}*(1-x5_ovshare)-x5_cogs_pa", "#,##0.00"),
        ("Просте LTV, EUR", lambda c, kk: f"={c}{{r0}}*{c}{{r1}}/x5_churn", "#,##0.00"),
        ("LTV / CAC", lambda c, kk: f"={c}{{r3}}/x5_cac", "0.00"),
        ("Product BE units", lambda c, kk: f"=x5_fixed/{c}{{r2}}", "0.00"),
        ("Operating BE units", lambda c, kk: f"=(x5_fixed+x5_spend)/{c}{{r2}}", "0.00"),
        ("Місяць product BE", lambda c, kk: month_formula(f"{c}{{r5}}", "x5_churn"), "0"),
        ("Місяць operating BE", lambda c, kk: month_formula(f"{c}{{r6}}", "x5_churn"), "0"),
    ]),
    ("cac", "CAC", [
        ("Acquisition spend, EUR", lambda c, kk: f"=x5_spend*{kk}", "#,##0.00"),
        ("CAC, EUR", lambda c, kk: f"={c}{{r0}}/x5_new", "#,##0.00"),
        ("LTV / CAC", lambda c, kk: f"=x5_ltv/{c}{{r1}}", "0.00"),
        ("CAC payback, міс.", lambda c, kk: f"={c}{{r1}}/x5_gp", "0.00"),
        ("Product BE units", lambda c, kk: "=x5_fixed/x5_cm", "0.00"),
        ("Operating BE units", lambda c, kk: f"=(x5_fixed+{c}{{r0}})/x5_cm", "0.00"),
        ("Місяць operating BE", lambda c, kk: month_formula(f"{c}{{r5}}", "x5_churn"), "0"),
    ]),
]


def month_formula(units, churn):
    t = f"ROUNDUP(LN(1-{units}*{churn}/x5_new)/LN(1-{churn})-1E-9,0)"
    return f'=IF({units}>=x5_new/{churn},"не досягнуто",IF({t}>x5_h,"не досягнуто",{t}))'


def sheet_sens(wb):
    ws = wb.create_sheet("Чутливість")
    ws.append(["Фактор / показник", "−20%", "База", "+20%", "Що змінилося"])
    ws.append(["Множник фактора", "=1-x5_delta", 1, "=1+x5_delta", "інші параметри незмінні"])
    for c in "BCD":
        ws[f"{c}2"].number_format = "0%"
        ws[f"{c}2"].fill = k.INPUT_FILL
    r = 3
    for key, title, lines in SENS:
        ws.cell(r, 1, title)
        ws.cell(r, 5, {"churn": "LTV/CAC; місяць BE", "price": "LTV/CAC; BE units", "cac": "LTV/CAC; payback"}[key])
        LAYOUT["sens_" + key] = r
        r0 = r + 1
        refs = {f"r{j}": r0 + j for j in range(len(lines))}
        for j, (label, f, fmt) in enumerate(lines):
            ws.cell(r0 + j, 1, label)
            for c in "BCD":
                ws[f"{c}{r0 + j}"] = f(c, f"{c}$2").format(**refs)
                ws[f"{c}{r0 + j}"].number_format = fmt
        r = r0 + len(lines)
    k.style_table(ws, 1, 1, r - 1, 5)
    for key, _, _ in SENS:
        ws.cell(LAYOUT["sens_" + key], 1).font = Font(name="Arial", size=10, bold=True)
    ws.cell(r + 1, 1, "Припущення для ціни: змінюються всі ціни прайсу (базовий платіж, ціна будівлі, ставка overage 55 EUR); COGS на акаунт 34,83 EUR незмінний; інші змінні витрати 6% від нової ARPA.")
    ws.cell(r + 2, 1, "Припущення для CAC: 7 нових клієнтів сталі, тому CAC змінюється разом з acquisition spend.")
    col_widths(ws, [34, 14, 14, 14, 22])
    k.page(ws, landscape=False)


def build_workbook(path):
    wb = Workbook()
    sheet_params(wb)
    sheet_assumptions(wb)
    sheet_funnel(wb)
    sheet_ue(wb)
    sheet_be(wb)
    sheet_forecast(wb)
    sheet_sens(wb)
    wb.save(path)


def close(a, b, tol=1e-6):
    return a is not None and not isinstance(a, str) and abs(a - b) < tol


def verify(path):
    tmp = Path(tempfile.mkdtemp())
    k.soffice(path, "xlsx", tmp)
    wb = load_workbook(tmp / path.name, data_only=True)
    ue, be, fc, sn, fu = wb["Unit economics"], wb["Break-even"], wb["Прогноз 12 міс."], wb["Чутливість"], wb["Воронка"]
    u = m.unit_economics()
    rc = m.revenue_churn()
    b = m.breakeven()
    f = m.forecast()
    fn = m.funnel()
    checks = {}
    pairs = {"x5_cac": u["cac"], "x5_arpa": u["arpa"], "x5_arpu_b": u["arpu_building"], "x5_cogs_pa": u["cogs_pa"],
             "x5_gp": u["gp_pa"], "x5_gm": u["gm"], "x5_ov": u["other_var"], "x5_cm": u["cm_pa"], "x5_cmr": u["cm_ratio"],
             "x5_lost_n": rc["lost_clients"], "x5_lost": rc["lost"], "x5_contr": rc["contraction"], "x5_grc": rc["gross"],
             "x5_nrc": rc["net"], "x5_ltv": u["ltv"], "x5_ltv_cac": u["ltv_cac"], "x5_payback": u["payback"],
             "x5_ltv_cm": u["ltv_cm"], "x5_ltv_rev": rc["ltv_rev"], "x5_ltv_disc": u["ltv_disc"], "x5_ltv_cap": u["ltv_cap"],
             "x5_payback_cm": u["payback_cm"]}
    for key, val in pairs.items():
        checks[key] = (ue[f"D{LAYOUT['ue_' + key]}"].value, val)
    for key, val in (("x5_be_p", b["product_units"]), ("x5_be_o", b["operating_units"]), ("x5_be_p_rev", b["product_revenue"]),
                     ("x5_be_o_rev", b["operating_revenue"]), ("x5_be_p_int", b["product_units_int"]),
                     ("x5_be_o_int", b["operating_units_int"]), ("x5_be_steady", f["steady"]),
                     ("x5_be_mos", b["mos_product_y3"]), ("x5_be_mos_o", f["mos_operating_m12"])):
        checks[key] = (be[f"D{LAYOUT['be_' + key]}"].value, val)
    for i, row in enumerate(f["rows"], start=2):
        checks[f"fc_cust_{i - 1}"] = (fc[f"F{i}"].value, row["customers"])
        checks[f"fc_cm_{i - 1}"] = (fc[f"K{i}"].value, row["cm"])
        checks[f"fc_cum_{i - 1}"] = (fc[f"P{i}"].value, row["cum_operating"])
    s0 = LAYOUT["fc_summary"]
    checks["fc_first_p"] = (fc[f"F{s0 + 1}"].value, f["first_product"])
    checks["fc_first_o"] = (fc[f"F{s0 + 2}"].value, f["first_operating"])
    checks["fc_unrecovered"] = (fc[f"F{s0 + 7}"].value, f["investment"] + f["presale_loss"] - f["cum_operating"])
    e = LAYOUT["funnel_extra"]
    checks["funnel_pilots"] = (fu["C3"].value, fn["pilots"])
    checks["funnel_pilot_cost"] = (fu[f"C{e + 4}"].value, fn["pilot_cost"])
    checks["funnel_marketing"] = (fu[f"C{e + 6}"].value, fn["marketing"])
    checks["funnel_spend"] = (fu[f"C{e + 7}"].value, P5["acq_spend"])
    _, _, sens = m.sensitivity5()
    texts = []
    for key, rows in (("churn", (("ltv_cac", 3), ("month_product", 5), ("month_operating", 6))),
                      ("price", (("ltv_cac", 5), ("product_units", 6), ("operating_units", 7), ("month_product", 8), ("month_operating", 9))),
                      ("cac", (("ltv_cac", 3), ("payback", 4), ("operating_units", 6), ("month_operating", 7)))):
        r = LAYOUT["sens_" + key]
        for j, s in zip("BCD", (-1, 0, 1)):
            for name, off in rows:
                want = sens[(key, s)][name]
                got = sn[f"{j}{r + off}"].value
                if want is None:
                    texts.append(got)
                else:
                    checks[f"sens_{key}_{name}_{s}"] = (got, want)
    ok = all(t == "не досягнуто" for t in texts)
    print("text checks:", texts)
    for key, (a, bb) in checks.items():
        good = close(a, bb, 1e-6)
        ok = ok and good
        print(f"{key}: LibreOffice={a} Python={bb} {'ok' if good else 'MISMATCH'}")
    print("checks:", len(checks) + len(texts))
    LAYOUT["n_checks"] = len(checks) + len(texts)
    LAYOUT["n_text"] = len(texts)
    (OUT / "verify.json").write_text(json.dumps({"checks": len(checks) + len(texts), "numeric": len(checks), "text": len(texts), "ok": ok}))
    shutil.copy(tmp / path.name, OUT / "lab5_recalculated.xlsx")
    shutil.rmtree(tmp, ignore_errors=True)
    return ok


PRINT_VIEWS = {
    "Unit economics": ("ue", "portrait", []),
    "Break-even": ("be", "portrait", []),
    "Прогноз 12 міс.": ("forecast", "landscape", ["S"]),
    "Чутливість": ("sens", "portrait", []),
    "Воронка": ("funnel", "portrait", []),
    "Параметри ЛР5": ("params", "portrait", []),
}


def render_sheets(path):
    import pymupdf
    from PIL import Image
    files = {}
    for title, (key, orient, hidden) in PRINT_VIEWS.items():
        tmp = Path(tempfile.mkdtemp())
        wb = load_workbook(path)
        for ws in wb.worksheets:
            ws.sheet_state = "visible" if ws.title == title else "hidden"
        wb.active = wb.sheetnames.index(title)
        ws = wb[title]
        k.split_columns(ws)
        for col in hidden:
            ws.column_dimensions[col].hidden = True
        for r in range(1, ws.max_row + 1):
            ws.row_dimensions[r].height = 16
        ws.row_dimensions[1].height = 30
        if key not in ("sens", "forecast"):
            for r in range(2, ws.max_row + 1):
                ratios = [len(str(ws.cell(r, c).value or "")) * 1.1 / (ws.column_dimensions[get_column_letter(c)].width or 10)
                          for c in range(1, min(ws.max_column, 6) + 1)
                          if isinstance(ws.cell(r, c).value, str) and not ws.cell(r, c).value.startswith("=")]
                longest = max(ratios, default=0)
                if longest > 1.0:
                    ws.row_dimensions[r].height = 15 * min(int(longest) + 1, 3)
        if key == "forecast":
            ws.row_dimensions[1].height = 44
        if key == "sens":
            ws.delete_rows(ws.max_row - 1, 2)
        ws.sheet_properties.pageSetUpPr.fitToPage = False
        ws.page_setup.scale = 100
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.orientation = orient
        src = tmp / f"view_{key}.xlsx"
        wb.save(src)
        k.soffice(src, "pdf", tmp)
        doc = pymupdf.open(tmp / f"view_{key}.pdf")
        pages = []
        for i, pg in enumerate(doc):
            pix = pg.get_pixmap(dpi=170)
            part = tmp / f"p{i}.png"
            pix.save(part)
            k.crop(part)
            pages.append(Image.open(part).convert("RGB"))
        width = max(im.width for im in pages)
        canvas = Image.new("RGB", (width, sum(im.height for im in pages)), "white")
        y = 0
        for im in pages:
            canvas.paste(im, (0, y))
            y += im.height
        out = OUT / f"sheet_{key}.png"
        canvas.save(out)
        files[key] = out
        shutil.rmtree(tmp, ignore_errors=True)
    return files


def eur(v, d=0):
    return f"{v:,.{d}f}".replace(",", " ").replace(".", ",")


def chart_funnel():
    fn = m.funnel()
    st = fn["stages"]
    events = ["кваліфікований лід:\nдзвінок з розрахунком економії",
              "пілот запущено:\nперша доба 15-хв даних",
              "пілот активний у 3-му місяці:\nреакція на алерт підтверджена",
              "перший оплачений рахунок:\nновий платний клієнт"]
    fig, ax = plt.subplots(figsize=(10.5, 5.6), dpi=200)
    top = st[0]["count"]
    h = 0.78
    for i, s in enumerate(st):
        w = 0.36 + 0.64 * s["count"] / top
        y = 3 - i
        nxt = st[i + 1]["count"] if i + 1 < len(st) else s["count"] * 0.8
        w2 = 0.36 + 0.64 * nxt / top
        poly = Polygon([(-w / 2, y + h / 2), (w / 2, y + h / 2), (w2 / 2, y - h / 2), (-w2 / 2, y - h / 2)],
                       closed=True, color=COLORS[i], alpha=0.92)
        ax.add_patch(poly)
        ax.text(0, y + 0.05, f"{s['stage']}: {eur(s['count'])}", ha="center", va="center", fontsize=14, color="white", fontweight="bold")
        ax.text(-0.62, y, events[i], ha="right", va="center", fontsize=11, color=INK)
        conv = "вхід воронки" if s["conv"] is None else f"конверсія {eur(s['conv'] * 100, 1)}%"
        ax.text(0.62, y + 0.1, conv, ha="left", va="center", fontsize=11.5, color=INK, fontweight="bold")
        src = ["B1 (П)", "B2 (П)", "B4 (П)", "варіант 14 (Д)"][i]
        ax.text(0.62, y - 0.2, src, ha="left", va="center", fontsize=10.5, color=MUTED)
    ax.set_xlim(-1.45, 1.05)
    ax.set_ylim(-0.6, 3.6)
    ax.axis("off")
    ax.set_title(f"Воронка базового місяця: лід у платного клієнта {eur(fn['lead_to_paid'] * 100, 1)}%, CPL {eur(fn['cpl'], 2)} EUR",
                 fontsize=13, color=INK)
    fig.tight_layout()
    fig.savefig(OUT / "fig_funnel.png")
    plt.close(fig)


def chart_forecast():
    f = m.forecast()
    rows = f["rows"]
    be = f["be"]
    x = np.arange(1, len(rows) + 1)
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(11.5, 4.9), dpi=200)
    cust = [r["customers"] for r in rows]
    a1.plot(x, cust, marker="o", color=COLORS[0], linewidth=2.2, label="клієнти на кінець місяця")
    a1.axhline(be["product_units"], color=COLORS[1], linewidth=1.8, label=f"product BE {eur(be['product_units'], 2)}")
    a1.axhline(be["operating_units"], color=COLORS[2], linewidth=1.8, label=f"operating BE {eur(be['operating_units'], 2)}")
    for mm, col in ((f["first_product"], COLORS[1]), (f["first_operating"], COLORS[2])):
        if mm:
            a1.plot([mm], [cust[mm - 1]], marker="o", markersize=11, markerfacecolor="none", markeredgecolor=col, markeredgewidth=2)
            a1.annotate(f"M{mm}: {eur(cust[mm - 1], 1)}", xy=(mm, cust[mm - 1]), xytext=(mm - 3.2, cust[mm - 1] + 6),
                        fontsize=10, color=INK, arrowprops={"arrowstyle": "->", "color": MUTED})
    a1.set_xticks(x)
    a1.set_xlabel("Місяць прогнозу (M1 = вересень 2027)")
    a1.set_ylabel("Активні клієнти")
    a1.set_ylim(0, 85)
    a1.set_title("База клієнтів і межі break-even", fontsize=12.5)
    a1.legend(loc="upper left", frameon=False, fontsize=10)
    k.axes(a1)
    cm = [r["cm"] for r in rows]
    a2.bar(x, cm, color=COLORS[0], width=0.65, label="contribution margin")
    a2.axhline(P5["fixed_opex"], color=COLORS[1], linewidth=1.8, label="fixed OPEX 9 000")
    a2.axhline(P5["fixed_opex"] + P5["acq_spend"], color=COLORS[2], linewidth=1.8, label="fixed OPEX + spend 12 400")
    a2.set_xticks(x)
    a2.set_xlabel("Місяць прогнозу")
    a2.set_ylabel("EUR за місяць")
    a2.set_ylim(0, 16500)
    a2.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: eur(v)))
    a2.set_title("Contribution margin і постійні витрати", fontsize=12.5)
    a2.legend(loc="upper left", frameon=False, fontsize=10)
    k.axes(a2)
    a2.grid(axis="x", visible=False)
    fig.tight_layout()
    fig.savefig(OUT / "fig_forecast.png")
    plt.close(fig)


def chart_sens():
    base, bbe, s = m.sensitivity5()
    b0 = base["ltv_cac"]
    labels = [("price", "Ціна ±20%"), ("churn", "Churn ±20%"), ("cac", "CAC ±20%")]
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(11.5, 4.6), dpi=200, gridspec_kw={"width_ratios": [1.15, 1]})
    for i, (key, lbl) in enumerate(labels[::-1]):
        lo, hi = s[(key, -1)]["ltv_cac"], s[(key, 1)]["ltv_cac"]
        a1.barh(i, lo - b0, left=b0, color=COLORS[0], height=0.55)
        a1.barh(i, hi - b0, left=b0, color=COLORS[1], height=0.55)
        for v in (lo, hi):
            a1.text(v + (0.15 if v > b0 else -0.15), i, eur(v, 2), va="center", ha="left" if v > b0 else "right", fontsize=10, color=INK)
    a1.set_yticks(range(3))
    a1.set_yticklabels([l for _, l in labels[::-1]])
    a1.axvline(b0, color=INK, linewidth=0.9)
    a1.set_xlim(8, 19.5)
    a1.set_xlabel("LTV / CAC")
    a1.set_title(f"LTV / CAC: база {eur(b0, 2)}", fontsize=12.5)
    k.axes(a1)
    a1.grid(axis="y", visible=False)
    a1.legend(handles=[Patch(color=COLORS[0], label="фактор мінус 20%"), Patch(color=COLORS[1], label="фактор плюс 20%")],
              loc="upper center", bbox_to_anchor=(0.5, -0.16), ncol=2, frameon=False, fontsize=10)
    x = np.arange(3)
    w = 0.36
    pu = [s[("price", j)]["product_units"] for j in (-1, 0, 1)]
    ou = [s[("price", j)]["operating_units"] for j in (-1, 0, 1)]
    a2.bar(x - w / 2, pu, width=w, color=COLORS[0], label="product BE units")
    a2.bar(x + w / 2, ou, width=w, color=COLORS[1], label="operating BE units")
    for j in range(3):
        a2.text(x[j] - w / 2, pu[j] - 2, eur(pu[j], 1), ha="center", va="top", fontsize=10, color="white", fontweight="bold", zorder=4)
        a2.text(x[j] + w / 2, ou[j] - 2, eur(ou[j], 1), ha="center", va="top", fontsize=10, color="white", fontweight="bold", zorder=4)
    m12 = m.forecast()["rows"][-1]["customers"]
    a2.axhline(m12, color=COLORS[2], linewidth=1.8, label=f"клієнти M12 {eur(m12, 1)}", zorder=3)
    a2.set_xticks(x)
    a2.set_xticklabels([f"ціна −20%\nARPA {eur(s[('price', -1)]['value'], 2)}", f"база\nARPA {eur(s[('price', 0)]['value'], 2)}",
                        f"ціна +20%\nARPA {eur(s[('price', 1)]['value'], 2)}"], fontsize=10)
    a2.set_ylim(0, 105)
    a2.set_ylabel("Клієнтів")
    a2.set_title("Break-even units за ціною", fontsize=12.5)
    a2.legend(loc="upper center", bbox_to_anchor=(0.5, -0.2), ncol=2, frameon=False, fontsize=10)
    k.axes(a2)
    a2.grid(axis="x", visible=False)
    fig.tight_layout()
    fig.savefig(OUT / "fig_sens.png")
    plt.close(fig)


def chart_ltv():
    u = m.unit_economics()
    rc = m.revenue_churn()
    items = [("Просте LTV\n(GM, logo churn)", u["ltv"]), ("LTV за CM", u["ltv_cm"]), ("LTV за gross\nrevenue churn", rc["ltv_rev"]),
             ("LTV з дисконтом\n10% річних", u["ltv_disc"]), ("LTV за 36 міс.", u["ltv_cap"])]
    fig, ax = plt.subplots(figsize=(10.5, 4.4), dpi=200)
    x = np.arange(len(items))
    ax.bar(x, [v for _, v in items], color=COLORS[0], width=0.6)
    for j, (_, v) in enumerate(items):
        ax.text(j, v + 120, f"{eur(v)}\nLTV/CAC {eur(v / u['cac'], 2)}", ha="center", fontsize=10, color=INK)
    ax.axhline(u["cac"], color=COLORS[1], linewidth=1.8)
    ax.text(4.45, u["cac"] + 120, f"CAC {eur(u['cac'], 2)}", ha="right", fontsize=10.5, color=INK)
    ax.axhline(3 * u["cac"], color=COLORS[3], linewidth=1.6)
    ax.text(4.45, 3 * u["cac"] + 120, f"3 × CAC {eur(3 * u['cac'])}", ha="right", fontsize=10.5, color=INK)
    ax.set_xticks(x)
    ax.set_xticklabels([n for n, _ in items], fontsize=10.5)
    ax.set_ylim(0, 8000)
    ax.set_ylabel("EUR на клієнта")
    ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: eur(v)))
    ax.set_title("Оцінки LTV одного клієнта за різних припущень", fontsize=12.5)
    k.axes(ax)
    ax.grid(axis="x", visible=False)
    fig.tight_layout()
    fig.savefig(OUT / "fig_ltv.png")
    plt.close(fig)


def main():
    OUT.mkdir(exist_ok=True)
    xlsx = OUT / "lab5_unit_economics.xlsx"
    build_workbook(xlsx)
    ok = verify(xlsx)
    print("workbook verified:", ok)
    render_sheets(OUT / "lab5_recalculated.xlsx")
    chart_funnel()
    chart_forecast()
    chart_sens()
    chart_ltv()
    if not ok:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
