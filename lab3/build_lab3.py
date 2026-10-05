import shutil
import sys
import tempfile
from datetime import datetime
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT.parent))
sys.path.insert(0, str(ROOT.parent / "lab1"))
import econ_model as m
import build_lab1 as k

OUT = ROOT / "out"
plt.rcParams.update({"font.size": 12.5})
COLORS = k.COLORS
INK, MUTED, GRID = k.INK, k.MUTED, k.GRID
PH = [p[0] for p in m.PHASE_NAMES]
ROLE_ROWS = {rid: 2 + i for i, (rid, _) in enumerate(m.LR3_ROLES)}
UA_MONTHS = ["січ", "лют", "бер", "кві", "тра", "чер", "лип", "сер", "вер", "жов", "лис", "гру"]
MCOL = [get_column_letter(2 + i) for i in range(m.LR3_MONTHS)]
LAYOUT = {}


def iso(s):
    return datetime.strptime(str(s), "%Y-%m-%d")


def col_widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def month_label(i):
    d = np.datetime64(m.PROJECT_START, "M") + i
    y, mo = int(str(d)[:4]), int(str(d)[5:7])
    return f"M{i + 1} {UA_MONTHS[mo - 1]} {str(y)[2:]}"


def sheet_params(wb):
    ws = wb.active
    ws.title = "Параметри ЛР3"
    ws.append(["Ім'я", "Показник", "Значення", "Одиниця", "Тип", "Пояснення"])
    bl = m.baseline()
    ph = m.phase_bounds()
    p = m.P3
    rows = [
        ("x_start", "Старт проєкту", iso(m.PROJECT_START), "дата", "Д", "ЛР № 2"),
        ("x_bl_days", "Baseline MVP", bl["days"], "роб. днів", "Д", "ЛР № 2, P80 з contingency"),
        ("x_ef11", "Кінець аналізу: EF пакета 1.1", ph[0][3], "роб. днів", "Р", "ЛР № 2, CPM"),
        ("x_ef12", "Кінець проєктування: EF пакета 1.2", ph[1][3], "роб. днів", "Р", "ЛР № 2, CPM"),
        ("x_fc", "Feature complete", ph[2][3], "роб. днів", "Р", "ЛР № 2, віха"),
        ("x_mr_days", "Management reserve за строком", m.MR_DAYS, "роб. днів", "Д", "ЛР № 2"),
        ("x_rate_BA", "Loaded rate BA/PM", p["rate_BA"], "EUR/міс.", "Д", "набір B"),
        ("x_rate_UX", "Loaded rate UX", p["rate_UX"], "EUR/міс.", "Д", "набір B"),
        ("x_rate_DEV", "Loaded rate Developer", p["rate_DEV"], "EUR/міс.", "Д", "набір B"),
        ("x_rate_QA", "Loaded rate QA", p["rate_QA"], "EUR/міс.", "Д", "набір B"),
        ("x_rate_OPS", "Loaded rate DevOps", p["rate_OPS"], "EUR/міс.", "Д", "набір B"),
        ("x_esv", "Ставка ЄСВ роботодавця", p["esv"], "частка", "Д", "закон про ЄСВ"),
        ("x_min_wage", "Мінімальна зарплата 2026", p["min_wage_uah"], "UAH", "Д", "закон про бюджет 2026"),
        ("x_cap_mult", "Максимальна база ЄСВ", p["esv_cap_mult"], "мін. зарплат", "Д", "закон про ЄСВ"),
        ("x_fx", "Курс", p["fx"], "UAH/EUR", "П", "лише для межі бази ЄСВ"),
        ("x_benefits", "Бенефіти на FTE", p["benefits"], "EUR/міс.", "П", "медичне страхування 50, навчання 30"),
        ("x_overhead", "Накладні на FTE", p["overhead"], "EUR/міс.", "П", "робоче місце 250, ноутбук 100, бухгалтерія і HR 200, офісне ПЗ 100"),
        ("x_days_avg", "Робочих днів у середньому місяці", p["days_month_avg"], "днів", "Р", "261 / 12"),
        ("x_r1_p", "R1: імовірність", p["r1_p"], "частка", "Д", "варіант 14"),
        ("x_r1_impact", "R1: фінансовий вплив", p["r1_impact"], "EUR", "Д", "варіант 14, 16 тис."),
        ("x_r1_month", "R1: місяць запланованої реакції", p["r1_month"], "M", "П", "після F3 на тест-стенді"),
        ("x_r2_p", "R2: імовірність", p["r2_p"], "частка", "П", "розрахунок через плинність"),
        ("x_r2_days", "R2: зсув фінішу", p["r2_days"], "роб. днів", "П", "передача знань і онбординг"),
        ("x_r2_turnover", "Річна плинність розробників", p["r2_turnover"], "частка", "П", ""),
        ("x_cloud_y1", "Cloud production Y1", p["cloud_y1"], "EUR/рік", "Д", "варіант 14"),
        ("x_g2", "Growth Y2", p["growth_y2"], "частка", "Д", "варіант 14"),
        ("x_g3", "Growth Y3", p["growth_y3"], "частка", "Д", "варіант 14"),
        ("x_alt_start", "Альтернатива: start", p["alt_start"], "EUR", "Д", "варіант 14"),
        ("x_alt_op", "Альтернатива: op", p["alt_op"], "EUR/рік", "Д", "варіант 14"),
        ("x_alt_month", "Місяць start-up альтернативи", p["alt_start_month"], "M", "П", "до релізу"),
        ("x_sup_ops", "Support: FTE DevOps", p["support_OPS"], "FTE", "П", ""),
        ("x_sup_ba", "Support: FTE BA/PM", p["support_BA"], "FTE", "П", ""),
        ("x_chg_dev", "Change: FTE Developer", p["change_DEV"], "FTE", "П", ""),
        ("x_chg_qa", "Change: FTE QA", p["change_QA"], "FTE", "П", ""),
        ("x_crisk_p", "Cloud risk: імовірність на рік", p["cloud_risk_p"], "частка", "П", ""),
        ("x_crisk_s", "Cloud risk: перевищення", p["cloud_risk_share"], "частка cloud", "П", ""),
        ("x_arisk_p", "Alt risk: імовірність на рік", p["alt_risk_p"], "частка", "П", ""),
        ("x_arisk_hw", "Alt risk: заміна вузла", p["alt_risk_hw"], "EUR", "П", ""),
        ("x_arisk_d", "Alt risk: днів DevOps", p["alt_risk_days"], "днів", "П", ""),
        ("x_ret_days", "Retirement: днів DevOps", p["retire_days"], "днів", "П", ""),
        ("x_ret_alt", "Retirement альтернативи: демонтаж", p["alt_retire_extra"], "EUR", "П", ""),
        ("x_disc", "Ставка дисконтування", p["discount"], "частка", "П", "довідково"),
    ]
    for i, r in enumerate(rows, start=2):
        ws.append(list(r))
        if r[4] in ("П", "Д"):
            ws.cell(i, 3).fill = k.INPUT_FILL
        k.name(wb, r[0], f"'Параметри ЛР3'!$C${i}")
        ws.cell(i, 3).number_format = "0.0000" if isinstance(r[2], float) and r[2] % 1 else "#,##0.##"
    ws.cell(2, 3).number_format = "DD.MM.YYYY"
    k.style_table(ws, 1, 1, len(rows) + 1, 6)
    top = len(rows) + 3
    ws.cell(top, 1, "Неробочі дні")
    for j, h in enumerate(m.HOLIDAYS):
        c = ws.cell(top + 1 + j, 1, iso(h))
        c.number_format = "DD.MM.YYYY"
        c.fill = k.INPUT_FILL
    k.name(wb, "x_holidays", f"'Параметри ЛР3'!$A${top + 1}:$A${top + len(m.HOLIDAYS)}")
    k.style_table(ws, top, 1, top + len(m.HOLIDAYS), 1)
    col_widths(ws, [13, 36, 12, 12, 6, 48])
    k.page(ws, landscape=False)


def sheet_assumptions(wb):
    ws = wb.create_sheet("Журнал припущень")
    ws.append(["ID", "Припущення", "Значення", "Тип", "Підстава або джерело"])
    for a in m.LR3_ASSUMPTIONS:
        ws.append([a[0], a[1], a[2], a[3], a[4]])
    k.style_table(ws, 1, 1, len(m.LR3_ASSUMPTIONS) + 1, 5)
    col_widths(ws, [6, 24, 60, 6, 44])
    k.page(ws)


def sheet_rates(wb):
    ws = wb.create_sheet("Ставки")
    ws.append(["Роль", "Loaded rate", "Накладні", "Бенефіти", "Оклад", "ЄСВ", "Сума складових", "Частка окладу", "Контроль"])
    for rid, name in m.LR3_ROLES:
        r = ROLE_ROWS[rid]
        rest = f"(B{r}-C{r}-D{r})"
        ws.append([name, f"=x_rate_{rid}", "=x_overhead", "=x_benefits",
                   f"=IF({rest}/(1+x_esv)<=$B$9,{rest}/(1+x_esv),{rest}-x_esv*$B$9)",
                   f"=x_esv*MIN(E{r},$B$9)", f"=C{r}+D{r}+E{r}+F{r}", f"=E{r}/B{r}",
                   f'=IF(ABS(G{r}-B{r})<0.01,"збігається","розбіжність")'])
        for c in range(2, 8):
            ws.cell(r, c).number_format = "#,##0.0"
        ws.cell(r, 8).number_format = "0.0%"
    k.style_table(ws, 1, 1, 6, 9)
    ws.cell(8, 1, "Показник")
    ws.cell(8, 2, "Значення")
    ws.cell(9, 1, "Максимальна база ЄСВ, EUR/міс.")
    ws.cell(9, 2, "=x_min_wage*x_cap_mult/x_fx")
    ws.cell(9, 2).number_format = "#,##0.0"
    k.style_table(ws, 8, 1, 9, 2)
    col_widths(ws, [30, 12, 11, 11, 11, 11, 15, 12, 13])
    k.page(ws)


def sheet_calendar(wb):
    ws = wb.create_sheet("Календар і фази")
    ws.append(["Фаза", "Початок, роб. день", "Кінець, роб. день", "Тривалість, роб. днів", "Тривалість, міс."])
    bounds = [("0", "x_ef11"), ("x_ef11", "x_ef12"), ("x_ef12", "x_fc"), ("x_fc", "x_bl_days")]
    first = 9
    last = first + m.LR3_MONTHS - 1
    pcols = ["G", "H", "I", "J"]
    for i, ((key, name), (a, b)) in enumerate(zip(m.PHASE_NAMES, bounds), start=2):
        ws.append([name, f"={a}", f"={b}", f"=C{i}-B{i}", f"=SUM({pcols[i - 2]}{first}:{pcols[i - 2]}{last})"])
        for c in (2, 3, 4):
            ws.cell(i, c).number_format = "0.00"
        ws.cell(i, 5).number_format = "0.0000"
        k.name(wb, f"x_ps_{key}", f"'Календар і фази'!$B${i}")
        k.name(wb, f"x_pe_{key}", f"'Календар і фази'!$C${i}")
        k.name(wb, f"x_pm_{key}", f"'Календар і фази'!$E${i}")
    ws.append(["Разом", "", "", "=SUM(D2:D5)", "=SUM(E2:E5)"])
    ws.cell(6, 4).number_format = "0.00"
    ws.cell(6, 5).number_format = "0.0000"
    k.name(wb, "x_months_total", "'Календар і фази'!$E$6")
    k.style_table(ws, 1, 1, 6, 5)
    hdr = ["Місяць", "Початок", "Кінець", "Роб. днів", "Наростом на початок", "Наростом на кінець",
           "Частка: аналіз", "Частка: проєктування", "Частка: розроблення", "Частка: тест і реліз", "Частка експлуатації"]
    for j, h in enumerate(hdr, start=1):
        ws.cell(first - 1, j, h)
    for i in range(m.LR3_MONTHS):
        r = first + i
        ws.cell(r, 1, month_label(i))
        ws.cell(r, 2, f"=EDATE(DATE(2026,11,1),{i})")
        ws.cell(r, 3, f"=EOMONTH(B{r},0)")
        ws.cell(r, 4, f"=NETWORKDAYS(MAX(B{r},x_start),C{r},x_holidays)")
        ws.cell(r, 5, 0 if i == 0 else f"=F{r - 1}")
        ws.cell(r, 6, f"=E{r}+D{r}")
        for key, col in zip(PH, pcols):
            ws[f"{col}{r}"] = f"=MAX(0,MIN(x_pe_{key},F{r})-MAX(x_ps_{key},E{r}))/D{r}"
            ws[f"{col}{r}"].number_format = "0.0000"
        ws.cell(r, 11, f"=IF(F{r}<=x_bl_days,0,IF(E{r}>=x_bl_days,1,(F{r}-x_bl_days)/D{r}))")
        ws.cell(r, 11).number_format = "0.0000"
        ws.cell(r, 2).number_format = "DD.MM.YYYY"
        ws.cell(r, 3).number_format = "DD.MM.YYYY"
    k.style_table(ws, first - 1, 1, last, 11)
    col_widths(ws, [16, 12, 12, 11, 12, 12, 11, 13, 12, 12, 12])
    k.page(ws)
    LAYOUT["cal_first"] = first


def sheet_labor(wb):
    ws = wb.create_sheet("Labor")
    hdr = ["Роль", "Loaded rate, EUR/міс.", "FTE: аналіз", "FTE: проєктування", "FTE: розроблення", "FTE: тест і реліз",
           "FTE-міс.: аналіз", "FTE-міс.: проєктування", "FTE-міс.: розроблення", "FTE-міс.: тест і реліз",
           "FTE-міс. разом", "Вартість, EUR", "FTE у ЛР № 2"]
    ws.append(hdr)
    for rid, name in m.LR3_ROLES:
        r = ROLE_ROWS[rid]
        vals = [name + (" (2 особи)" if rid == "DEV" else ""), f"=x_rate_{rid}"] + list(m.PHASE_FTE[rid])
        vals += [f"={c}{r}*x_pm_{key}" for c, key in zip("CDEF", PH)]
        vals += [f"=SUM(G{r}:J{r})", f"=B{r}*K{r}", m.LR2_FTE[rid]]
        ws.append(vals)
        for c in range(3, 7):
            ws.cell(r, c).fill = k.INPUT_FILL
        for c in range(7, 12):
            ws.cell(r, c).number_format = "0.000"
        ws.cell(r, 12).number_format = "#,##0.0"
    t = 2 + len(m.LR3_ROLES)
    ws.cell(t, 1, "Разом")
    for c in "CDEFGHIJKLM":
        ws[f"{c}{t}"] = f"=SUM({c}2:{c}{t - 1})"
    for c in range(7, 12):
        ws.cell(t, c).number_format = "0.000"
    ws.cell(t, 12).number_format = "#,##0.0"
    k.name(wb, "x_labor", f"Labor!$L${t}")
    k.style_table(ws, 1, 1, t, 13)
    c0 = t + 2
    ws.cell(c0, 1, "Вартість команди за повний місяць фази")
    for j, c in enumerate("CDEF"):
        ws[f"{c}{c0}"] = f"=SUMPRODUCT($B$2:$B${t - 1},{c}2:{c}{t - 1})"
        ws[f"{c}{c0}"].number_format = "#,##0"
    ws.cell(c0 + 1, 1, "Вартість фази, EUR")
    for c, g in zip("CDEF", "GHIJ"):
        ws[f"{c}{c0 + 1}"] = f"=SUMPRODUCT($B$2:$B${t - 1},{g}2:{g}{t - 1})"
        ws[f"{c}{c0 + 1}"].number_format = "#,##0"
    ws.cell(c0 + 2, 1, "Команда за FTE ЛР № 2, EUR/міс.")
    ws.cell(c0 + 2, 3, f"=SUMPRODUCT($B$2:$B${t - 1},$M$2:$M${t - 1})")
    ws.cell(c0 + 2, 3).number_format = "#,##0"
    k.name(wb, "x_team_month", f"Labor!$C${c0 + 2}")
    k.style_table(ws, c0, 1, c0 + 2, 6, header=False)
    col_widths(ws, [40, 11, 9, 13, 13, 10, 9, 13, 13, 10, 9, 11, 9])
    k.page(ws)
    LAYOUT["labor_total_row"] = t
    LAYOUT["phase_cost_row"] = c0


def sheet_baseline(wb):
    ws = wb.create_sheet("Baseline MVP")
    ws.append(["Стаття", "Позиція", "Місяць", "Сума, EUR"])
    items = [("Обладнання", n, mm, v) for n, mm, v in m.EQUIPMENT] + [("Ліцензії", n, mm, v) for n, mm, v in m.LICENSES] + \
            [("Cloud розроблення", "dev і stage, CI, тестовий брокер", i + 1, v) for i, v in enumerate(m.DEV_CLOUD)]
    for it in items:
        ws.append(list(it))
        ws.cell(ws.max_row, 4).fill = k.INPUT_FILL
    n = len(items) + 1
    k.style_table(ws, 1, 1, n, 4)
    rk0 = n + 2
    for j, h in enumerate(["ID", "Ризик", "Імовірність", "Вплив, EUR", "EMV, EUR", "Місяць реакції", "Довідково"], start=1):
        ws.cell(rk0, j, h)
    ws.cell(rk0 + 1, 1, "R1")
    ws.cell(rk0 + 1, 2, "Нестабільні лічильники (варіант)")
    ws.cell(rk0 + 1, 3, "=x_r1_p")
    ws.cell(rk0 + 1, 4, "=x_r1_impact")
    ws.cell(rk0 + 1, 6, "=x_r1_month")
    ws.cell(rk0 + 1, 7, "=x_team_month*x_mr_days/x_days_avg")
    ws.cell(rk0 + 2, 1, "R2")
    ws.cell(rk0 + 2, 2, "Вибуття розробника")
    ws.cell(rk0 + 2, 3, "=x_r2_p")
    ws.cell(rk0 + 2, 4, "=x_rate_DEV+x_team_month*x_r2_days/x_days_avg")
    ws.cell(rk0 + 2, 6, "немає")
    ws.cell(rk0 + 2, 7, "=1-(1-x_r2_turnover)^(2*x_months_total/12)")
    for r in (rk0 + 1, rk0 + 2):
        ws.cell(r, 5, f"=C{r}*D{r}")
        ws.cell(r, 3).number_format = "0%"
        for c in (4, 5):
            ws.cell(r, c).number_format = "#,##0.0"
    ws.cell(rk0 + 1, 7).number_format = "#,##0.0"
    ws.cell(rk0 + 2, 7).number_format = "0.0%"
    k.style_table(ws, rk0, 1, rk0 + 2, 7)
    b0 = rk0 + 4
    rows = [
        ("Loaded labor", "=x_labor", "ні, базова стаття", "аркуш Labor"),
        ("Обладнання", f'=SUMIF(A2:A{n},"Обладнання",D2:D{n})', "так, не входить у ставку", "варіант 14"),
        ("Ліцензії", f'=SUMIF(A2:A{n},"Ліцензії",D2:D{n})', "так, не входить у ставку", "варіант 14"),
        ("Cloud розроблення", f'=SUMIF(A2:A{n},"Cloud розроблення",D2:D{n})', "так", "варіант 14"),
        ("Base cost estimate MVP", f"=SUM(B{b0 + 1}:B{b0 + 4})", "", "labor + прямі"),
        ("Contingency", f"=SUM(E{rk0 + 1}:E{rk0 + 2})", "так", "EMV R1 + EMV R2"),
        ("Cost baseline MVP", f"=B{b0 + 5}+B{b0 + 6}", "", "base estimate + contingency"),
        ("Management reserve", f"=D{rk0 + 1}-E{rk0 + 1}", "окремо", "R1: вплив мінус EMV"),
        ("Project budget", f"=B{b0 + 7}+B{b0 + 8}", "", "cost baseline + MR"),
    ]
    for j, h in enumerate(["Стаття", "Сума, EUR", "Додаємо окремо?", "Підстава"], start=1):
        ws.cell(b0, j, h)
    for i, r in enumerate(rows, start=b0 + 1):
        for c, v in enumerate(r, start=1):
            ws.cell(i, c, v)
        ws.cell(i, 2).number_format = "#,##0.0"
    k.name(wb, "x_base", f"'Baseline MVP'!$B${b0 + 5}")
    k.name(wb, "x_cont", f"'Baseline MVP'!$B${b0 + 6}")
    k.name(wb, "x_cbl", f"'Baseline MVP'!$B${b0 + 7}")
    k.name(wb, "x_mr", f"'Baseline MVP'!$B${b0 + 8}")
    k.name(wb, "x_budget", f"'Baseline MVP'!$B${b0 + 9}")
    k.style_table(ws, b0, 1, b0 + len(rows), 4)
    col_widths(ws, [24, 50, 24, 26, 11, 12, 12])
    k.page(ws, landscape=False)
    LAYOUT.update({"items_n": n, "risk_row": rk0, "bl_row": b0})


def sheet_cash(wb):
    ws = wb.create_sheet("Cash burn")
    first = LAYOUT["cal_first"]
    n = LAYOUT["items_n"]
    rk0 = LAYOUT["risk_row"]
    pc = LAYOUT["phase_cost_row"]
    ws.cell(1, 1, "Стаття, EUR")
    for i, c in enumerate(MCOL):
        ws[f"{c}1"] = month_label(i)
    ws.cell(1, 14, "Разом")
    labels = ["Labor MVP", "Обладнання / ліцензії", "Cloud розроблення / operation", "Support і change після релізу",
              "Використання резерву", "Разом cash burn", "Наростаючим підсумком", "Альтернатива: разом cash burn"]
    for j, l in enumerate(labels, start=2):
        ws.cell(j, 1, l)
    for i, c in enumerate(MCOL):
        cr = first + i
        mn = i + 1
        ws[f"{c}2"] = f"=SUMPRODUCT('Календар і фази'!G{cr}:J{cr},Labor!$C${pc}:$F${pc})"
        ws[f"{c}3"] = f"=SUMIF('Baseline MVP'!$C$2:$C${n},{mn},'Baseline MVP'!$D$2:$D${n})-SUMIFS('Baseline MVP'!$D$2:$D${n},'Baseline MVP'!$C$2:$C${n},{mn},'Baseline MVP'!$A$2:$A${n},\"Cloud розроблення\")"
        ws[f"{c}4"] = f"=SUMIFS('Baseline MVP'!$D$2:$D${n},'Baseline MVP'!$C$2:$C${n},{mn},'Baseline MVP'!$A$2:$A${n},\"Cloud розроблення\")+x_cloud_y1/12*'Календар і фази'!K{cr}"
        ws[f"{c}5"] = f"=(TCO!$H$4+TCO!$H$5)/12*'Календар і фази'!K{cr}"
        ws[f"{c}6"] = f"=SUMIF('Baseline MVP'!$F${rk0 + 1}:$F${rk0 + 2},{mn},'Baseline MVP'!$E${rk0 + 1}:$E${rk0 + 2})"
        ws[f"{c}7"] = f"=SUM({c}2:{c}6)"
        ws[f"{c}8"] = f"={c}7" if i == 0 else f"={MCOL[i - 1]}8+{c}7"
        ws[f"{c}9"] = f"={c}7-x_cloud_y1/12*'Календар і фази'!K{cr}+x_alt_op/12*'Календар і фази'!K{cr}+IF({mn}=x_alt_month,x_alt_start,0)"
    for r in (2, 3, 4, 5, 6, 7, 9):
        ws.cell(r, 14, f"=SUM(B{r}:M{r})")
    for r in range(2, 10):
        for c in range(2, 15):
            ws.cell(r, c).number_format = "#,##0"
    k.style_table(ws, 1, 1, 9, 14)
    ws.cell(11, 1, "Показник")
    ws.cell(11, 2, "Значення")
    ws.cell(12, 1, "Contingency, EUR")
    ws.cell(12, 2, "=x_cont")
    ws.cell(13, 1, "Використано в плані M1...M12, EUR")
    ws.cell(13, 2, "=N6")
    ws.cell(14, 1, "Залишок contingency, EUR")
    ws.cell(14, 2, "=B12-B13")
    ws.cell(15, 1, "Management reserve (поза baseline), EUR")
    ws.cell(15, 2, "=x_mr")
    ws.cell(16, 1, "Місяць релізу")
    ws.cell(16, 2, '="M7, "&TEXT(WORKDAY(x_start,x_bl_days-1,x_holidays),"DD.MM.YYYY")')
    for r in range(12, 16):
        ws.cell(r, 2).number_format = "#,##0.0"
    k.style_table(ws, 11, 1, 16, 2)
    col_widths(ws, [42, 12] + [10] * 11 + [11])
    k.page(ws)


def sheet_tco(wb):
    ws = wb.create_sheet("TCO")
    ws.append(["Компонент", "Cloud Y1", "Cloud Y2", "Cloud Y3", "Альтернатива Y1", "Альтернатива Y2", "Альтернатива Y3", "Річна база", "Cloud Σ3", "Альтернатива Σ3"])
    ws.append(["MVP / start-up", "=x_cbl", 0, 0, "=x_cbl+x_alt_start", 0, 0])
    ws.append(["Operation", "=x_cloud_y1", "=B3*(1+x_g2)", "=C3*(1+x_g3)", "=x_alt_op", "=x_alt_op", "=x_alt_op"])
    ws.append(["Support", "=$H$4", "=$H$4", "=$H$4", "=$H$4", "=$H$4", "=$H$4", "=12*(x_sup_ops*x_rate_OPS+x_sup_ba*x_rate_BA)"])
    ws.append(["Change", "=$H$5", "=$H$5", "=$H$5", "=$H$5", "=$H$5", "=$H$5", "=12*(x_chg_dev*x_rate_DEV+x_chg_qa*x_rate_QA)"])
    alt_risk = "x_arisk_p*(x_arisk_hw+x_arisk_d*x_rate_OPS/x_days_avg)"
    ws.append(["Risk після релізу", "=x_crisk_p*x_crisk_s*B3", "=x_crisk_p*x_crisk_s*C3", "=x_crisk_p*x_crisk_s*D3", f"={alt_risk}", f"={alt_risk}", f"={alt_risk}"])
    ws.append(["Change і risk", "=B5+B6", "=C5+C6", "=D5+D6", "=E5+E6", "=F5+F6", "=G5+G6"])
    ws.append(["Retirement", 0, 0, "=x_ret_days*x_rate_OPS/x_days_avg", 0, 0, "=x_ret_days*x_rate_OPS/x_days_avg+x_ret_alt"])
    ws.append(["Разом за рік", "=B2+B3+B4+B7+B8", "=C2+C3+C4+C7+C8", "=D2+D3+D4+D7+D8", "=E2+E3+E4+E7+E8", "=F2+F3+F4+F7+F8", "=G2+G3+G4+G7+G8"])
    for r in range(2, 11):
        if r == 10:
            break
        ws.cell(r, 9, f"=SUM(B{r}:D{r})")
        ws.cell(r, 10, f"=SUM(E{r}:G{r})")
    ws.append(["TCO3", "", "", "=SUM(B9:D9)", "", "", "=SUM(E9:G9)", "", "=I9", "=J9"])
    ws.append(["Різниця: альтернатива мінус cloud", "", "", "", "", "", "", "", "", "=J10-I10"])
    ws.append(["PV при ставці дисконтування", "", "", "=B9+C9/(1+x_disc)+D9/(1+x_disc)^2", "", "", "=E9+F9/(1+x_disc)+G9/(1+x_disc)^2"])
    for r in range(2, 13):
        for c in range(2, 11):
            ws.cell(r, c).number_format = "#,##0"
    k.style_table(ws, 1, 1, 12, 10)
    k.name(wb, "x_tco_cloud", "TCO!$D$10")
    k.name(wb, "x_tco_alt", "TCO!$G$10")
    s0 = 14
    for j, h in enumerate(["Cloud Y1, множник", "Cloud Y1, EUR", "TCO3 cloud", "TCO3 альтернатива", "Різниця", "Рішення"], start=1):
        ws.cell(s0, j, h)
    ws.cell(s0 + 4, 1, "Беззбитковий множник")
    for i, kk in enumerate((0.8, 1.0, 1.2), start=s0 + 1):
        ws.cell(i, 1, kk)
        ws.cell(i, 2, f"=x_cloud_y1*A{i}")
        ws.cell(i, 3, f"=x_tco_cloud+(A{i}-1)*(I3+I6)")
        ws.cell(i, 4, "=x_tco_alt")
        ws.cell(i, 5, f"=D{i}-C{i}")
        ws.cell(i, 6, f'=IF(E{i}>0,"cloud","альтернатива")')
        for c in (2, 3, 4, 5):
            ws.cell(i, c).number_format = "#,##0"
    ws.cell(s0 + 4, 2, "=(x_tco_alt-(x_tco_cloud-(I3+I6)))/(I3+I6)")
    ws.cell(s0 + 4, 2).number_format = "0.000"
    ws.cell(s0 + 4, 3, "=x_cloud_y1*B18")
    ws.cell(s0 + 4, 3).number_format = "#,##0"
    k.style_table(ws, s0, 1, s0 + 4, 6)
    col_widths(ws, [32, 13, 13, 11, 14, 14, 14, 11, 11, 14])
    k.page(ws)


def build_workbook(path):
    wb = Workbook()
    sheet_params(wb)
    sheet_assumptions(wb)
    sheet_rates(wb)
    sheet_calendar(wb)
    sheet_labor(wb)
    sheet_baseline(wb)
    sheet_cash(wb)
    sheet_tco(wb)
    wb.save(path)


def close(a, b, tol=1e-6):
    return a is not None and not isinstance(a, str) and abs(a - b) < tol


def verify(path):
    tmp = Path(tempfile.mkdtemp())
    k.soffice(path, "xlsx", tmp)
    wb = load_workbook(tmp / path.name, data_only=True)
    lp = m.labor_plan()
    b = m.mvp_budget()
    cb = m.cash_burn()
    cba = m.cash_burn(scenario="alt")
    t = m.tco()
    rk = m.risks()
    bl0 = LAYOUT["bl_row"]
    r0 = LAYOUT["risk_row"]
    lt = LAYOUT["labor_total_row"]
    bs = wb["Baseline MVP"]
    cs = wb["Cash burn"]
    ts = wb["TCO"]
    ls = wb["Labor"]
    checks = {
        "labor": (ls[f"L{lt}"].value, lp["total"]),
        "base": (bs[f"B{bl0 + 5}"].value, b["base"]),
        "contingency": (bs[f"B{bl0 + 6}"].value, b["contingency"]),
        "cost_baseline": (bs[f"B{bl0 + 7}"].value, b["baseline"]),
        "mr": (bs[f"B{bl0 + 8}"].value, b["mr"]),
        "budget": (bs[f"B{bl0 + 9}"].value, b["budget"]),
        "r2_impact": (bs[f"D{r0 + 2}"].value, rk[1]["impact"]),
        "r2_p_calc": (bs[f"G{r0 + 2}"].value, rk[1]["p_calc"]),
        "r1_team": (bs[f"G{r0 + 1}"].value, rk[0]["team_part"]),
        "cash_sum": (cs["N7"].value, cb["sum"]),
        "cash_alt_sum": (cs["N9"].value, sum(cba["total"])),
        "tco_cloud": (ts["D10"].value, t["cloud"]["tco3"]),
        "tco_alt": (ts["G10"].value, t["alt"]["tco3"]),
        "diff": (ts["J11"].value, t["diff"]),
        "pv_cloud": (ts["D12"].value, t["cloud"]["pv"]),
        "pv_alt": (ts["G12"].value, t["alt"]["pv"]),
        "breakeven": (ts["B18"].value, m.breakeven_cloud_multiplier()),
    }
    for rid, _ in m.LR3_ROLES:
        checks[f"cost_{rid}"] = (ls[f"L{ROLE_ROWS[rid]}"].value, lp["roles"][rid]["cost"])
        checks[f"gross_{rid}"] = (wb["Ставки"][f"E{ROLE_ROWS[rid]}"].value, m.rate_parts(rid)["gross"])
    for i, c in enumerate(MCOL):
        checks[f"cash_M{i + 1}"] = (cs[f"{c}7"].value, cb["total"][i])
        checks[f"alt_M{i + 1}"] = (cs[f"{c}9"].value, cba["total"][i])
    for kk, row in enumerate(("A", "D", "B", "T"), start=2):
        checks[f"months_{row}"] = (wb["Календар і фази"][f"E{kk}"].value, lp["months"][row])
    ok = True
    for key, (a, e) in checks.items():
        good = close(a, e, 1e-4)
        ok = ok and good
        print(f"{key}: LibreOffice={a} Python={e} {'ok' if good else 'MISMATCH'}")
    shutil.copy(tmp / path.name, OUT / "lab3_recalculated.xlsx")
    shutil.rmtree(tmp, ignore_errors=True)
    return ok


PRINT_VIEWS = {
    "Ставки": ("rates", [], "portrait"),
    "Labor": ("labor", [], "landscape"),
    "Baseline MVP": ("baseline", [], "landscape"),
    "Cash burn": ("cash", [], "landscape"),
    "TCO": ("tco", [], "landscape"),
    "Календар і фази": ("calendar", [], "landscape"),
}


def render_sheets(path):
    import pymupdf
    from PIL import Image
    files = {}
    for title, (key, hidden, orient) in PRINT_VIEWS.items():
        tmp = Path(tempfile.mkdtemp())
        wb = load_workbook(path)
        for ws in wb.worksheets:
            ws.sheet_state = "visible" if ws.title == title else "hidden"
        wb.active = wb.sheetnames.index(title)
        ws = wb[title]
        k.split_columns(ws)
        for col in hidden:
            ws.column_dimensions[col].hidden = True
        for r in range(1, ws.max_row + 1):
            ws.row_dimensions[r].height = 15
        heads = {"rates": [1, 8], "labor": [1], "baseline": [1, LAYOUT["risk_row"], LAYOUT["bl_row"]], "cash": [1, 11],
                 "tco": [1, 14], "calendar": [1, LAYOUT["cal_first"] - 1]}[key]
        for r in heads:
            ws.row_dimensions[r].height = 32
        if key == "baseline":
            for r in range(1, ws.max_row + 1):
                if ws.cell(r, 2).value and len(str(ws.cell(r, 2).value)) > 50:
                    ws.row_dimensions[r].height = 28
        ws.sheet_properties.pageSetUpPr.fitToPage = False
        ws.page_setup.scale = 100
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.orientation = orient
        src = tmp / f"view_{key}.xlsx"
        wb.save(src)
        k.soffice(src, "pdf", tmp)
        doc = pymupdf.open(tmp / f"view_{key}.pdf")
        pages = []
        for i, pg in enumerate(doc):
            pix = pg.get_pixmap(dpi=170)
            part = tmp / f"p{i}.png"
            pix.save(part)
            k.crop(part)
            pages.append(Image.open(part).convert("RGB"))
        width = max(im.width for im in pages)
        canvas = Image.new("RGB", (width, sum(im.height for im in pages)), "white")
        y = 0
        for im in pages:
            canvas.paste(im, (0, y))
            y += im.height
        out = OUT / f"sheet_{key}.png"
        canvas.save(out)
        files[key] = out
        shutil.rmtree(tmp, ignore_errors=True)
    return files


def eur(v):
    return f"{v:,.0f}".replace(",", " ")


def chart_rates():
    names = [n for rid, n in m.LR3_ROLES]
    parts = [m.rate_parts(rid) for rid, _ in m.LR3_ROLES]
    keys = [("gross", "оклад"), ("esv", "ЄСВ 22%"), ("benefits", "бенефіти"), ("overhead", "накладні")]
    fig, ax = plt.subplots(figsize=(10, 4.4), dpi=200)
    left = np.zeros(len(names))
    for i, (kk, lbl) in enumerate(keys):
        vals = np.array([p[kk] for p in parts])
        ax.barh(names, vals, left=left, color=COLORS[i], label=lbl, height=0.6)
        for j, v in enumerate(vals):
            if v > 350:
                ax.text(left[j] + v / 2, j, eur(v), ha="center", va="center", fontsize=10, color="white")
        left += vals
    for j, v in enumerate(left):
        ax.text(v + 60, j, eur(v), va="center", fontsize=11, color=INK)
    ax.invert_yaxis()
    ax.set_xlim(0, 7000)
    ax.set_xlabel("EUR на місяць за 1 FTE")
    ax.xaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: eur(v)))
    ax.set_title("Склад loaded monthly rate, набір B")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.2), ncol=4, frameon=False)
    k.axes(ax)
    ax.grid(axis="y", visible=False)
    fig.tight_layout()
    fig.savefig(OUT / "fig_rates.png")
    plt.close(fig)


def chart_labor():
    lp = m.labor_plan()
    roles = [(rid, n) for rid, n in m.LR3_ROLES if rid != "UX"]
    phases = [n for _, n in m.PHASE_NAMES]
    fig, ax = plt.subplots(figsize=(10, 4.6), dpi=200)
    bottom = np.zeros(4)
    for i, (rid, n) in enumerate(roles):
        vals = np.array(lp["roles"][rid]["phase_cost"])
        ax.bar(phases, vals, bottom=bottom, color=COLORS[i], label=n, width=0.6)
        bottom += vals
    for j, v in enumerate(bottom):
        ax.text(j, v + 1200, eur(v), ha="center", fontsize=11, color=INK)
    ax.set_ylim(0, max(bottom) * 1.15)
    ax.set_ylabel("EUR")
    ax.set_title("Labor cost MVP за фазами і ролями")
    ax.legend(loc="upper left", frameon=False)
    ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: eur(v)))
    k.axes(ax)
    ax.grid(axis="x", visible=False)
    fig.tight_layout()
    fig.savefig(OUT / "fig_labor.png")
    plt.close(fig)


def chart_cash():
    cb = m.cash_burn()
    alt = m.cash_burn(scenario="alt")
    labels = [month_label(i).replace(" ", "\n", 1) for i in range(m.LR3_MONTHS)]
    x = np.arange(m.LR3_MONTHS)
    keys = [("labor", "labor MVP"), ("eq_lic", "обладнання і ліцензії"), ("cloud_op", "cloud розроблення і operation"),
            ("post", "support і change після релізу"), ("reserve", "використання contingency")]
    fig, ax = plt.subplots(figsize=(10, 5.4), dpi=200)
    bottom = np.zeros(m.LR3_MONTHS)
    for i, (kk, lbl) in enumerate(keys):
        vals = np.array(cb[kk])
        ax.bar(x, vals, bottom=bottom, color=COLORS[i], label=lbl, width=0.68)
        bottom += vals
    for j, v in enumerate(bottom):
        ax.text(j, v + 1600, f"{v / 1000:.1f}".replace(".", ","), ha="center", fontsize=10, color=INK)
    ax.plot(x, alt["total"], marker="D", linestyle="none", color=INK, markersize=5, label="альтернатива: разом за місяць")
    ax.text(5, alt["total"][5] + 1500, f"start-up {eur(m.P3['alt_start'])}", ha="center", fontsize=10, color=INK)
    rel = 6 + (1 - cb["release_frac"][6]) - 0.5
    ax.axvline(rel, color=MUTED, linestyle=":", linewidth=1.2)
    ax.text(rel + 0.1, 66000, "реліз 06.05.2027", fontsize=10.5, color=INK)
    ax.set_xticks(x)
    ax.set_xticklabels(labels, fontsize=10)
    ax.set_ylim(0, 72000)
    ax.set_ylabel("EUR за місяць (над стовпцями: тис. EUR)")
    ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: eur(v)))
    ax.set_title("Cash burn першого року, сценарій cloud, EUR")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.17), ncol=3, frameon=False, fontsize=10.5)
    k.axes(ax)
    ax.grid(axis="x", visible=False)
    fig.tight_layout()
    fig.savefig(OUT / "fig_cash.png")
    plt.close(fig)


def chart_tco():
    t = m.tco()
    b = m.mvp_budget()
    comp = [("MVP baseline", lambda s: b["baseline"]), ("start-up", lambda s: t[s]["start"]),
            ("operation", lambda s: t[s]["op_sum"]), ("support і change", lambda s: 3 * (m.annual_support() + m.annual_change())),
            ("risk і retirement", lambda s: t[s]["risk_sum"] + t[s]["retire"])]
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(10, 5.2), dpi=200, gridspec_kw={"width_ratios": [1, 1.25]})
    names = ["Cloud", "Альтернатива"]
    left = np.zeros(2)
    cols = COLORS
    for i, (lbl, f) in enumerate(comp):
        vals = np.array([f("cloud"), f("alt")])
        a1.bar(names, vals, bottom=left, color=cols[i], label=lbl, width=0.55)
        left += vals
    for j, v in enumerate(left):
        a1.text(j, v + 6000, eur(v), ha="center", fontsize=11, color=INK)
    a1.set_ylim(0, 540000)
    a1.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: f"{v / 1000:.0f}"))
    a1.set_ylabel("тис. EUR")
    a1.set_title("TCO3 за компонентами")
    a1.legend(loc="upper center", bbox_to_anchor=(0.5, -0.08), ncol=2, frameon=False, fontsize=10)
    k.axes(a1)
    a1.grid(axis="x", visible=False)
    for i, s in enumerate(("cloud", "alt")):
        cum = np.cumsum([0] + [y["total"] for y in t[s]["years"]])
        a2.plot(["старт Y1", "кінець Y1", "кінець Y2", "кінець Y3"], cum / 1000, marker="o", color=COLORS[i * 1], linewidth=2, label=names[i])
        a2.text(3.08, cum[-1] / 1000 + (6 if s == "alt" else -14), f"{cum[-1] / 1000:.1f}".replace(".", ","), fontsize=11, color=INK)
    a2.set_xlim(-0.2, 3.6)
    a2.set_ylabel("тис. EUR наростаючим підсумком")
    a2.set_title("Накопичені витрати за горизонт")
    a2.legend(loc="upper left", frameon=False)
    k.axes(a2)
    fig.tight_layout()
    fig.savefig(OUT / "fig_tco.png")
    plt.close(fig)


def chart_sens():
    base, rows = m.sensitivity()
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(10, 4.8), dpi=200)
    b0 = base["cloud"]["tco3"] / 1000
    order = sorted([r for r in rows if abs(r["cloud_hi"] - r["cloud_lo"]) > 1], key=lambda r: abs(r["cloud_hi"] - r["cloud_lo"]))
    for i, r in enumerate(order):
        lo, hi = r["cloud_lo"] / 1000 - b0, r["cloud_hi"] / 1000 - b0
        a1.barh(i, lo, color=COLORS[0], height=0.6)
        a1.barh(i, hi, color=COLORS[1], height=0.6)
    a1.set_yticks(range(len(order)))
    a1.set_yticklabels([r["label"] for r in order], fontsize=10)
    a1.axvline(0, color=INK, linewidth=0.8)
    a1.set_xlabel("Зміна TCO3 cloud, тис. EUR")
    a1.set_title("TCO3 cloud: " + f"{b0:.1f}".replace(".", ",") + " тис. EUR", fontsize=12.5)
    k.axes(a1)
    a1.grid(axis="y", visible=False)
    d0 = base["diff"] / 1000
    order2 = sorted([r for r in rows if abs(r["diff_hi"] - r["diff_lo"]) > 1], key=lambda r: abs(r["diff_hi"] - r["diff_lo"]))
    for i, r in enumerate(order2):
        a2.barh(i, r["diff_lo"] / 1000 - d0, left=d0, color=COLORS[0], height=0.6)
        a2.barh(i, r["diff_hi"] / 1000 - d0, left=d0, color=COLORS[1], height=0.6)
    a2.set_yticks(range(len(order2)))
    a2.set_yticklabels([r["label"] for r in order2], fontsize=10)
    a2.axvline(d0, color=INK, linewidth=0.8)
    a2.axvline(0, color=COLORS[3], linewidth=1.4, linestyle="--")
    a2.set_xlim(-5, 55)
    a2.set_xlabel("Альтернатива мінус cloud, тис. EUR")
    a2.set_title("Перевага cloud: " + f"{d0:.1f}".replace(".", ",") + " тис. EUR", fontsize=12.5)
    k.axes(a2)
    a2.grid(axis="y", visible=False)
    from matplotlib.patches import Patch
    fig.legend(handles=[Patch(color=COLORS[0], label="параметр мінус 20%"), Patch(color=COLORS[1], label="параметр плюс 20%")],
               loc="lower center", ncol=2, frameon=False, fontsize=10.5)
    fig.tight_layout(rect=(0, 0.07, 1, 1))
    fig.savefig(OUT / "fig_sens.png")
    plt.close(fig)


def main():
    OUT.mkdir(exist_ok=True)
    xlsx = OUT / "lab3_budget_tco.xlsx"
    build_workbook(xlsx)
    ok = verify(xlsx)
    print("workbook verified:", ok)
    render_sheets(OUT / "lab3_recalculated.xlsx")
    chart_rates()
    chart_labor()
    chart_cash()
    chart_tco()
    chart_sens()
    if not ok:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
